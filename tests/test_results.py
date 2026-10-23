import asyncio
import json
import sqlite3
import tempfile
import unittest
from datetime import datetime, timedelta
from io import BytesIO
from contextlib import contextmanager
from pathlib import Path
from unittest.mock import patch

from fastapi import BackgroundTasks, HTTPException

from backend.main import (
    app,
    _record_discovered_documents,
    _crawl_content_types,
    _next_monitoring_run,
    _claim_scheduled_crawls,
    _claim_topic_summary_schedule,
    create_review,
    create_source,
    crawl_schedule_status,
    delete_review,
    discover_crawl_documents,
    download_document_ai_analysis,
    download_document_pdf,
    generate_document_ai_analysis,
    generate_poc_document_summary,
    get_document_ai_analysis,
    get_document_summary,
    get_product_class_analysis,
    get_results_workflow,
    generate_summary,
    list_documents_view,
    list_findings,
    list_new_information_documents,
    get_new_information_document,
    list_results,
    list_reviews,
    list_source_documents,
    list_sources,
    export_topic_assessments,
    topic_summary_schedule,
    monitoring_worker,
    process_selected_document,
    run_manual_crawl,
    run_product_class_analysis,
    selected_processing_jobs,
    start_selected_document_processing,
    startup,
    test_source_endpoint,
    update_source,
    update_review,
)
from backend.ollama_client import build_summary_prompt, generate_document_summary
from openpyxl import load_workbook
from backend.models import (
    CrawlDiscoverRequest,
    CrawlProcessRequest,
    DocumentAIAnalysisRequest,
    ReviewCreate,
    ReviewUpdate,
    SourceCreate,
    SourceTestRequest,
    SourceUpdate,
)


class ResultsPersistenceTestCase(unittest.TestCase):
    def setUp(self):
        self.connection = sqlite3.connect(":memory:")
        self.connection.row_factory = sqlite3.Row
        self.connection.executescript(
            """
            CREATE TABLE sources (
                id INTEGER PRIMARY KEY,
                reference_no TEXT,
                authority TEXT,
                name TEXT,
                source_type TEXT,
                region TEXT,
                monitoring_frequency TEXT,
                website TEXT,
                safety_alert_type TEXT,
                url TEXT,
                content_types TEXT,
                instruction TEXT,
                active INTEGER,
                last_crawled_at TEXT,
                last_crawl_status TEXT,
                last_error_message TEXT,
                created_at TEXT,
                updated_at TEXT
            );
            CREATE TABLE documents (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                source_id INTEGER NOT NULL,
                title TEXT,
                url TEXT NOT NULL UNIQUE,
                meeting_date TEXT,
                publication_date TEXT,
                filename TEXT,
                discovered_date TEXT,
                last_crawled_at TEXT,
                sha256 TEXT,
                content_type TEXT,
                reference_id TEXT UNIQUE,
                status TEXT NOT NULL DEFAULT 'NEW',
                created_at TEXT NOT NULL
            );
            CREATE TABLE crawl_history (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                source_id INTEGER,
                started_at TEXT,
                completed_at TEXT,
                status TEXT,
                documents_found INTEGER DEFAULT 0,
                new_documents INTEGER,
                total_documents INTEGER,
                error_message TEXT
            );
            CREATE TABLE document_pages (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                document_id INTEGER NOT NULL,
                page_number INTEGER NOT NULL,
                page_text TEXT NOT NULL DEFAULT ''
            );
            CREATE TABLE extraction_history (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                document_id INTEGER NOT NULL,
                started_at TEXT,
                completed_at TEXT,
                status TEXT,
                pages_extracted INTEGER,
                error_message TEXT
            );
            CREATE TABLE reviews (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                finding_id INTEGER NOT NULL,
                stage TEXT NOT NULL DEFAULT 'INITIAL',
                reviewer TEXT,
                decision TEXT,
                comments TEXT,
                reviewed_at TEXT
            );
            CREATE TABLE findings (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                document_id INTEGER NOT NULL,
                product_name TEXT,
                active_substance TEXT,
                trigger_name TEXT,
                safety_topic TEXT,
                relevant_terms TEXT,
                evidence TEXT,
                section TEXT,
                page_number INTEGER,
                created_at TEXT,
                comparison_status TEXT
            );
            CREATE TABLE product_classes (
                id INTEGER PRIMARY KEY,
                reference_no TEXT,
                class_name TEXT NOT NULL,
                active INTEGER NOT NULL DEFAULT 1
            );
            CREATE TABLE products (
                id INTEGER PRIMARY KEY,
                product_class_id INTEGER NOT NULL,
                product_name TEXT NOT NULL,
                active INTEGER NOT NULL DEFAULT 1
            );
            CREATE TABLE summaries (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                finding_id INTEGER NOT NULL,
                product_name TEXT,
                active_substance TEXT,
                safety_topic TEXT,
                prac_assessment TEXT,
                outcome TEXT,
                summary TEXT,
                model TEXT,
                created_at TEXT
            );
            CREATE TABLE document_ai_analyses (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                document_id INTEGER NOT NULL,
                provider TEXT NOT NULL,
                analysis_type TEXT NOT NULL,
                summary TEXT,
                status TEXT NOT NULL,
                error_message TEXT,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                UNIQUE(document_id, provider, analysis_type)
            );
            CREATE TABLE audit_events (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                created_at TEXT NOT NULL,
                event_type TEXT NOT NULL,
                details TEXT NOT NULL,
                entity_type TEXT,
                entity_id INTEGER,
                source_id INTEGER,
                reference_id TEXT,
                action TEXT,
                ai_model TEXT,
                result TEXT,
                status TEXT NOT NULL DEFAULT 'COMPLETED'
            );
            CREATE TABLE product_class_ai_analyses (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                product_class_id INTEGER NOT NULL,
                provider TEXT NOT NULL,
                substance TEXT NOT NULL DEFAULT '',
                summary TEXT,
                status TEXT NOT NULL,
                error_message TEXT,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                UNIQUE(product_class_id, provider, substance)
            );
            INSERT INTO sources
                (id, reference_no, authority, name, url, content_types, active)
            VALUES (
                1,
                'EMA_PRAC_001',
                'EMA',
                'EMA PRAC',
                'https://example.test/prac',
                '["Meeting Minutes", "Agendas", "Highlights"]',
                1
            );
            """
        )

        @contextmanager
        def open_test_database():
            yield self.connection

        self.get_db_patch = patch(
            "backend.main.get_db",
            open_test_database,
        )
        self.get_db_patch.start()

    def tearDown(self):
        self.get_db_patch.stop()
        self.connection.close()

    def _add_new_information_document(self, document_id, title, product, date):
        self.connection.execute(
            """
            INSERT INTO documents
                (id, source_id, title, url, publication_date, filename,
                 discovered_date, content_type, reference_id, status, created_at)
            VALUES (?, 1, ?, ?, ?, ?, ?, ?, ?, 'PROCESSED', ?)
            """,
            (
                document_id,
                title,
                f"https://example.test/{document_id}.pdf",
                date,
                f"{document_id}.pdf",
                date,
                "Meeting Minutes",
                f"EMA_PRAC_001_{document_id}",
                date,
            ),
        )
        self.connection.execute(
            """
            INSERT INTO findings
                (document_id, product_name, trigger_name, safety_topic,
                 evidence, page_number, created_at, comparison_status)
            VALUES (?, ?, 'Safety signal', 'Adverse reaction', 'Exact quote', 2, ?, 'NEW')
            """,
            (document_id, product, date),
        )
        self.connection.commit()

    def test_initial_review_must_precede_final_review(self):
        self._add_new_information_document(1, "Safety update", "Alprazolam", "2026-09-25")

        with self.assertRaises(HTTPException) as error:
            create_review(ReviewCreate(
                finding_id=1,
                stage="FINAL",
                decision="RELEVANT",
            ))
        self.assertEqual(error.exception.status_code, 409)

        initial = create_review(ReviewCreate(
            finding_id=1,
            stage="INITIAL",
            decision="RELEVANT",
            reviewer="Initial reviewer",
        ))
        self.assertEqual(
            list_findings(
                limit=1000,
                document_id=None,
                trigger_name=None,
                comparison_status=None,
                review_status=None,
            )[0]["review_status"],
            "AWAITING_FINAL_REVIEW",
        )
        final = create_review(ReviewCreate(
            finding_id=1,
            stage="FINAL",
            decision="ESCALATE",
            reviewer="Final reviewer",
        ))
        self.assertTrue(initial["success"])
        self.assertTrue(final["success"])
        finding_state = list_findings(
            limit=1000,
            document_id=None,
            trigger_name=None,
            comparison_status=None,
            review_status=None,
        )[0]
        self.assertEqual(finding_state["review_status"], "REVIEWED")
        self.assertEqual(finding_state["initial_review"]["stage"], "INITIAL")
        self.assertEqual(finding_state["final_review"]["stage"], "FINAL")
        self.assertEqual([review["stage"] for review in list_reviews()], ["FINAL", "INITIAL"])

        with self.assertRaises(HTTPException) as error:
            delete_review(initial["id"])
        self.assertEqual(error.exception.status_code, 409)

    def test_final_review_cannot_be_edited_into_initial_after_finalization(self):
        self._add_new_information_document(1, "Safety update", "Alprazolam", "2026-09-25")
        initial = create_review(ReviewCreate(
            finding_id=1,
            stage="INITIAL",
            decision="RELEVANT",
        ))
        create_review(ReviewCreate(
            finding_id=1,
            stage="FINAL",
            decision="RELEVANT",
        ))

        with self.assertRaises(HTTPException) as error:
            update_review(initial["id"], ReviewUpdate(
                stage="INITIAL",
                decision="NOT_RELEVANT",
            ))
        self.assertEqual(error.exception.status_code, 409)

    def test_topic_export_contains_assessment_and_both_review_stages(self):
        self._add_new_information_document(1, "Safety update", "Alprazolam", "2026-09-25")
        self.connection.execute(
            "INSERT INTO product_classes (id, class_name, active) VALUES (1, 'Benzodiazepines', 1)"
        )
        self.connection.execute(
            "INSERT INTO products (id, product_class_id, product_name, active) VALUES (1, 1, 'Alprazolam', 1)"
        )
        self.connection.execute(
            """
            INSERT INTO summaries
                (finding_id, summary, prac_assessment, outcome, model, created_at)
            VALUES (1, 'Concise topic', 'Class impact; consider PI review', 'Review update', 'local-model', '2026-09-26')
            """
        )
        self.connection.commit()
        create_review(ReviewCreate(finding_id=1, stage="INITIAL", decision="RELEVANT"))
        create_review(ReviewCreate(finding_id=1, stage="FINAL", decision="ESCALATE"))

        response = export_topic_assessments()
        workbook = load_workbook(BytesIO(response.body), read_only=True)
        sheet = workbook["Topic Assessments"]
        headers = [cell.value for cell in next(sheet.iter_rows())]
        row = dict(zip(headers, next(sheet.iter_rows(min_row=2, values_only=True))))

        self.assertEqual(response.media_type, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertEqual(row["Concise Topic Summary"], "Concise topic")
        self.assertEqual(row["Configured Product Class"], "Benzodiazepines")
        self.assertEqual(row["Initial Decision"], "RELEVANT")
        self.assertEqual(row["Final Decision"], "ESCALATE")
        self.assertIn("Review History", workbook.sheetnames)

    def test_alternate_day_schedule_claim_persists_cadence(self):
        first_day = datetime(2026, 1, 5, 9, 0)
        first_claim = _claim_topic_summary_schedule(first_day)
        self.connection.commit()

        self.assertIsNotNone(first_claim)
        self.assertIsNone(_claim_topic_summary_schedule(first_day))
        self.connection.commit()
        self.assertIsNone(_claim_topic_summary_schedule(datetime(2026, 1, 6, 9, 0)))
        self.connection.commit()
        third_claim = _claim_topic_summary_schedule(datetime(2026, 1, 7, 9, 0))
        self.connection.commit()
        self.assertIsNotNone(third_claim)
        self.assertEqual(
            topic_summary_schedule()["schedule"],
            "Every other day at 09:00 local server time",
        )

    def test_topic_summary_prompt_uses_configured_class_and_pi_smpc_assessment(self):
        prompt = build_summary_prompt(
            product_name="Alprazolam",
            active_substance="Alprazolam",
            trigger_name="Safety signal",
            safety_topic="Adverse reaction",
            relevant_terms="",
            evidence="Exact source evidence",
            product_class_context="Alprazolam → Benzodiazepines",
        )

        self.assertIn("Alprazolam → Benzodiazepines", prompt)
        self.assertIn("PI/SmPC update", prompt)
        self.assertIn("proposed outcome for the human reviewer", prompt)

    def test_new_information_documents_are_filtered_and_paginated(self):
        self._add_new_information_document(1, "Alprazolam safety update", "Alprazolam", "2026-09-25")
        self._add_new_information_document(2, "Diazepam safety update", "Diazepam", "2026-09-26")
        self.connection.execute(
            "INSERT INTO product_classes (id, class_name, active) VALUES (1, 'Benzodiazepines', 1)"
        )
        self.connection.execute(
            "INSERT INTO products (id, product_class_id, product_name, active) VALUES (1, 1, 'Alprazolam', 1)"
        )
        self.connection.commit()

        response = list_new_information_documents(
            page=1,
            page_size=1,
            date_range="all",
            date_from=None,
            date_to=None,
            source_id=None,
            product_class_id=1,
            content_type="Meeting Minutes",
            trigger_name="Safety signal",
            safety_trigger=None,
            search="Alprazolam",
        )

        self.assertEqual(response["total"], 1)
        self.assertEqual(len(response["records"]), 1)
        self.assertEqual(response["records"][0]["title"], "Alprazolam safety update")
        self.assertEqual(response["records"][0]["products_found"], ["Alprazolam"])
        self.assertEqual(response["records"][0]["safety_trigger_count"], 1)
        self.assertEqual(response["stats"]["new_documents"], 1)
        self.assertEqual(response["stats"]["potential_safety_triggers"], 1)
        self.assertEqual(response["filters"]["sources"][0]["id"], 1)

        second_page = list_new_information_documents(
            page=2,
            page_size=1,
            date_range="all",
            date_from=None,
            date_to=None,
            source_id=None,
            content_type=None,
            product_class_id=None,
            trigger_name=None,
            safety_trigger=None,
            search=None,
        )
        self.assertEqual(second_page["total"], 2)
        self.assertEqual(second_page["records"][0]["title"], "Alprazolam safety update")

        custom_date = list_new_information_documents(
            page=1,
            page_size=5,
            date_range="custom",
            date_from="2026-09-26",
            date_to="2026-09-26",
            source_id=None,
            content_type=None,
            product_class_id=None,
            trigger_name=None,
            safety_trigger=None,
            search=None,
        )
        self.assertEqual(custom_date["total"], 1)
        self.assertEqual(custom_date["records"][0]["title"], "Diazepam safety update")

    def test_new_information_details_return_real_summary_products_and_evidence(self):
        self._add_new_information_document(3, "PRAC safety update", "Alprazolam", "2026-09-25")
        self.connection.execute(
            "INSERT INTO summaries (finding_id, summary, created_at) VALUES (1, 'Stored evidence summary', '2026-09-25')"
        )
        self.connection.commit()

        response = get_new_information_document(3)

        self.assertEqual(response["reference_id"], "EMA_PRAC_001_3")
        self.assertEqual(response["safety_information_summary"], "Stored evidence summary")
        self.assertEqual(response["products_found"], ["Alprazolam"])
        self.assertEqual(response["safety_trigger_count"], 1)
        self.assertEqual(response["safety_triggers"][0]["evidence"], "Exact quote")
        self.assertEqual(response["safety_triggers"][0]["page_number"], 2)

    def test_new_information_rejects_invalid_custom_date_range(self):
        with self.assertRaises(HTTPException) as context:
            list_new_information_documents(
                date_range="custom",
                date_from="2026-09-28",
                date_to="2026-09-01",
            )

        self.assertEqual(context.exception.status_code, 422)

    def test_new_information_api_defaults_to_all_time(self):
        parameters = app.openapi()["paths"]["/api/new-information"]["get"]["parameters"]
        date_range = next(parameter for parameter in parameters if parameter["name"] == "date_range")

        self.assertEqual(date_range["schema"]["default"], "all")

    def test_repeated_crawl_updates_existing_document_without_duplicate(self):
        source = {
            "id": 1,
            "reference_no": "EMA_PRAC_001",
        }
        first_document = {
            "title": "PRAC Meeting Minutes",
            "url": "https://example.test/minutes.pdf",
            "document_key": "https://example.test/minutes.pdf",
            "publication_date": "25-Sep-2026",
            "content_type": "Meeting Minutes",
        }

        new_count, _ = _record_discovered_documents(
            source,
            [first_document],
        )
        first_reference_id = first_document["reference_id"]
        self.assertEqual(new_count, 1)
        self.assertTrue(first_document["new_document"])
        self.assertEqual(
            first_reference_id,
            "EMA_PRAC_001_20260925_001",
        )

        repeated_document = {
            **first_document,
            "title": "PRAC Meeting Minutes - Updated link text",
            "publication_date": "04 September 2026",
            "status": "NEW",
            "new_document": True,
        }
        new_count, _ = _record_discovered_documents(
            source,
            [repeated_document],
        )

        row = self.connection.execute(
            "SELECT * FROM documents"
        ).fetchone()
        self.assertEqual(new_count, 0)
        self.assertFalse(repeated_document["new_document"])
        self.assertEqual(repeated_document["status"], "ALREADY_CRAWLED")
        self.assertEqual(self.connection.execute(
            "SELECT COUNT(*) FROM documents"
        ).fetchone()[0], 1)
        self.assertNotEqual(row["reference_id"], first_reference_id)
        self.assertEqual(
            row["reference_id"],
            "EMA_PRAC_001_20260904_001",
        )
        self.assertEqual(row["publication_date"], "04 September 2026")
        self.assertEqual(row["title"], "PRAC Meeting Minutes - Updated link text")

    def test_crawl_discovery_endpoint_does_not_persist_or_record_crawls(self):
        discovered = {
            "title": "PRAC Meeting Minutes",
            "url": "https://example.test/prac/minutes.pdf",
            "document_key": "https://example.test/prac/minutes.pdf",
            "filename": "minutes.pdf",
            "meeting_date": "8-11 June 2026",
            "publication_date": "25 September 2026",
            "content_type": "Meeting Minutes",
            "status": "NEW",
        }
        with patch(
            "backend.main.crawl_source",
            return_value={
                "success": True,
                "documents": [discovered.copy()],
                "documents_found": 1,
                "new_documents": 1,
            },
        ) as crawl:
            first_result = discover_crawl_documents(
                CrawlDiscoverRequest(source_id=1),
            )
            self.assertEqual(first_result["new_documents"], 1)
            self.assertEqual(
                first_result["documents"][0]["status"],
                "NEW",
            )
            second_result = discover_crawl_documents(
                CrawlDiscoverRequest(source_id=1),
            )

        self.assertEqual(second_result["new_documents"], 1)
        self.assertEqual(
            second_result["documents"][0]["status"],
            "NEW",
        )
        self.assertEqual(
            crawl.call_args.kwargs["content_types"],
            [
                "Meeting Minutes",
                "Agendas",
                "Highlights",
                "Other published documents",
            ],
        )
        self.assertEqual(
            self.connection.execute(
                "SELECT COUNT(*) FROM documents"
            ).fetchone()[0],
            0,
        )
        self.assertEqual(
            self.connection.execute(
                "SELECT COUNT(*) FROM crawl_history"
            ).fetchone()[0],
            0,
        )

    def test_crawl_discovery_returns_actual_groups_and_all_documents(self):
        discovered = [
            {
                "title": "July minutes",
                "url": "https://example.test/july.pdf",
                "filename": "july.pdf",
                "publication_date": "10 July 2026",
                "content_type": "Meeting Minutes",
            },
            {
                "title": "September minutes",
                "url": "https://example.test/september.pdf",
                "filename": "september.pdf",
                "publication_date": "25 September 2026",
                "content_type": "Meeting Minutes",
            },
            {
                "title": "Latest agenda",
                "url": "https://example.test/agenda.pdf",
                "filename": "agenda.pdf",
                "publication_date": "30 September 2026",
                "content_type": "Agendas",
            },
            {
                "title": "August minutes",
                "url": "https://example.test/august.pdf",
                "filename": "august.pdf",
                "publication_date": "18 August 2026",
                "content_type": "Meeting Minutes",
            },
            {
                "title": "Procedure rules",
                "url": "https://example.test/rules.pdf",
                "filename": "rules.pdf",
                "publication_date": "20 July 2026",
                "content_type": "Rules of Procedure",
            },
        ]
        with patch(
            "backend.main.crawl_source",
            return_value={
                "success": True,
                "documents": discovered,
                "documents_found": len(discovered),
                "new_documents": len(discovered),
            },
        ) as crawl:
            result = discover_crawl_documents(CrawlDiscoverRequest(source_id=1))

        crawl.assert_called_once()
        self.assertEqual(crawl.call_args.kwargs["source_id"], None)
        self.assertEqual(
            crawl.call_args.kwargs["content_types"],
            [
                "Meeting Minutes",
                "Agendas",
                "Highlights",
                "Other published documents",
            ],
        )
        self.assertEqual(result["documents_found"], 5)
        self.assertEqual(
            [item["title"] for item in result["documents"]],
            [
                "Latest agenda",
                "September minutes",
                "August minutes",
                "Procedure rules",
                "July minutes",
            ],
        )
        self.assertEqual(
            {group["content_type"]: group["count"] for group in result["groups"]},
            {
                "Agendas": 1,
                "Meeting Minutes": 3,
                "Rules of Procedure": 1,
            },
        )
        self.assertEqual(
            [item["reference_id"] for item in result["documents"]],
            [
                "EMA_PRAC_001_20260930_001",
                "EMA_PRAC_001_20260925_002",
                "EMA_PRAC_001_20260818_003",
                "EMA_PRAC_001_20260720_004",
                "EMA_PRAC_001_20260710_005",
            ],
        )
        self.assertEqual(
            self.connection.execute("SELECT COUNT(*) FROM documents").fetchone()[0],
            0,
        )

    def test_document_summary_generation_uses_extraction_without_ollama(self):
        self.connection.execute(
            """
            INSERT INTO documents
                (source_id, title, url, publication_date, content_type,
                 reference_id, status, created_at)
            VALUES (1, 'PRAC Meeting Minutes',
                    'https://example.test/minutes.pdf', '25 September 2026',
                    'Meeting Minutes', 'EMA_PRAC_001_20260925_001',
                    'PROCESSED', '2026-09-25T00:00:00')
            """
        )
        document_id = self.connection.execute(
            "SELECT id FROM documents WHERE url = ?",
            ("https://example.test/minutes.pdf",),
        ).fetchone()["id"]
        self.connection.execute(
            "INSERT INTO document_pages (document_id, page_number, page_text) VALUES (?, 1, ?)",
            (
                document_id,
                "The committee reviewed important safety concerns for the medicinal product. "
                "The committee recommended continued monitoring after the assessment.",
            ),
        )
        self.connection.execute(
            """
            INSERT INTO findings
                (document_id, product_name, safety_topic, trigger_name, evidence)
            VALUES (?, 'Diazepam', 'Adverse reactions', 'Safety signal', 'Monitor risk')
            """,
            (document_id,),
        )

        with patch("backend.main.test_ollama_connection") as ollama:
            result = generate_poc_document_summary(document_id)
            summary = get_document_summary(document_id)

        ollama.assert_not_called()
        self.assertEqual(result["status"], "COMPLETED")
        self.assertEqual(
            result["summary_method"],
            "Document Extraction / POC Summary",
        )
        self.assertIn("safety", " ".join(result["summary"]["key_points"]).lower())
        self.assertEqual(result["summary"]["products"], ["Diazepam"])
        self.assertEqual(summary["summary_status"], "AVAILABLE")
        self.assertEqual(summary["document"]["title"], "PRAC Meeting Minutes")
        self.assertIn("Key safety and regulatory points include", summary["executive_summary"])
        stored = self.connection.execute(
            """
            SELECT provider, status, summary
            FROM document_ai_analyses
            WHERE document_id = ?
            """,
            (document_id,),
        ).fetchone()
        self.assertEqual(stored["provider"], "document-extraction")
        self.assertEqual(stored["status"], "COMPLETED")
        self.assertIn('"method": "Document Extraction / POC Summary"', stored["summary"])

    def test_crawl_discovery_marks_previously_processed_document(self):
        self.connection.execute(
            """
            INSERT INTO documents
                (source_id, title, url, created_at, status)
            VALUES (1, 'PRAC Meeting Minutes',
                    'https://example.test/prac/minutes.pdf',
                    '2026-09-25T00:00:00', 'PROCESSED')
            """
        )
        discovered = {
            "title": "PRAC Meeting Minutes",
            "url": "https://example.test/prac/minutes.pdf",
            "document_key": "https://example.test/prac/minutes.pdf",
            "filename": "minutes.pdf",
            "publication_date": "25 September 2026",
            "content_type": "Meeting Minutes",
            "status": "NEW",
        }
        with patch(
            "backend.main.crawl_source",
            return_value={
                "success": True,
                "documents": [discovered],
                "documents_found": 1,
                "new_documents": 1,
            },
        ):
            result = discover_crawl_documents(
                CrawlDiscoverRequest(source_id=1),
            )

        self.assertEqual(result["new_documents"], 0)
        self.assertEqual(result["already_crawled"], 1)
        self.assertEqual(result["documents"][0]["status"], "ALREADY_CRAWLED")
        self.assertIsNotNone(result["documents"][0]["reference_id"])
        self.assertEqual(
            self.connection.execute(
                "SELECT COUNT(*) FROM documents"
            ).fetchone()[0],
            1,
        )

    def test_failed_crawl_discovery_does_not_record_a_crawl(self):
        with patch(
            "backend.main.crawl_source",
            side_effect=RuntimeError("source unavailable"),
        ):
            with self.assertRaises(HTTPException):
                discover_crawl_documents(
                    CrawlDiscoverRequest(source_id=1),
                )

        self.assertEqual(
            self.connection.execute(
                "SELECT COUNT(*) FROM documents"
            ).fetchone()[0],
            0,
        )
        self.assertEqual(
            self.connection.execute(
                "SELECT COUNT(*) FROM crawl_history"
            ).fetchone()[0],
            0,
        )

    def test_selected_processing_saves_content_type_reference_and_history(self):
        job_id = "selected-document-test"
        selected_processing_jobs[job_id] = {}
        try:
            with (
                patch("backend.main.download_pdf", return_value="test.pdf"),
                patch(
                    "backend.main.process_pdf_safe",
                    return_value={
                        "success": True,
                        "pages": [{"page_number": 1, "text": "Test page"}],
                        "sha256": "test-hash",
                    },
                ),
                patch("backend.main.detect_document", return_value=[]),
                patch("backend.main.extract_meeting_topics", return_value=[]),
            ):
                process_selected_document(
                    job_id,
                    1,
                    "https://example.test/prac/minutes.pdf",
                    "PRAC Meeting Minutes",
                    "minutes.pdf",
                    "8-11 June 2026",
                    "25 September 2026",
                    "Meeting Minutes",
                    None,
                )

            document = self.connection.execute(
                "SELECT * FROM documents WHERE id = 1"
            ).fetchone()
            history = self.connection.execute(
                "SELECT * FROM crawl_history WHERE source_id = 1"
            ).fetchone()
            self.assertEqual(selected_processing_jobs[job_id]["status"], "completed")
            self.assertEqual(document["status"], "PROCESSED")
            self.assertEqual(document["content_type"], "Meeting Minutes")
            self.assertEqual(
                document["reference_id"],
                "EMA_PRAC_001_20260925_001",
            )
            self.assertIsNotNone(document["last_crawled_at"])
            self.assertEqual(history["status"], "success")
            self.assertEqual(history["new_documents"], 1)
        finally:
            selected_processing_jobs.pop(job_id, None)

    def test_already_processed_selection_skips_download_extraction_and_report(self):
        self.connection.execute(
            """
            INSERT INTO documents
                (source_id, title, url, created_at, status)
            VALUES (1, 'Previously processed',
                    'https://example.test/already.pdf',
                    '2026-09-25T00:00:00', 'PROCESSED')
            """
        )
        self.connection.execute(
            "INSERT INTO document_pages (document_id, page_number, page_text) VALUES (1, 1, 'Existing text')"
        )
        job_id = "already-processed-test"
        selected_processing_jobs[job_id] = {}
        try:
            with (
                patch("backend.main.download_pdf") as download,
                patch("backend.main.process_pdf_safe") as extract,
                patch("backend.main.create_document_summary_pdf") as create_report,
            ):
                process_selected_document(
                    job_id,
                    1,
                    "https://example.test/already.pdf",
                    "Previously processed",
                    "already.pdf",
                    None,
                    None,
                    "Meeting Minutes",
                    None,
                )
            self.assertEqual(selected_processing_jobs[job_id]["status"], "completed")
            self.assertTrue(selected_processing_jobs[job_id]["already_processed"])
            download.assert_not_called()
            extract.assert_not_called()
            create_report.assert_not_called()
        finally:
            selected_processing_jobs.pop(job_id, None)

    def test_source_listing_includes_persisted_crawl_counts(self):
        self.connection.execute(
            """
            INSERT INTO documents
                (source_id, title, url, created_at, status)
            VALUES (1, 'Stored minutes', 'https://example.test/stored.pdf',
                    '2026-10-01T00:00:00', 'NEW')
            """
        )
        self.connection.execute(
            """
            INSERT INTO crawl_history
                (source_id, started_at, completed_at, status,
                 documents_found, new_documents, total_documents)
            VALUES (1, '2026-10-01T00:00:00', '2026-10-01T00:01:00',
                    'success', 1, 1, 1)
            """
        )

        source = list_sources()[0]

        self.assertEqual(source["crawl_count"], 1)
        self.assertEqual(source["crawl_status"], "success")
        self.assertEqual(source["latest_new_documents"], 1)
        self.assertEqual(source["latest_total_documents"], 1)
        self.assertEqual(source["total_documents"], 1)

    def test_source_update_creates_configuration_audit_event(self):
        update_source(1, SourceUpdate(monitoring_frequency="Weekly"))

        event = self.connection.execute(
            "SELECT event_type, details FROM audit_events"
        ).fetchone()

        self.assertEqual(event["event_type"], "Source Updated")
        self.assertIn("Monitoring Frequency", event["details"])
        self.connection.commit()
        self.assertEqual(_claim_scheduled_crawls(datetime.utcnow()), [])
        self.connection.commit()

    def test_automatic_monitoring_configuration_is_persisted_and_manual_disables_it(self):
        original = self.connection.execute(
            "SELECT url FROM sources WHERE id = 1"
        ).fetchone()["url"]
        update_source(
            1,
            SourceUpdate(
                monitoring_frequency="Weekly",
                instruction="Keep this source-specific instruction.",
                automatic_monitoring=True,
            ),
        )

        configured = crawl_schedule_status(1)
        source = self.connection.execute(
            "SELECT url, monitoring_frequency, instruction FROM sources WHERE id = 1"
        ).fetchone()
        self.assertTrue(configured["configured_enabled"])
        self.assertEqual(configured["frequency"], "Weekly")
        self.assertEqual(source["url"], original)
        self.assertEqual(source["monitoring_frequency"], "Weekly")
        self.assertEqual(source["instruction"], "Keep this source-specific instruction.")

        update_source(
            1,
            SourceUpdate(
                monitoring_frequency="Manual",
                automatic_monitoring=False,
            ),
        )
        manual = crawl_schedule_status(1)
        self.assertFalse(manual["configured_enabled"])
        self.assertIn("Manual mode", manual["message"])

    def test_source_frequency_updates_keep_an_enabled_schedule_in_sync(self):
        update_source(
            1,
            SourceUpdate(monitoring_frequency="Weekly", automatic_monitoring=True),
        )
        update_source(1, SourceUpdate(monitoring_frequency="Monthly"))

        schedule = crawl_schedule_status(1)
        self.assertTrue(schedule["configured_enabled"])
        self.assertEqual(schedule["frequency"], "Monthly")

    def test_unsupported_source_frequency_cannot_silently_break_an_enabled_schedule(self):
        update_source(
            1,
            SourceUpdate(monitoring_frequency="Weekly", automatic_monitoring=True),
        )
        with self.assertRaises(HTTPException) as error:
            update_source(1, SourceUpdate(monitoring_frequency="Daily"))

        source = self.connection.execute(
            "SELECT monitoring_frequency FROM sources WHERE id = 1"
        ).fetchone()
        self.assertEqual(error.exception.status_code, 422)
        self.assertEqual(source["monitoring_frequency"], "Weekly")

    def test_scheduled_crawl_claim_is_due_once_and_does_not_start_a_crawl(self):
        now = datetime.utcnow()
        self.connection.execute(
            """
            INSERT INTO audit_events
                (created_at, event_type, details, source_id, status)
            VALUES (?, 'CRAWL_SCHEDULE_CONFIGURED',
                    '{"enabled": true, "frequency": "Weekly"}', 1, 'COMPLETED')
            """,
            ((now - timedelta(days=8)).isoformat(),),
        )
        self.connection.commit()

        claims = _claim_scheduled_crawls(now)
        self.connection.commit()
        self.assertEqual(len(claims), 1)
        self.assertEqual(claims[0]["source_id"], 1)
        self.assertEqual(_claim_scheduled_crawls(now), [])
        self.connection.commit()

    def test_monthly_schedule_uses_calendar_month_boundaries(self):
        january_end = datetime(2026, 1, 31, 9, 0)
        february_run = _next_monitoring_run(january_end, "Monthly")
        march_run = _next_monitoring_run(february_run, "Monthly")

        self.assertEqual(february_run, datetime(2026, 2, 28, 9, 0))
        self.assertEqual(march_run, datetime(2026, 3, 31, 9, 0))

    def test_schedule_reconfiguration_waits_for_the_new_frequency_interval(self):
        now = datetime.utcnow()
        self.connection.execute(
            """
            INSERT INTO audit_events (created_at, event_type, details, source_id, status)
            VALUES (?, 'CRAWL_SCHEDULE_CONFIGURED',
                    '{"enabled": true, "frequency": "Weekly"}', 1, 'COMPLETED')
            """,
            ((now - timedelta(hours=1)).isoformat(),),
        )
        self.connection.execute(
            """
            INSERT INTO audit_events (created_at, event_type, details, source_id, status)
            VALUES (?, 'SCHEDULED_SOURCE_CRAWL', '{}', 1, 'COMPLETED')
            """,
            ((now - timedelta(days=8)).isoformat(),),
        )
        self.connection.commit()

        self.assertEqual(_claim_scheduled_crawls(now), [])
        self.connection.commit()

    def test_document_discovery_groups_actual_types_and_removes_duplicate_urls(self):
        discovered_documents = [
            {
                "title": "PRAC Agenda",
                "url": "https://example.test/agenda.pdf",
                "document_key": "https://example.test/agenda.pdf",
                "content_type": "Agendas",
                "publication_date": "2026-10-01",
            },
            {
                "title": "PRAC Agenda duplicate",
                "url": "https://example.test/agenda.pdf",
                "document_key": "https://example.test/agenda.pdf",
                "content_type": "Agendas",
                "publication_date": "2026-10-01",
            },
            {
                "title": "Meeting highlights",
                "url": "https://example.test/highlights",
                "document_key": "https://example.test/highlights",
                "content_type": "Highlights",
                "publication_date": "2026-10-02",
            },
        ]
        with patch(
            "backend.main.crawl_source",
            return_value={"success": True, "documents": discovered_documents},
        ):
            result = discover_crawl_documents(CrawlDiscoverRequest(source_id=1))

        self.assertEqual(result["documents_found"], 2)
        self.assertEqual(
            {group["content_type"]: group["count"] for group in result["groups"]},
            {"Agendas": 1, "Highlights": 1},
        )
        self.assertEqual(
            {document["url"] for document in result["documents"]},
            {"https://example.test/agenda.pdf", "https://example.test/highlights"},
        )

    def test_document_processing_rejects_html_even_if_filename_claims_pdf(self):
        with self.assertRaises(HTTPException) as error:
            start_selected_document_processing(
                CrawlProcessRequest(
                    source_id=1,
                    document_url="https://example.test/article",
                    filename="article.pdf",
                ),
                BackgroundTasks(),
            )

        self.assertEqual(error.exception.status_code, 400)
        self.assertIn("PDF", error.exception.detail)

    def test_source_update_persists_editor_fields_without_changing_other_settings(self):
        update_source(
            1,
            SourceUpdate(
                name="Updated PRAC source",
                source_type="PRAC Signals",
                authority="Updated authority",
                url="https://example.test/updated-prac",
                content_types=["Agendas", "Highlights"],
                monitoring_frequency="Weekly",
                instruction="Updated instructions",
                active=False,
            ),
        )

        source = self.connection.execute(
            "SELECT * FROM sources WHERE id = 1"
        ).fetchone()
        self.assertEqual(source["name"], "Updated PRAC source")
        self.assertEqual(source["source_type"], "PRAC Signals")
        self.assertEqual(source["authority"], "Updated authority")
        self.assertEqual(source["url"], "https://example.test/updated-prac")
        self.assertEqual(
            json.loads(source["content_types"]),
            ["Agendas", "Highlights"],
        )
        self.assertEqual(source["monitoring_frequency"], "Weekly")
        self.assertEqual(source["instruction"], "Updated instructions")
        self.assertEqual(source["active"], 0)

    def test_source_creation_persists_all_supported_settings(self):
        result = create_source(
            SourceCreate(
                reference_no="EMA_SIGNAL_002",
                authority="European Medicines Agency",
                name="EMA Safety Signals",
                source_type="PRAC Signals",
                region="EEA",
                monitoring_frequency="Weekly",
                website="EMA",
                safety_alert_type="PRAC recommendation on signals",
                url="https://example.test/signals",
                content_types=["PRAC recommendation on signals"],
                instruction="Monitor safety signal recommendations.",
                active=False,
            )
        )

        source = self.connection.execute(
            "SELECT * FROM sources WHERE id = ?",
            (result["id"],),
        ).fetchone()
        self.assertEqual(source["reference_no"], "EMA_SIGNAL_002")
        self.assertEqual(source["name"], "EMA Safety Signals")
        self.assertEqual(source["source_type"], "PRAC Signals")
        self.assertEqual(source["monitoring_frequency"], "Weekly")
        self.assertEqual(source["url"], "https://example.test/signals")
        self.assertEqual(
            json.loads(source["content_types"]),
            ["PRAC recommendation on signals"],
        )
        self.assertEqual(source["instruction"], "Monitor safety signal recommendations.")
        self.assertEqual(source["active"], 0)

    def test_source_test_uses_configured_content_types(self):
        with patch(
            "backend.main.test_source",
            return_value={
                "url": "https://example.test/signals",
                "status": "SUCCESS",
                "pdf_links_found": 1,
            },
        ) as source_test:
            result = test_source_endpoint(
                SourceTestRequest(
                    url="https://example.test/signals",
                    content_types=["PRAC recommendation on signals"],
                )
            )

        source_test.assert_called_once_with(
            "https://example.test/signals",
            content_types=["PRAC recommendation on signals"],
        )
        self.assertTrue(result["success"])

    def test_startup_preserves_existing_builtin_source_settings(self):
        self.connection.execute(
            """
            UPDATE sources
            SET authority = 'User-edited authority',
                name = 'User-edited source name',
                source_type = 'User-edited type',
                region = 'User region',
                monitoring_frequency = 'Daily',
                website = 'User website',
                safety_alert_type = 'User alert',
                url = 'https://user.example.test/prac',
                content_types = '["User content"]',
                instruction = 'User instructions',
                active = 0,
                created_at = 'original-created',
                updated_at = 'original-updated'
            WHERE id = 1
            """
        )
        default_sources = [
            {
                "reference_no": "EMA_PRAC_001",
                "name": "Default PRAC",
                "source_type": "PRAC",
                "region": "EEA",
                "monitoring_frequency": "Monthly",
                "website": "EMA",
                "safety_alert_type": "PRAC",
                "url": "https://default.example.test/prac",
                "content_types": ["Meeting Minutes"],
                "instruction": "Default instruction",
                "active": True,
            },
            {
                "reference_no": "EMA_SIGNAL_001",
                "name": "Default signals",
                "source_type": "PRAC Signals",
                "region": "EEA",
                "monitoring_frequency": "Monthly",
                "website": "EMA",
                "safety_alert_type": "Signals",
                "url": "https://default.example.test/signals",
                "content_types": ["PRAC recommendation on signals"],
                "instruction": "Default signal instruction",
                "active": True,
            }
        ]

        with (
            patch("backend.main.init_db"),
            patch("backend.main.DEFAULT_SOURCE_CONFIGS", default_sources),
            patch("backend.main.DEFAULT_PRODUCT_CLASSES", []),
            patch("backend.main._topic_summary_scheduler_thread", None),
            patch("backend.main.threading.Thread"),
        ):
            startup()

        source = self.connection.execute(
            "SELECT * FROM sources WHERE id = 1"
        ).fetchone()
        self.assertEqual(source["authority"], "User-edited authority")
        self.assertEqual(source["name"], "User-edited source name")
        self.assertEqual(source["source_type"], "User-edited type")
        self.assertEqual(source["region"], "User region")
        self.assertEqual(source["monitoring_frequency"], "Daily")
        self.assertEqual(source["website"], "User website")
        self.assertEqual(source["safety_alert_type"], "User alert")
        self.assertEqual(source["url"], "https://user.example.test/prac")
        self.assertEqual(json.loads(source["content_types"]), ["User content"])
        self.assertEqual(source["instruction"], "User instructions")
        self.assertEqual(source["active"], 0)
        self.assertEqual(source["created_at"], "original-created")
        self.assertEqual(source["updated_at"], "original-updated")
        self.assertEqual(
            self.connection.execute("SELECT COUNT(*) FROM sources").fetchone()[0],
            2,
        )
        inserted = self.connection.execute(
            "SELECT * FROM sources WHERE reference_no = 'EMA_SIGNAL_001'"
        ).fetchone()
        self.assertEqual(inserted["name"], "Default signals")
        self.assertEqual(inserted["url"], "https://default.example.test/signals")
        self.assertEqual(inserted["monitoring_frequency"], "Monthly")

    def test_crawl_content_selection_must_be_configured(self):
        source = {
            "content_types": ["Meeting Minutes", "Agendas"],
            "safety_alert_type": "PRAC",
        }

        self.assertEqual(
            _crawl_content_types(source, ["Agendas"]),
            ["Agendas"],
        )
        with self.assertRaises(HTTPException):
            _crawl_content_types(source, ["Highlights"])
        with self.assertRaises(HTTPException):
            _crawl_content_types(source, [])

    def test_source_documents_are_source_scoped_with_publication_metadata(self):
        self.connection.execute(
            """
            INSERT INTO sources (id, reference_no, authority, name, url, active)
            VALUES (2, 'EMA_SIGNAL_001', 'EMA', 'EMA Signals',
                    'https://example.test/signals', 1)
            """
        )
        self.connection.executemany(
            """
            INSERT INTO documents
                (source_id, title, url, publication_date, content_type,
                 reference_id, discovered_date, status, created_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            [
                (
                    1,
                    "PRAC Meeting Minutes",
                    "https://example.test/prac/documents/minutes.pdf",
                    "25 September 2026",
                    "Meeting Minutes",
                    "EMA_PRAC_001_20260925_001",
                    "2026-09-30T00:00:00",
                    "PROCESSED",
                    "2026-09-30T00:00:00",
                ),
                (
                    2,
                    "Signal recommendation",
                    "https://example.test/signals/one.pdf",
                    "26 September 2026",
                    "PRAC recommendation on signals",
                    None,
                    "2026-09-30T00:00:00",
                    "NEW",
                    "2026-09-30T00:00:00",
                ),
                (
                    2,
                    "Signal management",
                    "https://example.test/signals/navigation",
                    "Signal management",
                    "PRAC recommendation on signals",
                    None,
                    "2026-09-30T00:00:00",
                    "NEW",
                    "2026-09-30T00:00:00",
                ),
            ],
        )

        prac_documents = list_source_documents(source_id=1, limit=1000)
        signal_documents = list_source_documents(source_id=2, limit=1000)

        self.assertEqual(len(prac_documents), 1)
        self.assertEqual(prac_documents[0]["publication_date"], "25 September 2026")
        self.assertEqual(prac_documents[0]["reference_id"], "EMA_PRAC_001_20260925_001")
        self.assertEqual(prac_documents[0]["working_status"], "Yes")
        self.assertEqual(len(signal_documents), 1)
        self.assertEqual(signal_documents[0]["source_id"], 2)
        self.assertEqual(signal_documents[0]["reference_id"], "EMA_SIGNAL_001_20260926_002")
        self.assertEqual(signal_documents[0]["working_status"], "Pending")
        source_totals = {source["id"]: source for source in list_sources()}
        self.assertEqual(source_totals[2]["total_documents"], 1)

    def test_source_documents_returns_404_for_unknown_source(self):
        with self.assertRaises(HTTPException) as error:
            list_source_documents(source_id=999, limit=1000)

        self.assertEqual(error.exception.status_code, 404)

    def test_results_include_real_extraction_and_ai_status_fields(self):
        self.connection.execute(
            """
            INSERT INTO documents
                (source_id, title, url, publication_date, content_type,
                 reference_id, discovered_date, status, created_at)
            VALUES (
                1, 'Published minutes',
                'https://example.test/prac/documents/minutes.pdf',
                '25 September 2026', 'Meeting Minutes',
                'EMA_PRAC_001_20260925_001', '2026-09-29T10:30:00',
                'PROCESSED', '2026-09-28T10:30:00'
            )
            """
        )
        document_id = self.connection.execute(
            "SELECT id FROM documents WHERE reference_id = ?",
            ("EMA_PRAC_001_20260925_001",),
        ).fetchone()["id"]
        self.connection.execute(
            "INSERT INTO findings (document_id, active_substance) VALUES (?, ?)",
            (document_id, "Diazepam, Lorazepam"),
        )
        self.connection.execute(
            """
            INSERT INTO document_pages (document_id, page_number, page_text)
            VALUES (?, 1, 'Stored rare-content-keyword document text')
            """,
            (document_id,),
        )
        self.connection.execute(
            """
            INSERT INTO extraction_history
                (document_id, started_at, completed_at, status, pages_extracted)
            VALUES (?, '2026-09-28T10:00:00', '2026-09-28T10:30:00', 'SUCCESS', 1)
            """,
            (document_id,),
        )

        results = list_results(
            source_id=1,
            content_type=None,
            search=None,
            limit=1000,
        )

        self.assertEqual(len(results), 1)
        self.assertEqual(results[0]["publication_date"], "25 September 2026")
        self.assertEqual(results[0]["reference_id"], "EMA_PRAC_001_20260925_001")
        self.assertEqual(
            results[0]["products_found"],
            ["Diazepam", "Lorazepam"],
        )
        self.assertEqual(results[0]["safety_trigger_status"], "Found")
        self.assertIsNone(results[0]["ai_analysis_status"])
        self.assertEqual(results[0]["crawled_status"], "Yes")
        self.assertEqual(results[0]["date_crawled"], "2026-09-28T10:30:00")
        keyword_results = list_results(
            source_id=1,
            content_type=None,
            search="rare-content-keyword",
            limit=1000,
        )
        self.assertEqual([item["id"] for item in keyword_results], [document_id])
        document_view = list_documents_view(
            page=1,
            page_size=10,
            source_id=1,
            content_type="Meeting Minutes",
            status="working",
            date_from=None,
            date_to=None,
            search=None,
        )
        self.assertEqual(
            document_view["documents"][0]["products_found"],
            ["Diazepam", "Lorazepam"],
        )
        self.assertEqual(
            document_view["documents"][0]["safety_trigger_status"],
            "Detected",
        )

    def test_results_workflow_returns_only_saved_analysis_review_and_audit_data(self):
        self.connection.execute(
            """
            INSERT INTO documents
                (source_id, title, url, publication_date, content_type,
                 reference_id, discovered_date, status, created_at)
            VALUES (
                1, 'Reviewed minutes',
                'https://example.test/prac/documents/reviewed-minutes.pdf',
                '25 September 2026', 'Meeting Minutes',
                'EMA_PRAC_001_20260925_002', '2026-09-28T10:30:00',
                'PROCESSED', '2026-09-28T10:30:00'
            )
            """
        )
        document_id = self.connection.execute(
            "SELECT id FROM documents WHERE reference_id = ?",
            ("EMA_PRAC_001_20260925_002",),
        ).fetchone()["id"]
        self.connection.execute(
            "UPDATE sources SET instruction = ? WHERE id = 1",
            ("Check for Meeting minutes; download the minutes and summarize the topics.",),
        )
        self.connection.execute(
            """
            INSERT INTO findings
                (document_id, product_name, safety_topic, evidence, page_number)
            VALUES (?, 'Example medicine', 'Signal monitoring', 'Stored evidence', 2)
            """,
            (document_id,),
        )
        finding_id = self.connection.execute(
            "SELECT id FROM findings WHERE document_id = ?",
            (document_id,),
        ).fetchone()["id"]
        self.connection.execute(
            """
            INSERT INTO document_ai_analyses
                (document_id, provider, analysis_type, summary, status,
                 created_at, updated_at)
            VALUES (?, 'ollama', 'summary', 'Saved analysis text', 'COMPLETED',
                    '2026-09-29T09:00:00', '2026-09-29T09:00:00')
            """,
            (document_id,),
        )
        self.connection.execute(
            """
            INSERT INTO reviews
                (finding_id, stage, reviewer, decision, comments, reviewed_at)
            VALUES (?, 'INITIAL', 'Recorded reviewer', 'RELEVANT',
                    'Existing review note', '2026-09-29T10:00:00')
            """,
            (finding_id,),
        )
        self.connection.execute(
            """
            INSERT INTO audit_events
                (created_at, event_type, details, entity_type, entity_id,
                 action, ai_model, result, status)
            VALUES ('2026-09-29T09:00:00', 'AI Analysis (Summary)',
                    'Saved analysis generated', 'document', ?,
                    'AI Analysis (Summary)', 'test-model', 'Completed', 'COMPLETED')
            """,
            (document_id,),
        )
        before = self.connection.total_changes

        workflow = get_results_workflow(document_id)

        self.assertEqual(workflow["document"]["source_instruction"], "Check for Meeting minutes; download the minutes and summarize the topics.")
        self.assertEqual(workflow["analysis"]["summary"], "Saved analysis text")
        self.assertEqual(workflow["analysis"]["model"], "test-model")
        self.assertEqual(workflow["findings"][0]["evidence"], "Stored evidence")
        self.assertEqual(workflow["reviews"][0]["reviewer"], "Recorded reviewer")
        self.assertEqual(workflow["audit_events"][0]["event_type"], "AI Analysis (Summary)")
        self.assertFalse(workflow["reviewer_directory"]["configured"])
        self.assertFalse(workflow["notifications_configured"])
        self.assertEqual(self.connection.total_changes, before)

    def test_product_class_analysis_filters_findings_by_class_and_substance(self):
        self.connection.executemany(
            """
            INSERT INTO product_classes (id, reference_no, class_name, active)
            VALUES (?, ?, ?, 1)
            """,
            [
                (1, "CLASS_A", "Class A"),
                (2, "CLASS_B", "Class B"),
            ],
        )
        self.connection.executemany(
            """
            INSERT INTO products (id, product_class_id, product_name, active)
            VALUES (?, ?, ?, 1)
            """,
            [(1, 1, "Substance A"), (2, 1, "Substance A2"), (3, 2, "Substance B")],
        )
        for doc_id, title, drug, evidence in (
            (1, "Class A document", "Substance A", "Evidence from class A"),
            (2, "Class B document", "Substance B", "Evidence from class B"),
        ):
            self.connection.execute(
                """
                INSERT INTO documents
                    (id, source_id, title, url, publication_date, content_type,
                     reference_id, status, created_at)
                VALUES (?, 1, ?, ?, '2026-10-01', 'Safety notice', ?, 'PROCESSED', '2026-10-02')
                """,
                (doc_id, title, f"https://example.test/{doc_id}.pdf", f"REF-{doc_id}"),
            )
            self.connection.execute(
                """
                INSERT INTO findings
                    (document_id, product_name, trigger_name, safety_topic,
                     evidence, page_number)
                VALUES (?, ?, 'Safety Signal', 'Signal monitoring', ?, 2)
                """,
                (doc_id, drug, evidence),
            )

        all_findings = get_product_class_analysis(1)
        substance_findings = get_product_class_analysis(1, substance="Substance A")
        other_class = get_product_class_analysis(2)

        self.assertEqual(all_findings["product_class"]["class_name"], "Class A")
        self.assertEqual(len(all_findings["findings"]), 1)
        self.assertEqual(all_findings["findings"][0]["document_title"], "Class A document")
        self.assertEqual(all_findings["findings"][0]["category"], "Safety Signals")
        self.assertEqual(len(substance_findings["findings"]), 1)
        self.assertEqual(substance_findings["findings"][0]["product_name"], "Substance A")
        self.assertEqual(other_class["findings"][0]["document_title"], "Class B document")
        self.assertNotEqual(
            all_findings["findings"][0]["document_id"],
            other_class["findings"][0]["document_id"],
        )

    def test_product_class_ai_analysis_persists_mocked_provider_result(self):
        self.connection.execute(
            """
            INSERT INTO product_classes (id, reference_no, class_name, active)
            VALUES (1, 'CLASS_A', 'Class A', 1)
            """
        )
        self.connection.execute(
            """
            INSERT INTO products (id, product_class_id, product_name, active)
            VALUES (1, 1, 'Substance A', 1)
            """
        )
        self.connection.execute(
            """
            INSERT INTO documents
                (id, source_id, title, url, publication_date, content_type,
                 reference_id, status, created_at)
            VALUES (1, 1, 'Class A document', 'https://example.test/1.pdf',
                    '2026-10-01', 'Safety notice', 'REF-1',
                    'PROCESSED', '2026-10-02')
            """
        )
        self.connection.execute(
            """
            INSERT INTO findings
                (document_id, product_name, safety_topic, evidence, page_number)
            VALUES (1, 'Substance A', 'Safety topic', 'Stored source evidence', 2)
            """
        )
        with patch(
            "backend.main.list_ai_providers",
            return_value={"providers": [
                {"id": "ollama", "available": False, "message": "Local provider unavailable."},
                {"id": "openai", "available": True, "message": "Configured."},
            ]},
        ), patch(
            "backend.main.generate_openai_document_summary",
            return_value={
                "status": "SUCCESS",
                "model": "test-model",
                "summary": "Mocked class summary based on stored evidence.",
            },
        ) as generator:
            result = run_product_class_analysis(1, provider="openai", substance=None)

        self.assertEqual(result["status"], "COMPLETED")
        self.assertEqual(result["summary"], "Mocked class summary based on stored evidence.")
        generator.assert_called_once()
        persisted = self.connection.execute(
            """
            SELECT provider, summary, status
            FROM product_class_ai_analyses
            WHERE product_class_id = 1
            """
        ).fetchone()
        self.assertEqual(tuple(persisted), (
            "openai",
            "Mocked class summary based on stored evidence.",
            "COMPLETED",
        ))
        event = self.connection.execute(
            "SELECT event_type, status FROM audit_events WHERE entity_type = 'product_class'"
        ).fetchone()
        self.assertEqual(tuple(event), ("Product Class AI Analysis", "COMPLETED"))

    def test_documents_view_supports_database_pagination_filters_and_stats(self):
        self.connection.executemany(
            """
            INSERT INTO documents
                (source_id, title, url, publication_date, content_type,
                 reference_id, discovered_date, status, created_at)
            VALUES (1, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            [
                (
                    "Meeting minutes",
                    "https://example.test/minutes.pdf",
                    "25 September 2026",
                    "Meeting Minutes",
                    "EMA_PRAC_001_20260925_001",
                    "2026-09-28T10:30:00",
                    "PROCESSED",
                    "2026-09-28T10:30:00",
                ),
                (
                    "Meeting agenda",
                    "https://example.test/agenda.pdf",
                    "18 September 2026",
                    "Agendas",
                    "EMA_PRAC_001_20260918_001",
                    "2026-09-28T11:30:00",
                    "NEW",
                    "2026-09-28T11:30:00",
                ),
            ],
        )
        minutes_id = self.connection.execute(
            "SELECT id FROM documents WHERE title = 'Meeting minutes'"
        ).fetchone()["id"]
        self.connection.execute(
            "INSERT INTO findings (document_id, product_name) VALUES (?, ?)",
            (minutes_id, "Diazepam"),
        )

        result = list_documents_view(
            page=1,
            page_size=1,
            source_id=None,
            content_type=None,
            status=None,
            date_from=None,
            date_to=None,
            search=None,
        )
        self.assertEqual(result["total"], 2)
        self.assertEqual(len(result["documents"]), 1)
        self.assertEqual(result["page_size"], 1)
        self.assertEqual(result["documents"][0]["title"], "Meeting minutes")
        self.assertEqual(result["stats"]["total_documents"], 2)
        self.assertEqual(result["stats"]["safety_triggers"], 1)
        self.assertEqual(result["stats"]["active_sources"], 1)

        filtered = list_documents_view(
            page=1,
            page_size=10,
            source_id=1,
            content_type="Meeting Minutes",
            status="working",
            date_from="2026-09-25",
            date_to="2026-09-25",
            search="Diazepam",
        )
        self.assertEqual(filtered["total"], 1)
        self.assertEqual(filtered["documents"][0]["title"], "Meeting minutes")
        self.assertEqual(filtered["documents"][0]["products_found"], ["Diazepam"])
        self.assertEqual(filtered["documents"][0]["safety_trigger_status"], "Detected")

    def test_documents_view_is_source_scoped_and_classifies_signal_documents_from_metadata(self):
        now = datetime.utcnow().isoformat()
        self.connection.execute(
            """
            INSERT INTO sources
                (id, reference_no, authority, name, url, content_types, active)
            VALUES (
                2, 'EMA_SIGNAL_001', 'EMA',
                'EMA PRAC Signals', 'https://example.test/signals',
                '["Safety Signals", "PRAC Recommendations", "Agendas"]', 1
            )
            """
        )
        self.connection.executemany(
            """
            INSERT INTO documents
                (source_id, title, url, filename, publication_date,
                 discovered_date, content_type, status, created_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            [
                (
                    1, "Stored agenda", "https://example.test/agenda.pdf",
                    None, "25 September 2026", now, "Agendas", "NEW", now,
                ),
                (
                    2, "Signal management", "https://example.test/signal-management",
                    None, "20 September 2026", now,
                    "PRAC recommendation on signals", "NEW", now,
                ),
                (
                    2, "List of safety signals discussed since September 2012",
                    "https://example.test/signals#list-of-safety-signals",
                    None, "18 September 2026", now,
                    "PRAC recommendation on signals", "NEW", now,
                ),
                (
                    2, "PRAC recommendations on safety signals",
                    "https://example.test/prac-recommendations",
                    None, "17 September 2026", now,
                    "PRAC recommendation on signals", "PROCESSED", now,
                ),
                (
                    2, "Other published documents - new-product-information-wording-extracts.pdf",
                    "https://example.test/new-product-information-wording-extracts.pdf",
                    "local.pdf", "16 September 2026", now,
                    "Other published documents", "NEW", now,
                ),
                (
                    2, "PRAC work plan 2026",
                    "https://example.test/work-plan.pdf", None,
                    "15 September 2026", now, None, "NEW", now,
                ),
            ],
        )
        self.connection.commit()

        monitoring = list_documents_view(
            page=1,
            page_size=10,
            source_id=1,
            content_type=None,
            status=None,
            date_from=None,
            date_to=None,
            search=None,
        )
        signals = list_documents_view(
            page=1,
            page_size=10,
            source_id=2,
            content_type=None,
            status=None,
            date_from=None,
            date_to=None,
            search=None,
        )

        self.assertEqual(monitoring["total"], 1)
        self.assertEqual(monitoring["stats"]["total_documents"], 1)
        self.assertEqual(monitoring["stats"]["pdf_documents"], 1)
        self.assertEqual(monitoring["stats"]["html_documents"], 0)
        self.assertEqual(monitoring["filters"]["content_types"], ["Agendas"])
        self.assertEqual(signals["total"], 5)
        self.assertEqual(signals["stats"]["total_documents"], 5)
        self.assertEqual(signals["stats"]["pdf_documents"], 2)
        self.assertEqual(signals["stats"]["html_documents"], 3)
        self.assertEqual(
            set(signals["filters"]["content_types"]),
            {
                "Other published documents",
                "PRAC Recommendations",
                "Safety Signals",
                "Signal Management",
                "Work Plan",
            },
        )
        by_title = {item["title"]: item for item in signals["documents"]}
        self.assertEqual(by_title["Signal management"]["content_type"], "Signal Management")
        self.assertEqual(
            by_title["List of safety signals discussed since September 2012"]["content_type"],
            "Safety Signals",
        )
        self.assertEqual(
            by_title["PRAC recommendations on safety signals"]["content_type"],
            "PRAC Recommendations",
        )
        self.assertEqual(by_title["PRAC work plan 2026"]["content_type"], "Work Plan")
        stored_type = self.connection.execute(
            "SELECT content_type FROM documents WHERE title = 'PRAC work plan 2026'"
        ).fetchone()["content_type"]
        self.assertIsNone(stored_type)

        signal_management = list_documents_view(
            page=1,
            page_size=10,
            source_id=2,
            content_type="Signal Management",
            status=None,
            date_from=None,
            date_to=None,
            search=None,
        )
        self.assertEqual(signal_management["total"], 1)

    def test_documents_view_distinguishes_downloaded_processed_and_analyzed(self):
        self.connection.execute(
            """
            INSERT INTO documents
                (source_id, title, url, filename, content_type,
                 discovered_date, status, created_at)
            VALUES (1, 'Stored PDF', 'https://example.test/stored.pdf',
                    'stored.pdf', 'Meeting Minutes', '2026-09-25',
                    'PROCESSED', '2026-09-25')
            """
        )
        document_id = self.connection.execute(
            "SELECT id FROM documents WHERE title = 'Stored PDF'"
        ).fetchone()["id"]
        self.connection.execute(
            """
            INSERT INTO document_ai_analyses
                (document_id, provider, analysis_type, status,
                 created_at, updated_at)
            VALUES (?, 'ollama', 'summary', 'COMPLETED',
                    '2026-09-25', '2026-09-25')
            """,
            (document_id,),
        )
        self.connection.commit()

        with tempfile.TemporaryDirectory() as directory:
            (Path(directory) / "stored.pdf").write_bytes(b"%PDF-1.4 test")
            with patch("backend.main.PDF_DIR", Path(directory)):
                result = list_documents_view(
                    page=1,
                    page_size=10,
                    source_id=1,
                    content_type=None,
                    status="downloaded",
                    date_from=None,
                    date_to=None,
                    search=None,
                )
                self.assertEqual(result["total"], 1)
                item = result["documents"][0]
                self.assertEqual(item["downloaded_status"], "Yes")
                self.assertTrue(item["has_local_pdf"])
                self.assertEqual(item["processing_status"], "Processed")
                self.assertEqual(item["analysis_status"], "Analyzed")
                self.assertEqual(item["file_format"], "PDF")

    def test_documents_view_rejects_invalid_date_ranges(self):
        with self.assertRaises(HTTPException) as error:
            list_documents_view(
                page=1,
                page_size=10,
                source_id=None,
                content_type=None,
                status=None,
                date_from="2026-10-02",
                date_to="2026-10-01",
                search=None,
            )
        self.assertEqual(error.exception.status_code, 422)

    def test_document_ai_analysis_is_persisted_per_document_and_provider(self):
        self.connection.executemany(
            """
            INSERT INTO documents
                (source_id, title, url, status, created_at)
            VALUES (1, ?, ?, 'PROCESSED', '2026-09-28T10:30:00')
            """,
            [
                ("First document", "https://example.test/first.pdf"),
                ("Second document", "https://example.test/second.pdf"),
            ],
        )
        first_id = self.connection.execute(
            "SELECT id FROM documents WHERE title = 'First document'"
        ).fetchone()["id"]
        second_id = self.connection.execute(
            "SELECT id FROM documents WHERE title = 'Second document'"
        ).fetchone()["id"]
        self.connection.execute(
            "INSERT INTO document_pages (document_id, page_number, page_text) VALUES (?, 1, ?)",
            (first_id, "Source text for first document."),
        )

        with patch(
            "backend.main.generate_document_summary",
            return_value={
                "status": "SUCCESS",
                "model": "test-model",
                "summary": "A persisted summary for the first document.",
            },
        ) as generator:
            response = generate_document_ai_analysis(
                first_id,
                DocumentAIAnalysisRequest(provider="ollama"),
            )

        self.assertEqual(generator.call_count, 1)
        self.assertIn("first document", generator.call_args.args[1])
        self.assertEqual(response["status"], "COMPLETED")
        self.assertEqual(
            response["summary"],
            "A persisted summary for the first document.",
        )
        self.assertEqual(
            get_document_ai_analysis(first_id, provider="ollama")["status"],
            "COMPLETED",
        )
        self.assertEqual(
            get_document_ai_analysis(second_id, provider="ollama")["status"],
            "NOT_GENERATED",
        )
        audit_event = self.connection.execute(
            "SELECT action, ai_model, status FROM audit_events WHERE entity_id = ?",
            (first_id,),
        ).fetchone()
        self.assertEqual(audit_event["action"], "AI Analysis (Summary)")
        self.assertEqual(audit_event["ai_model"], "test-model")
        self.assertEqual(audit_event["status"], "COMPLETED")

    def test_document_ai_provider_failure_is_not_reported_as_success(self):
        self.connection.execute(
            """
            INSERT INTO documents (source_id, title, url, status, created_at)
            VALUES (1, 'Unanalyzed document', 'https://example.test/one.pdf',
                    'PROCESSED', '2026-09-28T10:30:00')
            """
        )
        document_id = self.connection.execute(
            "SELECT id FROM documents WHERE title = 'Unanalyzed document'"
        ).fetchone()["id"]
        self.connection.execute(
            "INSERT INTO document_pages (document_id, page_number, page_text) VALUES (?, 1, ?)",
            (document_id, "Extracted source text."),
        )

        with patch(
            "backend.main.generate_document_summary",
            return_value={
                "status": "UNAVAILABLE",
                "error": "Configured Ollama model is unavailable.",
            },
        ):
            with self.assertRaises(HTTPException) as error:
                generate_document_ai_analysis(
                    document_id,
                    DocumentAIAnalysisRequest(provider="ollama"),
                )

        self.assertEqual(error.exception.status_code, 503)
        persisted = get_document_ai_analysis(document_id, provider="ollama")
        self.assertEqual(persisted["status"], "UNAVAILABLE")
        self.assertIsNone(persisted["summary"])

    def test_openai_document_summary_is_persisted_with_model_provenance(self):
        self.connection.execute(
            """
            INSERT INTO documents (source_id, title, url, status, created_at)
            VALUES (1, 'OpenAI document', 'https://example.test/openai.pdf',
                    'PROCESSED', '2026-09-28T10:30:00')
            """
        )
        document_id = self.connection.execute(
            "SELECT id FROM documents WHERE title = 'OpenAI document'"
        ).fetchone()["id"]
        self.connection.execute(
            "INSERT INTO document_pages (document_id, page_number, page_text) VALUES (?, 1, ?)",
            (document_id, "Source text sent only after user requested OpenAI summaries."),
        )

        with patch(
            "backend.main.generate_openai_document_summary",
            return_value={
                "status": "SUCCESS",
                "model": "gpt-test",
                "summary": "OpenAI summary based on the stored document.",
            },
        ) as generator:
            response = generate_document_ai_analysis(
                document_id,
                DocumentAIAnalysisRequest(provider="openai"),
            )

        self.assertEqual(generator.call_count, 1)
        self.assertEqual(response["provider"], "openai")
        self.assertEqual(response["summary"], "OpenAI summary based on the stored document.")

    def test_openai_finding_summary_is_saved_with_model_provenance(self):
        self._add_new_information_document(
            1, "Finding summary document", "Alprazolam", "2026-09-26"
        )
        with patch(
            "backend.main.openai_provider_status",
            return_value={"available": True},
        ), patch(
            "backend.main.generate_openai_finding_summary",
            return_value={
                "status": "SUCCESS",
                "model": "gpt-test",
                "summary": "Evidence-only summary.",
                "prac_assessment": "Assessment from evidence.",
                "outcome": "No outcome stated.",
            },
        ):
            response = generate_summary(1, provider="openai")

        self.assertTrue(response["success"])
        self.assertEqual(response["model"], "gpt-test")
        saved = self.connection.execute(
            "SELECT summary, model FROM summaries WHERE finding_id = 1"
        ).fetchone()
        self.assertEqual(saved["summary"], "Evidence-only summary.")
        self.assertEqual(saved["model"], "gpt-test")

    def test_document_without_extracted_text_persists_analysis_failure(self):
        self.connection.execute(
            """
            INSERT INTO documents
                (source_id, title, url, status, created_at)
            VALUES (1, 'Unextracted document', 'https://example.test/one.pdf',
                    'NEW', '2026-09-28T10:30:00')
            """
        )
        document_id = self.connection.execute(
            "SELECT id FROM documents WHERE title = 'Unextracted document'"
        ).fetchone()["id"]

        with self.assertRaises(HTTPException) as error:
            generate_document_ai_analysis(
                document_id,
                DocumentAIAnalysisRequest(provider="ollama"),
            )

        self.assertEqual(error.exception.status_code, 422)
        persisted = get_document_ai_analysis(document_id, provider="ollama")
        self.assertEqual(persisted["status"], "FAILED")
        self.assertIn("no extracted text", persisted["error_message"])

    def test_document_ai_summary_download_contains_saved_document_metadata(self):
        self.connection.execute(
            """
            INSERT INTO documents
                (source_id, title, url, publication_date, reference_id,
                 status, created_at)
            VALUES (1, 'Published minutes', 'https://example.test/minutes.pdf',
                    '25 September 2026', 'EMA_PRAC_001_20260925_001',
                    'PROCESSED', '2026-09-28T10:30:00')
            """
        )
        document_id = self.connection.execute(
            "SELECT id FROM documents WHERE title = 'Published minutes'"
        ).fetchone()["id"]
        self.connection.execute(
            """
            INSERT INTO document_ai_analyses
                (document_id, provider, analysis_type, summary, status,
                 created_at, updated_at)
            VALUES (?, 'ollama', 'summary', 'Persisted AI summary',
                    'COMPLETED', '2026-10-01T00:00:00', '2026-10-01T00:00:00')
            """,
            (document_id,),
        )

        response = download_document_ai_analysis(document_id, provider="ollama")

        self.assertEqual(response.status_code, 200)
        self.assertIn("attachment;", response.headers["content-disposition"])
        self.assertIn("EMA_PRAC_001_20260925_001", response.body.decode())
        self.assertIn("Persisted AI summary", response.body.decode())

    def test_original_pdf_download_fetches_only_from_configured_source_host(self):
        self.connection.execute(
            """
            INSERT INTO documents
                (source_id, title, url, status, created_at)
            VALUES (
                1, 'Remote minutes',
                'https://example.test/prac/documents/minutes.pdf',
                'NEW', '2026-09-28T10:30:00'
            )
            """
        )
        document_id = self.connection.execute(
            "SELECT id FROM documents WHERE title = 'Remote minutes'"
        ).fetchone()["id"]
        with tempfile.TemporaryDirectory() as directory, patch(
            "backend.main.PDF_DIR", Path(directory)
        ), patch("backend.main.download_pdf") as downloader:
            def write_pdf(url, destination):
                destination_path = Path(destination)
                destination_path.parent.mkdir(parents=True, exist_ok=True)
                destination_path.write_bytes(b"%PDF-1.4 test")
                return destination

            downloader.side_effect = write_pdf
            response = download_document_pdf(document_id)

            self.assertEqual(response.media_type, "application/pdf")
            self.assertEqual(downloader.call_args.args[0],
                             "https://example.test/prac/documents/minutes.pdf")
            downloaded_path = Path(response.path)
            self.assertTrue(downloaded_path.is_file())

            asyncio.run(response.background())
            self.assertFalse(downloaded_path.exists())

        self.connection.execute(
            """
            INSERT INTO documents
                (source_id, title, url, status, created_at)
            VALUES (1, 'External PDF', 'https://other.example/external.pdf',
                    'NEW', '2026-09-28T10:30:00')
            """
        )
        external_id = self.connection.execute(
            "SELECT id FROM documents WHERE title = 'External PDF'"
        ).fetchone()["id"]
        with self.assertRaises(HTTPException) as error:
            download_document_pdf(external_id)
        self.assertEqual(error.exception.status_code, 400)

    def test_ollama_document_summary_uses_one_document_text_without_fallback(self):
        with patch(
            "backend.ollama_client.test_ollama_connection",
            return_value={
                "available": True,
                "model_available": True,
                "model": "test-model",
            },
        ), patch(
            "backend.ollama_client._call_ollama",
            return_value={
                "status": "SUCCESS",
                "model": "test-model",
                "data": {"summary": "Evidence-based document summary."},
            },
        ) as call:
            result = generate_document_summary(
                "One stored document",
                "Text extracted from that document only.",
            )

        self.assertEqual(result["status"], "SUCCESS")
        self.assertEqual(result["summary"], "Evidence-based document summary.")
        self.assertIn("One stored document", call.call_args.args[0])
        self.assertIn("Text extracted from that document only.", call.call_args.args[0])

    def test_manual_crawl_persists_success_counts_and_deduplicates(self):
        discovered = {
            "title": "PRAC Agenda",
            "url": "https://example.test/prac/agenda.pdf",
            "document_key": "https://example.test/prac/agenda.pdf",
            "publication_date": "25 September 2026",
            "content_type": "Agendas",
        }

        def crawl_result(*_args, **_kwargs):
            return {
                "success": True,
                "documents": [discovered.copy()],
                "documents_found": 1,
            }

        with patch("backend.main.crawl_source", side_effect=crawl_result) as crawler:
            first_result = run_manual_crawl(
                CrawlDiscoverRequest(source_id=1, content_types=["Agendas"]),
            )
            second_result = run_manual_crawl(
                CrawlDiscoverRequest(source_id=1, content_types=["Agendas"]),
            )

        self.assertEqual(crawler.call_args.kwargs["content_types"], ["Agendas"])
        self.assertEqual(first_result["new_documents"], 1)
        self.assertEqual(second_result["new_documents"], 0)
        self.assertEqual(
            self.connection.execute("SELECT COUNT(*) FROM documents").fetchone()[0],
            1,
        )
        self.assertEqual(
            self.connection.execute(
                "SELECT COUNT(*) FROM crawl_history WHERE status = 'success'"
            ).fetchone()[0],
            2,
        )

    def test_manual_crawl_failure_is_recorded_for_source_status(self):
        with patch(
            "backend.main.crawl_source",
            side_effect=RuntimeError("crawler unavailable"),
        ):
            with self.assertRaises(HTTPException) as error:
                run_manual_crawl(CrawlDiscoverRequest(source_id=1))

        self.assertEqual(error.exception.status_code, 502)
        source = list_sources()[0]
        self.assertEqual(source["crawl_count"], 0)
        self.assertEqual(source["crawl_status"], "failed")
        self.assertEqual(source["latest_new_documents"], 0)
        self.assertEqual(
            self.connection.execute(
                "SELECT COUNT(*) FROM crawl_history WHERE status = 'failed'"
            ).fetchone()[0],
            1,
        )

    def test_monitoring_worker_processes_sqlite_source_rows(self):
        job = {
            "sources_processed": 0,
            "sources_failed": 0,
            "sources_without_new_documents": 0,
        }
        with patch.dict("backend.main.monitoring_jobs", {"test-job": job}), patch(
            "backend.main.crawl_source",
            return_value={"success": True, "documents": [], "documents_found": 0},
        ):
            monitoring_worker("test-job")

        self.assertEqual(job["status"], "completed")
        self.assertEqual(job["sources_processed"], 1)
        self.assertEqual(job["sources_failed"], 0)
        self.assertEqual(job["sources_without_new_documents"], 1)
        self.assertEqual(
            self.connection.execute(
                "SELECT last_crawl_status FROM sources WHERE id = 1"
            ).fetchone()[0],
            "success",
        )


if __name__ == "__main__":
    unittest.main()
