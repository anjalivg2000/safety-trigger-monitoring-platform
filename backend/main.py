import logging
import json
import os
import uuid
import re
import sys
import threading
import fitz
from calendar import monthrange
from copy import copy
from html import escape as html_escape
from datetime import date, datetime, time as datetime_time, timedelta
from io import BytesIO
from pathlib import Path
from typing import Optional
from urllib.parse import urlparse

from openpyxl import Workbook
from fastapi import FastAPI, HTTPException, BackgroundTasks, Query
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, Response
from starlette.background import BackgroundTask
from fastapi.staticfiles import StaticFiles

from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from reportlab.lib import colors

if __package__ in (None, ""):
    PROJECT_ROOT = Path(__file__).resolve().parent.parent
    if str(PROJECT_ROOT) not in sys.path:
        sys.path.insert(0, str(PROJECT_ROOT))
    from backend.database import get_db, init_db
    from backend.models import (
        SourceCreate,
        SourceUpdate,
        SourceResponse,
        ProductClassCreate,
        ProductClassUpdate,
        ProductClassResponse,
        ProductCreate,
        ProductUpdate,
        ProductResponse,
        SourceTestRequest,
        CrawlDiscoverRequest,
        CrawlProcessRequest,
        DocumentAIAnalysisRequest,
        DocumentResponse,
        DocumentPageResponse,
        FindingResponse,
        SummaryResponse,
        ReviewCreate,
        ReviewUpdate,
        ReviewResponse,
        TriggerConfig,
    )
    from backend.crawler import (
        DEFAULT_SOURCE_CONFIGS,
        EMAPRAC_URL,
        crawl_source,
        discover_meeting_pdfs,
        download_pdf,
        classify_content_type,
        deduplicate_documents,
        normalize_content_types,
        test_source,
    )
    from backend.product_config import DEFAULT_PRODUCT_CLASSES
    from backend.pdf_processor import (
        process_pdf_safe,
        get_pdf_metadata,
    )
    from backend.detector import (
        detect_document,
        load_trigger_config,
    )
    from backend.comparator import (
        compare_finding,
        remove_duplicate_findings,
    )
    from backend.ollama_client import (
        generate_document_summary,
        generate_summary_from_finding,
        test_ollama_connection,
    )
    from backend.openai_client import (
        generate_document_summary as generate_openai_document_summary,
        generate_finding_summary as generate_openai_finding_summary,
        openai_provider_status,
    )
    from backend.topic_extractor import (
        extract_meeting_topics,
        extract_section_summaries,
    )
else:
    from .database import get_db, init_db
    from .models import (
        SourceCreate,
        SourceUpdate,
        SourceResponse,
        ProductClassCreate,
        ProductClassUpdate,
        ProductClassResponse,
        ProductCreate,
        ProductUpdate,
        ProductResponse,
        SourceTestRequest,
        CrawlDiscoverRequest,
        CrawlProcessRequest,
        DocumentAIAnalysisRequest,
        DocumentResponse,
        DocumentPageResponse,
        FindingResponse,
        SummaryResponse,
        ReviewCreate,
        ReviewUpdate,
        ReviewResponse,
        TriggerConfig,
    )
    from .crawler import (
        DEFAULT_SOURCE_CONFIGS,
        EMAPRAC_URL,
        crawl_source,
        discover_meeting_pdfs,
        download_pdf,
        classify_content_type,
        deduplicate_documents,
        normalize_content_types,
        test_source,
    )
    from .product_config import DEFAULT_PRODUCT_CLASSES
    from .pdf_processor import (
        process_pdf_safe,
        get_pdf_metadata,
    )
    from .detector import (
        detect_document,
        load_trigger_config,
    )
    from .comparator import (
        compare_finding,
        remove_duplicate_findings,
    )
    from .ollama_client import (
        generate_document_summary,
        generate_summary_from_finding,
        test_ollama_connection,
    )
    from .openai_client import (
        generate_document_summary as generate_openai_document_summary,
        generate_finding_summary as generate_openai_finding_summary,
        openai_provider_status,
    )
    from .topic_extractor import (
        extract_meeting_topics,
        extract_section_summaries,
    )

# ---------------------------------------------------------
# Logging
# ---------------------------------------------------------

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s | %(levelname)s | %(name)s | %(message)s",
)

logger = logging.getLogger("safety-trigger-poc")

# ---------------------------------------------------------
# Paths
# ---------------------------------------------------------

BASE_DIR = Path(__file__).resolve().parent.parent
FRONTEND_DIR = BASE_DIR / "frontend"
PDF_DIR = BASE_DIR / "data" / "pdfs"
SUMMARY_DIR = BASE_DIR / "data" / "summaries"

SUMMARY_DIR.mkdir(
    parents=True,
    exist_ok=True
)

PRAC_SECTION_XPATH = "//h2[normalize-space()='PRAC minutes']"

PRAC_MEETING_XPATH = (
    "//a[contains(@href, "
    "'/documents/minutes/minutes-prac-meeting-')]"
)
PDF_DIR.mkdir(parents=True, exist_ok=True)

# ---------------------------------------------------------
# FastAPI
# ---------------------------------------------------------

app = FastAPI(
    title="Safety Trigger POC",
    description="EMA PRAC Safety Trigger Monitoring POC",
    version="1.0.0",
)
_topic_summary_scheduler_stop = threading.Event()
_topic_summary_scheduler_thread = None
_crawl_scheduler_stop = threading.Event()
_crawl_scheduler_thread = None

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# ---------------------------------------------------------
# Startup
# ---------------------------------------------------------

@app.on_event("startup")
def startup():
    logger.info("Initializing database...")
    init_db()

    with get_db() as conn:
        for source_config in DEFAULT_SOURCE_CONFIGS:
            reference_no = source_config["reference_no"]
            url = source_config["url"]

            existing = conn.execute(
                """
                SELECT id FROM sources
                WHERE reference_no = ? OR url = ?
                LIMIT 1
                """,
                (reference_no, url),
            ).fetchone()

            payload = {
                "reference_no": reference_no,
                "authority": "European Medicines Agency",
                "name": source_config["name"],
                "source_type": source_config["source_type"],
                "region": source_config["region"],
                "monitoring_frequency": source_config["monitoring_frequency"],
                "website": source_config["website"],
                "safety_alert_type": source_config["safety_alert_type"],
                "url": url,
                "content_types": source_config["content_types"],
                "instruction": source_config.get("instruction", ""),
                "active": int(bool(source_config.get("active", True))),
                "updated_at": datetime.utcnow().isoformat(),
            }

            if not existing:
                conn.execute(
                    """
                    INSERT INTO sources
                    (
                        reference_no,
                        authority,
                        name,
                        source_type,
                        region,
                        monitoring_frequency,
                        website,
                        safety_alert_type,
                        url,
                        content_types,
                        instruction,
                        active,
                        created_at,
                        updated_at
                    )
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        payload["reference_no"],
                        payload["authority"],
                        payload["name"],
                        payload["source_type"],
                        payload["region"],
                        payload["monitoring_frequency"],
                        payload["website"],
                        payload["safety_alert_type"],
                        payload["url"],
                        json.dumps(payload["content_types"]),
                        payload["instruction"],
                        payload["active"],
                        datetime.utcnow().isoformat(),
                        payload["updated_at"],
                    ),
                )

        for class_config in DEFAULT_PRODUCT_CLASSES:
            class_reference = class_config["reference_no"]
            existing_class = conn.execute(
                "SELECT id FROM product_classes WHERE reference_no = ? LIMIT 1",
                (class_reference,),
            ).fetchone()

            class_payload = {
                "reference_no": class_reference,
                "class_name": class_config["name"],
                "active": int(bool(class_config.get("active", True))),
                "updated_at": datetime.utcnow().isoformat(),
            }

            if existing_class:
                conn.execute(
                    """
                    UPDATE product_classes
                    SET reference_no = ?,
                        class_name = ?,
                        active = ?,
                        updated_at = ?
                    WHERE id = ?
                    """,
                    (
                        class_payload["reference_no"],
                        class_payload["class_name"],
                        class_payload["active"],
                        class_payload["updated_at"],
                        existing_class["id"],
                    ),
                )
                class_id = existing_class["id"]
            else:
                cursor = conn.execute(
                    """
                    INSERT INTO product_classes
                    (reference_no, class_name, active, created_at, updated_at)
                    VALUES (?, ?, ?, ?, ?)
                    """,
                    (
                        class_payload["reference_no"],
                        class_payload["class_name"],
                        class_payload["active"],
                        datetime.utcnow().isoformat(),
                        class_payload["updated_at"],
                    ),
                )
                class_id = cursor.lastrowid

            for product_index, product_name in enumerate(class_config["products"], start=1):
                product_reference = f"{class_reference}_P{product_index:03d}"
                existing_product = conn.execute(
                    "SELECT id FROM products WHERE reference_no = ? LIMIT 1",
                    (product_reference,),
                ).fetchone()

                payload = {
                    "product_class_id": class_id,
                    "reference_no": product_reference,
                    "product_name": product_name,
                    "active": 1,
                    "updated_at": datetime.utcnow().isoformat(),
                }

                if existing_product:
                    conn.execute(
                        """
                        UPDATE products
                        SET product_class_id = ?,
                            reference_no = ?,
                            product_name = ?,
                            active = ?,
                            updated_at = ?
                        WHERE id = ?
                        """,
                        (
                            payload["product_class_id"],
                            payload["reference_no"],
                            payload["product_name"],
                            payload["active"],
                            payload["updated_at"],
                            existing_product["id"],
                        ),
                    )
                else:
                    conn.execute(
                        """
                        INSERT INTO products
                        (product_class_id, reference_no, product_name, active, created_at, updated_at)
                        VALUES (?, ?, ?, ?, ?, ?)
                        """,
                        (
                            payload["product_class_id"],
                            payload["reference_no"],
                            payload["product_name"],
                            payload["active"],
                            datetime.utcnow().isoformat(),
                            payload["updated_at"],
                        ),
                    )

    logger.info("Default EMA source configuration loaded.")
    logger.info("Default product configuration loaded.")

    global _topic_summary_scheduler_thread
    if _topic_summary_scheduler_thread is None or not _topic_summary_scheduler_thread.is_alive():
        _topic_summary_scheduler_stop.clear()
        _topic_summary_scheduler_thread = threading.Thread(
            target=_topic_summary_scheduler_loop,
            name="topic-summary-scheduler",
            daemon=True,
        )
        _topic_summary_scheduler_thread.start()

    global _crawl_scheduler_thread
    if _crawl_scheduler_thread is None or not _crawl_scheduler_thread.is_alive():
        _crawl_scheduler_stop.clear()
        _crawl_scheduler_thread = threading.Thread(
            target=_crawl_scheduler_loop,
            name="configured-crawl-scheduler",
            daemon=True,
        )
        _crawl_scheduler_thread.start()


@app.on_event("shutdown")
def shutdown_topic_summary_scheduler():
    _topic_summary_scheduler_stop.set()
    if _topic_summary_scheduler_thread and _topic_summary_scheduler_thread.is_alive():
        _topic_summary_scheduler_thread.join(timeout=3)
    _crawl_scheduler_stop.set()
    if _crawl_scheduler_thread and _crawl_scheduler_thread.is_alive():
        _crawl_scheduler_thread.join(timeout=3)


def _next_monitoring_run(anchor: datetime, frequency: str) -> Optional[datetime]:
    normalized = " ".join(str(frequency or "").strip().casefold().split())
    if normalized == "every alternate day":
        return anchor + timedelta(days=2)
    if normalized == "weekly":
        return anchor + timedelta(days=7)
    if normalized == "monthly":
        next_month = anchor.month % 12 + 1
        next_year = anchor.year + (anchor.month // 12)
        last_day = monthrange(next_year, next_month)[1]
        day = last_day if anchor.day == monthrange(anchor.year, anchor.month)[1] else min(anchor.day, last_day)
        return anchor.replace(year=next_year, month=next_month, day=day)
    return None


def _schedule_timestamp(value: str) -> datetime:
    return datetime.fromisoformat(str(value).replace("Z", "+00:00")).replace(tzinfo=None)


def _claim_scheduled_crawls(now: datetime) -> list[dict]:
    claims = []
    with get_db() as conn:
        conn.execute("BEGIN IMMEDIATE")
        sources = conn.execute(
            """
            SELECT id, reference_no, name, monitoring_frequency
            FROM sources
            WHERE active = 1
            """
        ).fetchall()

        for source in sources:
            source_id = int(source["id"])
            configured = conn.execute(
                """
                SELECT id, created_at, details
                FROM audit_events
                WHERE source_id = ? AND event_type = 'CRAWL_SCHEDULE_CONFIGURED'
                ORDER BY id DESC
                LIMIT 1
                """,
                (source_id,),
            ).fetchone()
            if not configured:
                continue

            try:
                schedule = json.loads(configured["details"])
                if (
                    not schedule.get("enabled")
                    or _next_monitoring_run(now, schedule.get("frequency")) is None
                ):
                    continue
            except (json.JSONDecodeError, AttributeError, TypeError):
                logger.exception("Invalid automatic crawl schedule for source %s", source_id)
                continue

            previous_run = conn.execute(
                """
                SELECT id, created_at, status
                FROM audit_events
                WHERE source_id = ? AND event_type = 'SCHEDULED_SOURCE_CRAWL'
                ORDER BY id DESC
                LIMIT 1
                """,
                (source_id,),
            ).fetchone()
            if previous_run and previous_run["status"] == "RUNNING":
                started_at = _schedule_timestamp(previous_run["created_at"])
                if now - started_at < timedelta(hours=1):
                    continue
                conn.execute(
                    """
                    UPDATE audit_events
                    SET status = 'FAILED',
                        result = 'Scheduled crawl was interrupted; no duplicate run was started.'
                    WHERE id = ?
                    """,
                    (int(previous_run["id"]),),
                )
                previous_run = conn.execute(
                    """
                    SELECT id, created_at, status
                    FROM audit_events
                    WHERE source_id = ? AND event_type = 'SCHEDULED_SOURCE_CRAWL'
                    ORDER BY id DESC
                    LIMIT 1
                    """,
                    (source_id,),
                ).fetchone()

            anchor = _schedule_timestamp(configured["created_at"])
            if previous_run:
                anchor = max(anchor, _schedule_timestamp(previous_run["created_at"]))
            next_run = _next_monitoring_run(anchor, schedule.get("frequency"))
            if next_run is None or now < next_run:
                continue

            details = json.dumps({
                "frequency": schedule["frequency"],
                "scheduled_at": now.isoformat(),
                "status": "RUNNING",
            })
            cursor = conn.execute(
                """
                INSERT INTO audit_events
                    (created_at, event_type, details, entity_type, entity_id,
                     source_id, reference_id, action, result, status)
                VALUES (?, 'SCHEDULED_SOURCE_CRAWL', ?, 'source', ?, ?, ?,
                        'SCHEDULED_CRAWL', 'Scheduled source discovery started.', 'RUNNING')
                """,
                (
                    now.isoformat(),
                    details,
                    source_id,
                    source_id,
                    source["reference_no"] or source["name"],
                ),
            )
            claims.append({
                "event_id": int(cursor.lastrowid),
                "source_id": source_id,
                "name": source["name"],
            })
    return claims


def _run_scheduled_source_crawl(claim: dict) -> None:
    event_id = int(claim["event_id"])
    source_id = int(claim["source_id"])
    try:
        result = run_manual_crawl(CrawlDiscoverRequest(source_id=source_id))
        succeeded = bool(result.get("success"))
        message = result.get("message") or result.get("error") or "Scheduled source discovery finished."
        details = json.dumps({
            "documents_found": result.get("documents_found", 0),
            "new_documents": result.get("new_documents", 0),
            "status": "COMPLETED" if succeeded else "FAILED",
        })
        status = "COMPLETED" if succeeded else "FAILED"
    except Exception as exc:
        logger.exception("Scheduled source crawl failed for source %s", source_id)
        details = json.dumps({"status": "FAILED", "error": str(exc)})
        status = "FAILED"
        message = str(exc)

    with get_db() as conn:
        conn.execute(
            """
            UPDATE audit_events
            SET details = ?, status = ?, result = ?
            WHERE id = ?
            """,
            (details, status, message, event_id),
        )
    logger.info("Scheduled source crawl finished for %s with status %s", claim["name"], status)


def _crawl_scheduler_loop() -> None:
    while not _crawl_scheduler_stop.wait(30):
        try:
            for claim in _claim_scheduled_crawls(datetime.utcnow()):
                _run_scheduled_source_crawl(claim)
        except Exception:
            logger.exception("Configured crawl scheduler failed.")


@app.get("/api/crawl/schedule/status/{source_id}")
def crawl_schedule_status(source_id: int):
    with get_db() as conn:
        source = conn.execute(
            "SELECT id, active, monitoring_frequency FROM sources WHERE id = ?",
            (source_id,),
        ).fetchone()
        if not source:
            raise HTTPException(status_code=404, detail="Source not found.")
        configured = conn.execute(
            """
            SELECT created_at, details FROM audit_events
            WHERE source_id = ? AND event_type = 'CRAWL_SCHEDULE_CONFIGURED'
            ORDER BY id DESC LIMIT 1
            """,
            (source_id,),
        ).fetchone()
        latest_run = conn.execute(
            """
            SELECT created_at, status, result FROM audit_events
            WHERE source_id = ? AND event_type = 'SCHEDULED_SOURCE_CRAWL'
            ORDER BY id DESC LIMIT 1
            """,
            (source_id,),
        ).fetchone()

    schedule = json.loads(configured["details"]) if configured else {
        "enabled": False,
        "frequency": source["monitoring_frequency"],
    }
    schedule_supported = _next_monitoring_run(datetime.utcnow(), schedule.get("frequency")) is not None
    thread_running = bool(_crawl_scheduler_thread and _crawl_scheduler_thread.is_alive())
    source_active = bool(source["active"])
    configured_enabled = bool(schedule.get("enabled")) and schedule_supported
    enabled = configured_enabled and thread_running and source_active
    anchor = configured["created_at"] if configured else None
    if latest_run:
        latest_run_at = _schedule_timestamp(latest_run["created_at"])
        anchor = max(anchor, latest_run_at) if anchor else latest_run_at
    next_run = _next_monitoring_run(_schedule_timestamp(anchor), schedule.get("frequency")) if (
        anchor and enabled
    ) else None
    next_run_at = f"{next_run.isoformat()}Z" if next_run else None
    if not configured_enabled:
        message = (
            "Manual mode does not schedule runs."
            if str(schedule.get("frequency") or "").casefold() == "manual"
            else "Automatic monitoring is not enabled. No scheduled runs will occur."
        )
    elif not source_active:
        message = "Automatic monitoring is paused because this source is inactive."
    elif not thread_running:
        message = "Automatic monitoring is configured but the scheduler is not running."
    else:
        message = f"Automatic monitoring is enabled ({schedule['frequency']}); scheduled runs use UTC."

    return {
        "source_id": source_id,
        "configured_enabled": configured_enabled,
        "enabled": enabled,
        "scheduler_running": thread_running,
        "frequency": schedule.get("frequency"),
        "next_run_at": next_run_at,
        "last_run": dict(latest_run) if latest_run else None,
        "message": message,
    }


def _claim_topic_summary_schedule(now: datetime) -> Optional[int]:
    if now.time() < datetime_time(hour=9):
        return None

    scheduled_date = now.date()
    with get_db() as conn:
        conn.execute("BEGIN IMMEDIATE")
        latest = conn.execute(
            """
            SELECT details FROM audit_events
            WHERE event_type = 'TOPIC_SUMMARY_SCHEDULE'
            ORDER BY id DESC
            LIMIT 1
            """
        ).fetchone()
        if latest:
            try:
                previous_date = date.fromisoformat(
                    json.loads(latest["details"])["scheduled_date"]
                )
            except (ValueError, KeyError, TypeError, json.JSONDecodeError):
                previous_date = now.date()
            if (scheduled_date - previous_date).days < 2:
                return None

        details = json.dumps({
            "scheduled_date": scheduled_date.isoformat(),
            "status": "RUNNING",
        })
        cursor = conn.execute(
            """
            INSERT INTO audit_events
                (created_at, event_type, details, entity_type, action, status)
            VALUES (?, 'TOPIC_SUMMARY_SCHEDULE', ?, 'FINDING', 'ALTERNATE_DAY', 'RUNNING')
            """,
            (now.isoformat(), details),
        )
        return int(cursor.lastrowid)


def _run_scheduled_topic_summaries(schedule_event_id: int, now: datetime) -> None:
    generated = 0
    failures = []
    provider_status = test_ollama_connection()
    if not provider_status.get("available") or not provider_status.get("model_available"):
        failures.append(
            provider_status.get("message")
            or f"Configured Ollama model '{os.getenv('OLLAMA_MODEL', 'llama3.2')}' is unavailable."
        )
    else:
        with get_db() as conn:
            pending = conn.execute(
                """
                SELECT f.id
                FROM findings f
                WHERE TRIM(COALESCE(f.evidence, '')) <> ''
                  AND NOT EXISTS (
                      SELECT 1 FROM summaries s WHERE s.finding_id = f.id
                  )
                ORDER BY f.created_at, f.id
                """
            ).fetchall()

        for row in pending:
            try:
                result = generate_summary(int(row["id"]), provider="ollama")
                if result.get("success"):
                    generated += 1
                else:
                    failures.append(
                        f"Finding {row['id']}: {result.get('message') or result.get('status')}"
                    )
            except Exception as exc:
                logger.exception("Scheduled topic summary failed for finding %s", row["id"])
                failures.append(f"Finding {row['id']}: {exc}")

    status = "FAILED" if failures and generated == 0 else "PARTIAL" if failures else "COMPLETED"
    details = json.dumps({
        "scheduled_date": now.date().isoformat(),
        "generated": generated,
        "failed": len(failures),
        "errors": failures[:20],
    })
    with get_db() as conn:
        conn.execute(
            """
            UPDATE audit_events
            SET details = ?, status = ?, result = ?
            WHERE id = ?
            """,
            (details, status, f"Generated {generated}; failed {len(failures)}.", schedule_event_id),
        )
    logger.info(
        "Alternate-day topic summary run finished: generated=%s failed=%s",
        generated,
        len(failures),
    )


def _topic_summary_scheduler_loop() -> None:
    while not _topic_summary_scheduler_stop.wait(30):
        now = datetime.now()
        try:
            event_id = _claim_topic_summary_schedule(now)
            if event_id is not None:
                _run_scheduled_topic_summaries(event_id, now)
        except Exception:
            logger.exception("Alternate-day topic summary scheduler failed.")


@app.get("/api/topic-summary/schedule")
def topic_summary_schedule():
    now = datetime.now()
    with get_db() as conn:
        latest = conn.execute(
            """
            SELECT created_at, details, status, result
            FROM audit_events
            WHERE event_type = 'TOPIC_SUMMARY_SCHEDULE'
            ORDER BY id DESC
            LIMIT 1
            """
        ).fetchone()

    last_run = None
    if latest:
        details = json.loads(latest["details"])
        last_run = {
            "created_at": latest["created_at"],
            "status": latest["status"],
            "result": latest["result"],
            "generated": details.get("generated", 0),
            "failed": details.get("failed", 0),
            "errors": details.get("errors", []),
        }
        next_date = date.fromisoformat(details["scheduled_date"]) + timedelta(days=2)
    else:
        next_date = now.date()
        if now.time() >= datetime_time(hour=9):
            next_date += timedelta(days=1)
    return {
        "schedule": "Every other day at 09:00 local server time",
        "next_eligible_at": datetime.combine(next_date, datetime_time(hour=9)).isoformat(),
        "last_run": last_run,
        "workflow_targets": {
            "publications": 20,
            "related_items_priority": 15,
            "turnaround_minutes": 15,
        },
    }


# ---------------------------------------------------------
# Frontend
# ---------------------------------------------------------

if FRONTEND_DIR.exists():
    app.mount(
        "/frontend",
        StaticFiles(directory=str(FRONTEND_DIR)),
        name="frontend",
    )


@app.get("/app.js", include_in_schema=False)
def serve_app_js():
    app_js = FRONTEND_DIR / "app.js"
    if not app_js.exists():
        raise HTTPException(status_code=404, detail="frontend/app.js not found")
    return FileResponse(
        app_js,
        media_type="application/javascript",
        headers={"Cache-Control": "no-cache"},
    )


@app.get("/style.css", include_in_schema=False)
def serve_style_css():
    css_file = FRONTEND_DIR / "style.css"
    if not css_file.exists():
        raise HTTPException(status_code=404, detail="frontend/style.css not found")
    return FileResponse(
        css_file,
        media_type="text/css",
        headers={"Cache-Control": "no-cache"},
    )


@app.get("/", include_in_schema=False)
def root():
    index_file = FRONTEND_DIR / "index.html"

    if not index_file.exists():
        return {
            "status": "ok",
            "message": "Safety Trigger POC backend is running.",
        }

    return FileResponse(index_file)


# ---------------------------------------------------------
# Health
# ---------------------------------------------------------

@app.get("/api/health")
def health():
    return {
        "status": "healthy",
        "application": "Safety Trigger POC",
        "timestamp": datetime.utcnow().isoformat(),
    }


# ---------------------------------------------------------
# Dashboard
# ---------------------------------------------------------

@app.get("/api/dashboard")
def dashboard():
    with get_db() as conn:
        source_counts = conn.execute(
            """
            SELECT
                COUNT(*) AS total,
                COALESCE(SUM(CASE WHEN active = 1 THEN 1 ELSE 0 END), 0) AS active,
                COALESCE(SUM(CASE WHEN active = 0 THEN 1 ELSE 0 END), 0) AS inactive
            FROM sources
            """
        ).fetchone()

        document_counts = conn.execute(
            """
            SELECT
                COUNT(*) AS total,
                COALESCE(SUM(CASE
                    WHEN UPPER(COALESCE(status, '')) IN ('PROCESSED', 'EXTRACTED')
                    THEN 1 ELSE 0 END), 0) AS working,
                COALESCE(SUM(CASE
                    WHEN UPPER(COALESCE(status, '')) IN
                        ('FAILED', 'EXTRACTION_FAILED', 'ERROR')
                    THEN 1 ELSE 0 END), 0) AS non_working,
                COALESCE(SUM(CASE
                    WHEN UPPER(COALESCE(status, '')) = 'NEW'
                    THEN 1 ELSE 0 END), 0) AS new_documents,
                COALESCE(SUM(CASE
                    WHEN last_crawled_at IS NOT NULL AND last_crawled_at != ''
                    THEN 1 ELSE 0 END), 0) AS crawled
            FROM documents
            """
        ).fetchone()

        crawl_counts = conn.execute(
            """
            SELECT COUNT(*) AS executions
            FROM crawl_history
            """
        ).fetchone()

        new_information = conn.execute(
            """
            SELECT COUNT(*) AS count
            FROM findings
            WHERE UPPER(COALESCE(comparison_status, ''))
                IN ('NEW', 'NEW_INFORMATION')
            """
        ).fetchone()["count"]

        pending_reviews = conn.execute(
            """
            SELECT COUNT(*) AS count
            FROM findings f
            WHERE NOT EXISTS (
                SELECT 1 FROM reviews r WHERE r.finding_id = f.id
            )
            """
        ).fetchone()["count"]

        review_workflow_counts = conn.execute(
            """
            SELECT
                (
                    SELECT COUNT(*)
                    FROM findings f
                    WHERE NOT EXISTS (
                        SELECT 1 FROM reviews r
                        WHERE r.finding_id = f.id AND r.stage = 'INITIAL'
                    )
                ) AS pending_initial_reviews,
                (
                    SELECT COUNT(*)
                    FROM findings f
                    WHERE EXISTS (
                        SELECT 1 FROM reviews ri
                        WHERE ri.finding_id = f.id AND ri.stage = 'INITIAL'
                    )
                    AND NOT EXISTS (
                        SELECT 1 FROM reviews rf
                        WHERE rf.finding_id = f.id AND rf.stage = 'FINAL'
                    )
                ) AS pending_final_reviews
            """
        ).fetchone()

        review_counts = conn.execute(
            """
            SELECT COUNT(*) AS total,
                   COUNT(DISTINCT finding_id) AS completed
            FROM reviews
            """
        ).fetchone()

        summaries = conn.execute(
            "SELECT COUNT(*) AS count FROM summaries"
        ).fetchone()["count"]

        safety_triggers = conn.execute(
            "SELECT COUNT(*) AS count FROM findings"
        ).fetchone()["count"]

        unique_products = conn.execute(
            """
            SELECT COUNT(DISTINCT lower(trim(
                COALESCE(NULLIF(trim(product_name), ''),
                         NULLIF(trim(active_substance), ''))
            ))) AS count
            FROM findings
            WHERE trim(COALESCE(
                NULLIF(trim(product_name), ''),
                NULLIF(trim(active_substance), ''),
                ''
            )) != ''
            """
        ).fetchone()["count"]

        content_type_rows = conn.execute(
            """
            SELECT content_type, COUNT(*) AS count
            FROM documents
            WHERE content_type IS NOT NULL AND trim(content_type) != ''
            GROUP BY lower(trim(content_type))
            """
        ).fetchall()

        recent_crawl_rows = conn.execute(
            """
            SELECT
                ch.id,
                ch.source_id,
                ch.started_at,
                ch.completed_at,
                ch.status,
                ch.documents_found,
                ch.new_documents,
                ch.total_documents,
                ch.error_message,
                s.reference_no AS source_reference_no,
                s.name AS source_name
            FROM crawl_history ch
            LEFT JOIN sources s ON s.id = ch.source_id
            ORDER BY COALESCE(ch.completed_at, ch.started_at, '') DESC, ch.id DESC
            LIMIT 8
            """
        ).fetchall()

        monitoring_status_rows = conn.execute(
            """
            SELECT id, reference_no, name, monitoring_frequency, active
            FROM sources
            ORDER BY active DESC, name COLLATE NOCASE, id
            """
        ).fetchall()

        recent_document_rows = conn.execute(
            """
            SELECT
                d.id,
                d.title,
                d.content_type,
                d.publication_date,
                d.last_crawled_at,
                d.status,
                d.url,
                s.name AS source_name
            FROM documents d
            LEFT JOIN sources s ON s.id = d.source_id
            ORDER BY COALESCE(
                d.created_at, d.discovered_date, d.last_crawled_at, ''
            ) DESC, d.id DESC
            LIMIT 10
            """
        ).fetchall()

        crawled_today = conn.execute(
            """
            SELECT COUNT(DISTINCT source_id) AS count
            FROM crawl_history
            WHERE source_id IS NOT NULL
              AND substr(COALESCE(completed_at, started_at, ''), 1, 10)
                    = date('now')
            """
        ).fetchone()["count"]

        source_health_rows = conn.execute(
            """
            SELECT
                s.id,
                s.active,
                s.last_crawl_status,
                s.last_error_message,
                (
                    SELECT ch.status
                    FROM crawl_history ch
                    WHERE ch.source_id = s.id
                    ORDER BY COALESCE(ch.completed_at, ch.started_at, '') DESC,
                             ch.id DESC
                    LIMIT 1
                ) AS latest_crawl_status
            FROM sources s
            """
        ).fetchall()

        last_crawl_row = conn.execute(
            """
            SELECT completed_at, started_at
            FROM crawl_history
            ORDER BY COALESCE(completed_at, started_at, '') DESC, id DESC
            LIMIT 1
            """
        ).fetchone()

        latest_documents_count = conn.execute(
            """
            SELECT COUNT(*) AS count
            FROM documents
            WHERE substr(COALESCE(created_at, discovered_date, ''), 1, 10)
                >= date('now', '-6 days')
            """
        ).fetchone()["count"]

        source_rows = conn.execute(
            "SELECT content_types FROM sources"
        ).fetchall()

    content_type_counts = {}
    for row in content_type_rows:
        content_type_counts[str(row["content_type"]).strip().casefold()] = row["count"]

    standard_content_types = {
        "meeting minutes": "Meeting Minutes",
        "agendas": "Agendas",
        "highlights": "Highlights",
    }
    content_type_summary = {
        label: content_type_counts.get(key, 0)
        for key, label in standard_content_types.items()
    }
    content_type_summary["Latest Documents"] = latest_documents_count

    configured_content_types = set()
    for row in source_rows:
        configured_content_types.update(normalize_content_types(row["content_types"]))
    other_content_types = sorted(
        configured_content_types
        | {
            str(row["content_type"]).strip()
            for row in content_type_rows
            if str(row["content_type"]).strip().casefold()
            not in standard_content_types
        },
        key=str.casefold,
    )
    content_type_summary["other_configured"] = [
        {
            "content_type": content_type,
            "count": content_type_counts.get(content_type.casefold(), 0),
        }
        for content_type in other_content_types
        if content_type.casefold() not in standard_content_types
    ]

    recent_crawl_activity = []
    for row in recent_crawl_rows:
        item = dict(row)
        item["source_id"] = item["source_reference_no"] or item["source_id"]
        item["result"] = (
            f"{item['new_documents']} new documents"
            if item["new_documents"] is not None
            else "New document count unavailable"
        )
        recent_crawl_activity.append(item)

    error_sources = [
        row
        for row in source_health_rows
        if (
            str(row["last_error_message"] or "").strip()
            or str(
                row["latest_crawl_status"] or row["last_crawl_status"] or ""
            ).upper()
            in {"FAILED", "ERROR"}
        )
    ]
    system_operational = not error_sources
    if error_sources:
        attention_count = len(error_sources)
        system_message = (
            f"{attention_count} source requires attention"
            if attention_count == 1
            else f"{attention_count} sources require attention"
        )
    else:
        attention_count = 0
        system_message = "System is running normally."

    latest_crawl_at = (
        last_crawl_row["completed_at"] or last_crawl_row["started_at"]
        if last_crawl_row
        else None
    )
    source_status_summary = {
        "active_sources": source_counts["active"],
        "inactive_sources": source_counts["inactive"],
        "sources_crawled_today": crawled_today,
        "sources_with_errors": attention_count,
    }
    document_summary = {
        "total_documents": document_counts["total"],
        "new_documents": document_counts["new_documents"],
        "documents_crawled": document_counts["crawled"],
        "working_documents": document_counts["working"],
        "non_working_documents": document_counts["non_working"],
    }

    return {
        "total_sources": source_counts["total"],
        "active_sources": source_counts["active"],
        "total_documents": document_counts["total"],
        "new_documents": document_counts["new_documents"],
        "crawl_executions": crawl_counts["executions"],
        "last_crawl_at": latest_crawl_at,
        "safety_triggers": safety_triggers,
        "unique_products": unique_products,
        "recent_crawl_activity": recent_crawl_activity,
        "document_summary": document_summary,
        "content_type_summary": content_type_summary,
        "source_status_summary": source_status_summary,
        "system_status": {
            "operational": system_operational,
            "attention_sources": attention_count,
            "message": system_message,
        },
        "sources": source_counts["active"],
        "documents": document_counts["total"],
        "processed_documents": document_counts["working"],
        "findings": safety_triggers,
        "new_information": new_information,
        "pending_reviews": pending_reviews,
        "pending_initial_reviews": review_workflow_counts["pending_initial_reviews"],
        "pending_final_reviews": review_workflow_counts["pending_final_reviews"],
        "approved_reviews": None,
        "rejected_reviews": None,
        "review_outcomes_note": (
            "The current review workflow records relevance decisions, "
            "not approval or rejection outcomes."
        ),
        "completed_reviews": review_counts["completed"],
        "reviews": review_counts["total"],
        "summaries": summaries,
        "recent_crawls": [
            {
                **dict(row),
                "next_scheduled_run": None,
            }
            for row in recent_crawl_rows
        ],
        "monitoring_status": [dict(row) for row in monitoring_status_rows],
        "recent_documents": [dict(row) for row in recent_document_rows],
    }


# ---------------------------------------------------------
# Sources
# ---------------------------------------------------------

def _record_audit_event(
    conn,
    event_type: str,
    details: str,
    *,
    entity_type: Optional[str] = None,
    entity_id: Optional[int] = None,
    source_id: Optional[int] = None,
    reference_id: Optional[str] = None,
    action: Optional[str] = None,
    ai_model: Optional[str] = None,
    result: Optional[str] = None,
    status: str = "COMPLETED",
) -> None:
    conn.execute(
        """
        INSERT INTO audit_events
            (created_at, event_type, details, entity_type, entity_id,
             source_id, reference_id, action, ai_model, result, status)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            datetime.utcnow().isoformat(),
            event_type,
            details,
            entity_type,
            entity_id,
            source_id,
            reference_id,
            action,
            ai_model,
            result,
            status,
        ),
    )


def serialize_source_row(row):
    data = dict(row)
    data["content_types"] = normalize_content_types(data.get("content_types"))
    data["reference_no"] = data.get("reference_no") or data.get("id")
    return data


@app.get("/api/sources")
def list_sources():
    with get_db() as conn:
        rows = conn.execute(
            """
            SELECT s.*,
                (
                    SELECT COUNT(*)
                    FROM crawl_history ch
                    WHERE ch.source_id = s.id
                      AND LOWER(COALESCE(ch.status, '')) IN ('success', 'completed')
                ) AS crawl_count,
                (
                    SELECT ch.status
                    FROM crawl_history ch
                    WHERE ch.source_id = s.id
                    ORDER BY COALESCE(ch.completed_at, ch.started_at) DESC, ch.id DESC
                    LIMIT 1
                ) AS crawl_status,
                (
                    SELECT ch.new_documents
                    FROM crawl_history ch
                    WHERE ch.source_id = s.id
                    ORDER BY COALESCE(ch.completed_at, ch.started_at) DESC, ch.id DESC
                    LIMIT 1
                ) AS latest_new_documents,
                (
                    SELECT ch.total_documents
                    FROM crawl_history ch
                    WHERE ch.source_id = s.id
                    ORDER BY COALESCE(ch.completed_at, ch.started_at) DESC, ch.id DESC
                    LIMIT 1
                ) AS latest_total_documents,
                (
                    SELECT COUNT(*)
                    FROM documents d
                    WHERE d.source_id = s.id
                      AND (
                          (
                              d.content_type IN ('Meeting Minutes', 'Agendas')
                              AND LOWER(COALESCE(d.url, ''))
                                  LIKE '%/documents/%.pdf%'
                          )
                          OR (
                              d.content_type = 'Highlights'
                              AND LOWER(COALESCE(d.url, ''))
                                  LIKE '%/en/news/meeting-highlights-%'
                          )
                          OR (
                              d.content_type = 'PRAC recommendation on signals'
                              AND LOWER(COALESCE(d.url, '')) LIKE '%.pdf%'
                          )
                          OR COALESCE(d.content_type, '') NOT IN (
                              'Meeting Minutes',
                              'Agendas',
                              'Highlights',
                              'PRAC recommendation on signals'
                          )
                      )
                ) AS total_documents
            FROM sources s
            ORDER BY s.id
            """
        ).fetchall()

    return [serialize_source_row(row) for row in rows]


def _crawl_content_types(source_data, requested_content_types):
    configured_content_types = normalize_content_types(
        source_data.get("content_types")
    )
    if not configured_content_types:
        configured_content_types = normalize_content_types(
            [source_data.get("safety_alert_type")]
        )

    if requested_content_types is None:
        return configured_content_types

    selected_content_types = normalize_content_types(requested_content_types)
    if not selected_content_types:
        raise HTTPException(
            status_code=400,
            detail="Select at least one configured content type.",
        )
    unconfigured_types = set(selected_content_types) - set(configured_content_types)
    if unconfigured_types:
        raise HTTPException(
            status_code=400,
            detail="Selected content types must be configured for this source.",
        )
    return selected_content_types


@app.post("/api/sources")
def create_source(source: SourceCreate):
    now = datetime.utcnow().isoformat()
    reference_no = source.reference_no or f"{source.website.upper()}_{source.name.strip().upper().replace(' ', '_')[:12]}_001"

    try:
        with get_db() as conn:
            existing = conn.execute(
                "SELECT id FROM sources WHERE reference_no = ? OR url = ? LIMIT 1",
                (reference_no, str(source.url)),
            ).fetchone()

            if existing:
                raise HTTPException(
                    status_code=400,
                    detail="A source with this reference number or URL already exists.",
                )

            cursor = conn.execute(
                """
                INSERT INTO sources
                (
                    reference_no,
                    authority,
                    name,
                    source_type,
                    region,
                    monitoring_frequency,
                    website,
                    safety_alert_type,
                    url,
                    content_types,
                    instruction,
                    active,
                    created_at,
                    updated_at
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    reference_no,
                    source.authority,
                    source.name,
                    source.source_type,
                    source.region,
                    source.monitoring_frequency,
                    source.website,
                    source.safety_alert_type,
                    str(source.url),
                    json.dumps(source.content_types),
                    source.instruction,
                    source.active,
                    now,
                    now,
                ),
            )

            source_id = cursor.lastrowid
            _record_audit_event(
                conn,
                "Source Created",
                reference_no,
                entity_type="source",
                entity_id=source_id,
                source_id=source_id,
                reference_id=reference_no,
            )

        return {
            "success": True,
            "id": source_id,
            "message": "Source created successfully.",
        }

    except HTTPException:
        raise
    except Exception as exc:
        logger.exception("Failed to create source")
        raise HTTPException(
            status_code=400,
            detail=str(exc),
        )


@app.patch("/api/sources/{source_id}")
def update_source(source_id: int, source: SourceUpdate):
    updates = source.dict(exclude_unset=True)
    automatic_monitoring = updates.pop("automatic_monitoring", None)
    if not updates and automatic_monitoring is None:
        raise HTTPException(status_code=400, detail="No source fields supplied.")

    if "content_types" in updates:
        updates["content_types"] = json.dumps(updates["content_types"])

    if updates:
        updates["updated_at"] = datetime.utcnow().isoformat()
    assignments = ", ".join(f"{field} = ?" for field in updates)
    values = list(updates.values()) + [source_id]

    try:
        with get_db() as conn:
            existing = conn.execute(
                """
                SELECT reference_no, name, active, monitoring_frequency
                FROM sources WHERE id = ?
                """,
                (source_id,),
            ).fetchone()
            if not existing:
                raise HTTPException(status_code=404, detail="Source not found.")
            source_active_after = bool(updates.get("active", existing["active"]))
            frequency = updates.get("monitoring_frequency", existing["monitoring_frequency"])
            previous_schedule = None
            previous_config = None
            if automatic_monitoring is not None or "monitoring_frequency" in updates:
                previous_schedule = conn.execute(
                    """
                    SELECT details FROM audit_events
                    WHERE source_id = ? AND event_type = 'CRAWL_SCHEDULE_CONFIGURED'
                    ORDER BY id DESC LIMIT 1
                    """,
                    (source_id,),
                ).fetchone()
                previous_config = (
                    json.loads(previous_schedule["details"]) if previous_schedule else None
                )
            schedule_enabled = (
                bool(automatic_monitoring)
                if automatic_monitoring is not None
                else bool(previous_config and previous_config.get("enabled"))
            )
            if schedule_enabled and _next_monitoring_run(datetime.utcnow(), frequency) is None:
                raise HTTPException(
                    status_code=422,
                    detail="Automatic monitoring requires Every alternate day, Weekly, or Monthly frequency.",
                )
            if updates:
                cursor = conn.execute(
                    f"UPDATE sources SET {assignments} WHERE id = ?",
                    values,
                )
                if cursor.rowcount == 0:
                    raise HTTPException(status_code=404, detail="Source not found.")
            if updates:
                active_value = updates.get("active", existing["active"])
                if "active" in updates and bool(active_value) != bool(existing["active"]):
                    event_type = "Source Activated" if active_value else "Source Deactivated"
                    details = f"{existing['reference_no'] or existing['name']} – Set {'to active' if active_value else 'to inactive'}"
                else:
                    event_type = "Source Updated"
                    changed_fields = ", ".join(
                        field.replace("_", " ").title()
                        for field in updates
                        if field != "updated_at"
                    )
                    details = f"{existing['reference_no'] or existing['name']} – Updated {changed_fields}"
                _record_audit_event(
                    conn,
                    event_type,
                    details,
                    entity_type="source",
                    entity_id=source_id,
                    source_id=source_id,
                    reference_id=updates.get("reference_no", existing["reference_no"]),
                )
            update_schedule = automatic_monitoring is not None or (
                previous_config is not None
                and "monitoring_frequency" in updates
                and previous_config.get("frequency") != frequency
            )
            if update_schedule:
                if automatic_monitoring is not None:
                    schedule_enabled = bool(automatic_monitoring)
                if (
                    previous_config is None
                    or bool(previous_config.get("enabled")) != schedule_enabled
                    or previous_config.get("frequency") != frequency
                ):
                    schedule_details = json.dumps({
                        "enabled": schedule_enabled,
                        "frequency": frequency,
                        "source_name": existing["name"],
                    })
                    _record_audit_event(
                        conn,
                        "CRAWL_SCHEDULE_CONFIGURED",
                        schedule_details,
                        entity_type="source",
                        entity_id=source_id,
                        source_id=source_id,
                        reference_id=updates.get("reference_no", existing["reference_no"]),
                        action="ENABLED" if schedule_enabled else "DISABLED",
                        result=(
                            f"Automatic monitoring {'enabled' if schedule_enabled else 'disabled'}"
                            f" for {frequency} frequency."
                        ),
                    )
    except HTTPException:
        raise
    except Exception as exc:
        logger.exception("Failed to update source")
        raise HTTPException(status_code=400, detail=str(exc))

    message = "Source updated successfully."
    if automatic_monitoring is True:
        if not source_active_after:
            message = "Source settings saved. Automatic monitoring is configured but paused because the source is inactive."
        elif _crawl_scheduler_thread and _crawl_scheduler_thread.is_alive():
            message = "Source settings saved. Automatic monitoring is enabled using the configured frequency."
        else:
            message = "Source settings saved. Automatic monitoring is configured but is not active because the scheduler is not running."
    elif automatic_monitoring is False:
        message = "Source settings saved. Automatic monitoring is disabled; no scheduled runs will occur."
    return {
        "success": True,
        "id": source_id,
        "message": message,
    }


@app.post("/api/sources/test")
def test_source_endpoint(request: SourceTestRequest):
    try:
        result = test_source(
            str(request.url),
            content_types=request.content_types,
        )

        return {
            "success": str(result.get("status", "")).upper() == "SUCCESS",
            **result,
        }

    except Exception as exc:
        logger.exception("Source test failed")
        return {
            "success": False,
            "status": "error",
            "message": str(exc),
        }


@app.get("/api/product-classes")
def list_product_classes():
    with get_db() as conn:
        classes = conn.execute(
            "SELECT * FROM product_classes ORDER BY id"
        ).fetchall()

        rows = []
        for item in classes:
            products = conn.execute(
                "SELECT * FROM products WHERE product_class_id = ? ORDER BY id",
                (item["id"],),
            ).fetchall()

            rows.append({
                "id": item["id"],
                "reference_no": item["reference_no"],
                "class_name": item["class_name"],
                "active": bool(item["active"]),
                "created_at": item["created_at"],
                "updated_at": item["updated_at"],
                "products": [
                    {
                        "id": product["id"],
                        "product_class_id": product["product_class_id"],
                        "reference_no": product["reference_no"],
                        "product_name": product["product_name"],
                        "active": bool(product["active"]),
                        "created_at": product["created_at"],
                        "updated_at": product["updated_at"],
                    }
                    for product in products
                ],
            })

    return rows


_PRODUCT_CLASS_ANALYSIS_CATEGORIES = (
    "Overview",
    "Safety Signals",
    "Adverse Events",
    "Class Effects & Mechanism",
    "Regulatory Actions",
    "Product Information",
    "Interactions & Special Populations",
    "Key Recommendations",
    "Source Documents",
)


def _product_class_finding_category(finding: dict) -> str:
    category_text = " ".join(
        str(finding.get(field) or "")
        for field in ("section", "trigger_name", "safety_topic")
    ).casefold()
    if any(term in category_text for term in ("regulatory", "label change", "warning", "regulator")):
        return "Regulatory Actions"
    if any(term in category_text for term in ("adverse event", "adverse reaction", "side effect", "reaction")):
        return "Adverse Events"
    if any(term in category_text for term in ("interaction", "pregnan", "elderly", "paediatric", "pediatric", "special population")):
        return "Interactions & Special Populations"
    if any(term in category_text for term in ("mechanism", "class effect", "pharmacolog")):
        return "Class Effects & Mechanism"
    if any(term in category_text for term in ("product information", "sm pc", "smpc", "label", "package leaflet")):
        return "Product Information"
    if any(term in category_text for term in ("recommendation", "recommended", "recommend")):
        return "Key Recommendations"
    if any(term in category_text for term in ("signal", "safety", "trigger")):
        return "Safety Signals"
    return "Overview"


def _product_class_findings(conn, class_id: int, substance: Optional[str] = None):
    parameters: list = [class_id]
    substance_clause = ""
    if substance:
        substance_clause = """
            AND (
                lower(trim(COALESCE(f.product_name, ''))) = lower(trim(?))
                OR lower(trim(COALESCE(f.active_substance, ''))) = lower(trim(?))
            )
        """
        parameters.extend((substance, substance))
    rows = conn.execute(
        f"""
        SELECT
            f.id, f.document_id, f.product_name, f.active_substance,
            f.trigger_name, f.safety_topic, f.relevant_terms, f.evidence,
            f.page_number, f.comparison_status, f.created_at,
            d.title AS document_title, d.url AS document_url,
            d.publication_date, d.reference_id, d.content_type,
            s.id AS source_id, s.name AS source_name, s.url AS source_url,
            (
                SELECT sm.summary
                FROM summaries sm
                WHERE sm.finding_id = f.id
                ORDER BY sm.created_at DESC, sm.id DESC
                LIMIT 1
            ) AS finding_summary
        FROM findings f
        JOIN documents d ON d.id = f.document_id
        JOIN sources s ON s.id = d.source_id
        WHERE EXISTS (
            SELECT 1
            FROM products p
            WHERE p.product_class_id = ?
              AND lower(trim(p.product_name)) IN (
                  lower(trim(COALESCE(f.product_name, ''))),
                  lower(trim(COALESCE(f.active_substance, '')))
              )
        )
        {substance_clause}
        ORDER BY
            CASE WHEN date(d.publication_date) IS NULL THEN 1 ELSE 0 END,
            date(d.publication_date) DESC,
            f.id DESC
        LIMIT 2000
        """,
        parameters,
    ).fetchall()
    findings = [dict(row) for row in rows]
    for finding in findings:
        finding["category"] = _product_class_finding_category(finding)
    return findings


@app.get("/api/product-classes/{class_id}/analysis")
def get_product_class_analysis(
    class_id: int,
    substance: Optional[str] = None,
    provider: Optional[str] = None,
):
    with get_db() as conn:
        product_class = conn.execute(
            "SELECT id, reference_no, class_name, active FROM product_classes WHERE id = ?",
            (class_id,),
        ).fetchone()
        if not product_class:
            raise HTTPException(status_code=404, detail="Product class not found.")
        products = conn.execute(
            """
            SELECT id, product_name, active
            FROM products
            WHERE product_class_id = ?
            ORDER BY product_name COLLATE NOCASE, id
            """,
            (class_id,),
        ).fetchall()
        findings = _product_class_findings(conn, class_id, substance)
        parameters = [class_id, substance or ""]
        analysis_filter = "product_class_id = ? AND substance = ?"
        if provider:
            analysis_filter += " AND provider = ?"
            parameters.append(provider.strip().lower())
        analysis = conn.execute(
            f"""
            SELECT provider, substance, summary, status, error_message,
                   created_at, updated_at
            FROM product_class_ai_analyses
            WHERE {analysis_filter}
            ORDER BY updated_at DESC, id DESC
            LIMIT 1
            """,
            parameters,
        ).fetchone()

    return {
        "product_class": dict(product_class),
        "products": [dict(row) for row in products],
        "selected_substance": substance or "",
        "categories": list(_PRODUCT_CLASS_ANALYSIS_CATEGORIES),
        "findings": findings,
        "analysis": dict(analysis) if analysis else None,
    }


@app.post("/api/product-classes/{class_id}/analysis")
def run_product_class_analysis(
    class_id: int,
    provider: str = Query(...),
    substance: Optional[str] = None,
):
    provider_id = provider.strip().lower()
    if provider_id not in {"ollama", "openai"}:
        raise HTTPException(status_code=400, detail="Unsupported AI provider.")
    provider_status = next(
        (
            item for item in list_ai_providers()["providers"]
            if item["id"] == provider_id
        ),
        None,
    )
    if not provider_status or not provider_status.get("available"):
        raise HTTPException(
            status_code=503,
            detail=(
                provider_status.get("message")
                if provider_status
                else "The selected AI provider is not configured."
            ),
        )

    with get_db() as conn:
        product_class = conn.execute(
            "SELECT id, class_name FROM product_classes WHERE id = ?",
            (class_id,),
        ).fetchone()
        if not product_class:
            raise HTTPException(status_code=404, detail="Product class not found.")
        findings = _product_class_findings(conn, class_id, substance)
        if not findings:
            raise HTTPException(
                status_code=422,
                detail="No stored findings match this product class and substance.",
            )
        now = datetime.utcnow().isoformat()
        conn.execute(
            """
            INSERT INTO product_class_ai_analyses
                (product_class_id, provider, substance, summary, status,
                 error_message, created_at, updated_at)
            VALUES (?, ?, ?, NULL, 'GENERATING', NULL, ?, ?)
            ON CONFLICT(product_class_id, provider, substance)
            DO UPDATE SET summary = NULL, status = 'GENERATING',
                          error_message = NULL, created_at = excluded.created_at,
                          updated_at = excluded.updated_at
            """,
            (class_id, provider_id, substance or "", now, now),
        )

    evidence = "\n\n".join(
        (
            f"Finding category: {finding['category']}\n"
            f"Product: {finding.get('product_name') or 'Not specified'}\n"
            f"Active substance: {finding.get('active_substance') or 'Not specified'}\n"
            f"Safety topic: {finding.get('safety_topic') or finding.get('trigger_name') or 'Not specified'}\n"
            f"Source document: {finding.get('document_title') or 'Not specified'} "
            f"({finding.get('publication_date') or 'Not specified'})\n"
            f"Evidence: {finding.get('evidence') or 'Not specified'}"
        )
        for finding in findings
    )
    generator = (
        generate_openai_document_summary
        if provider_id == "openai"
        else generate_document_summary
    )
    try:
        result = generator(
            f"{product_class['class_name']} drug-class findings",
            evidence,
            "Summarize only the stored findings and quotations provided. "
            "Do not add external or inferred claims.",
        )
    except Exception as exc:
        logger.exception("Product-class AI analysis failed for class %s", class_id)
        result = {"status": "ERROR", "error": "The configured AI provider failed."}

    completed_at = datetime.utcnow().isoformat()
    success = result.get("status") == "SUCCESS" and str(result.get("summary") or "").strip()
    status = "COMPLETED" if success else (
        "UNAVAILABLE" if result.get("status") in {"UNAVAILABLE", "MODEL_NOT_INSTALLED"}
        else "FAILED"
    )
    message = str(result.get("error") or "") if not success else None
    with get_db() as conn:
        conn.execute(
            """
            UPDATE product_class_ai_analyses
            SET summary = ?, status = ?, error_message = ?, updated_at = ?
            WHERE product_class_id = ? AND provider = ? AND substance = ?
            """,
            (
                str(result["summary"]).strip() if success else None,
                status,
                message,
                completed_at,
                class_id,
                provider_id,
                substance or "",
            ),
        )
        _record_audit_event(
            conn,
            "Product Class AI Analysis",
            f"AI analysis for {product_class['class_name']}"
            + (f" — {substance}" if substance else ""),
            entity_type="product_class",
            entity_id=class_id,
            action="Product Class AI Analysis",
            ai_model=result.get("model") or provider_id,
            result="Analysis completed" if success else message,
            status=status,
        )
    if not success:
        raise HTTPException(
            status_code=503 if status == "UNAVAILABLE" else 502,
            detail=message or "Product-class AI analysis failed.",
        )
    return {
        "success": True,
        "product_class_id": class_id,
        "provider": provider_id,
        "substance": substance or "",
        "summary": str(result["summary"]).strip(),
        "status": status,
        "updated_at": completed_at,
    }


@app.post("/api/product-classes")
def create_product_class(payload: ProductClassCreate):
    reference_no = payload.reference_no or f"PROD_CLASS_{payload.class_name.upper().replace(' ', '_')[:24]}_{datetime.utcnow().strftime('%Y%m%d%H%M%S')}"
    now = datetime.utcnow().isoformat()

    with get_db() as conn:
        existing = conn.execute(
            "SELECT id FROM product_classes WHERE reference_no = ? LIMIT 1",
            (reference_no,),
        ).fetchone()
        if existing:
            raise HTTPException(status_code=400, detail="A product class with this reference already exists.")

        cursor = conn.execute(
            """
            INSERT INTO product_classes
            (reference_no, class_name, active, created_at, updated_at)
            VALUES (?, ?, ?, ?, ?)
            """,
            (reference_no, payload.class_name, int(bool(payload.active)), now, now),
        )
        class_id = cursor.lastrowid
        _record_audit_event(
            conn,
            "Product Class Added",
            f"Added {payload.class_name}",
            entity_type="product_class",
            entity_id=class_id,
            reference_id=reference_no,
        )

    return {"success": True, "id": class_id, "reference_no": reference_no, "message": "Product class created successfully."}


@app.patch("/api/product-classes/{class_id}")
def update_product_class(class_id: int, payload: ProductClassUpdate):
    updates = payload.dict(exclude_unset=True)
    if not updates:
        raise HTTPException(status_code=400, detail="No product class fields supplied.")

    if "active" in updates:
        updates["active"] = int(bool(updates["active"]))

    updates["updated_at"] = datetime.utcnow().isoformat()
    assignments = ", ".join(f"{field} = ?" for field in updates)
    values = list(updates.values()) + [class_id]

    with get_db() as conn:
        existing = conn.execute(
            "SELECT reference_no, class_name, active FROM product_classes WHERE id = ?",
            (class_id,),
        ).fetchone()
        if not existing:
            raise HTTPException(status_code=404, detail="Product class not found.")
        cursor = conn.execute(
            f"UPDATE product_classes SET {assignments} WHERE id = ?",
            values,
        )
        if cursor.rowcount == 0:
            raise HTTPException(status_code=404, detail="Product class not found.")
        if "active" in updates and bool(updates["active"]) != bool(existing["active"]):
            event_type = "Product Class Activated" if updates["active"] else "Product Class Deactivated"
            details = f"{existing['class_name']} – Set {'to active' if updates['active'] else 'to inactive'}"
        else:
            event_type = "Product Class Updated"
            details = f"Updated {existing['class_name']}"
        _record_audit_event(
            conn,
            event_type,
            details,
            entity_type="product_class",
            entity_id=class_id,
            reference_id=existing["reference_no"],
        )

    return {"success": True, "id": class_id, "message": "Product class updated successfully."}


@app.post("/api/products")
def create_product(payload: ProductCreate):
    now = datetime.utcnow().isoformat()
    reference_no = payload.reference_no or f"PROD_{payload.product_class_id}_{payload.product_name.upper().replace(' ', '_')[:20]}_{datetime.utcnow().strftime('%Y%m%d%H%M%S')}"

    with get_db() as conn:
        existing = conn.execute(
            "SELECT id FROM products WHERE reference_no = ? LIMIT 1",
            (reference_no,),
        ).fetchone()
        if existing:
            raise HTTPException(status_code=400, detail="A product with this reference already exists.")

        class_exists = conn.execute(
            "SELECT id FROM product_classes WHERE id = ? LIMIT 1",
            (payload.product_class_id,),
        ).fetchone()
        if not class_exists:
            raise HTTPException(status_code=404, detail="Product class not found.")

        cursor = conn.execute(
            """
            INSERT INTO products
            (product_class_id, reference_no, product_name, active, created_at, updated_at)
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            (
                payload.product_class_id,
                reference_no,
                payload.product_name,
                int(bool(payload.active)),
                now,
                now,
            ),
        )
        product_id = cursor.lastrowid
        product_class = conn.execute(
            "SELECT class_name FROM product_classes WHERE id = ?",
            (payload.product_class_id,),
        ).fetchone()
        _record_audit_event(
            conn,
            "Product Added",
            f"Added {payload.product_name} to {product_class['class_name']}",
            entity_type="product",
            entity_id=product_id,
            reference_id=reference_no,
        )

    return {"success": True, "id": product_id, "reference_no": reference_no, "message": "Product created successfully."}


@app.patch("/api/products/{product_id}")
def update_product(product_id: int, payload: ProductUpdate):
    updates = payload.dict(exclude_unset=True)
    if not updates:
        raise HTTPException(status_code=400, detail="No product fields supplied.")

    if "active" in updates:
        updates["active"] = int(bool(updates["active"]))

    updates["updated_at"] = datetime.utcnow().isoformat()
    assignments = ", ".join(f"{field} = ?" for field in updates)
    values = list(updates.values()) + [product_id]

    with get_db() as conn:
        existing = conn.execute(
            """
            SELECT p.reference_no, p.product_name, p.active, pc.class_name
            FROM products p
            JOIN product_classes pc ON pc.id = p.product_class_id
            WHERE p.id = ?
            """,
            (product_id,),
        ).fetchone()
        if not existing:
            raise HTTPException(status_code=404, detail="Product not found.")
        cursor = conn.execute(
            f"UPDATE products SET {assignments} WHERE id = ?",
            values,
        )
        if cursor.rowcount == 0:
            raise HTTPException(status_code=404, detail="Product not found.")
        updated_name = updates.get("product_name", existing["product_name"])
        _record_audit_event(
            conn,
            "Product Updated",
            f"Updated {existing['product_name']} to {updated_name} in {existing['class_name']}",
            entity_type="product",
            entity_id=product_id,
            reference_id=existing["reference_no"],
        )

    return {"success": True, "id": product_id, "message": "Product updated successfully."}


@app.delete("/api/product-classes/{class_id}")
def delete_product_class(class_id: int):
    with get_db() as conn:
        product_class = conn.execute(
            "SELECT reference_no, class_name FROM product_classes WHERE id = ?",
            (class_id,),
        ).fetchone()
        if not product_class:
            raise HTTPException(status_code=404, detail="Product class not found.")
        cursor = conn.execute(
            "UPDATE product_classes SET active = 0, updated_at = ? WHERE id = ?",
            (datetime.utcnow().isoformat(), class_id),
        )
        if cursor.rowcount == 0:
            raise HTTPException(status_code=404, detail="Product class not found.")
        _record_audit_event(
            conn,
            "Product Class Deactivated",
            f"Deactivated {product_class['class_name']}",
            entity_type="product_class",
            entity_id=class_id,
            reference_id=product_class["reference_no"],
        )

    return {"success": True, "id": class_id, "message": "Product class deactivated."}


@app.delete("/api/products/{product_id}")
def delete_product(product_id: int):
    with get_db() as conn:
        product = conn.execute(
            """
            SELECT p.reference_no, p.product_name, pc.class_name
            FROM products p
            JOIN product_classes pc ON pc.id = p.product_class_id
            WHERE p.id = ?
            """,
            (product_id,),
        ).fetchone()
        if not product:
            raise HTTPException(status_code=404, detail="Product not found.")
        cursor = conn.execute(
            "UPDATE products SET active = 0, updated_at = ? WHERE id = ?",
            (datetime.utcnow().isoformat(), product_id),
        )
        if cursor.rowcount == 0:
            raise HTTPException(status_code=404, detail="Product not found.")
        _record_audit_event(
            conn,
            "Product Removed",
            f"Removed {product['product_name']} from {product['class_name']}",
            entity_type="product",
            entity_id=product_id,
            reference_id=product["reference_no"],
        )

    return {"success": True, "id": product_id, "message": "Product deactivated."}


# ---------------------------------------------------------
# Crawl & Monitor
# ---------------------------------------------------------

@app.get("/api/ai/providers")
def list_ai_providers():
    try:
        connection = test_ollama_connection()
    except Exception as exc:
        logger.exception("Unable to check configured AI providers")
        connection = {
            "available": False,
            "model_available": False,
            "message": str(exc),
        }

    openai_status = openai_provider_status()
    return {
        "providers": [
            {
                "id": "ollama",
                "name": "Ollama (Local)",
                "configured": bool(os.getenv("OLLAMA_BASE_URL")),
                "available": bool(
                    connection.get("available")
                    and connection.get("model_available")
                ),
                "model": connection.get("model"),
                "message": connection.get("message"),
            },
            {
                "id": "openai",
                "name": "OpenAI",
                **openai_status,
            },
        ]
    }


@app.get("/api/documents/{document_id}/ai-analysis")
def get_document_ai_analysis(
    document_id: int,
    provider: Optional[str] = None,
):
    with get_db() as conn:
        document = conn.execute(
            "SELECT id FROM documents WHERE id = ?",
            (document_id,),
        ).fetchone()
        if not document:
            raise HTTPException(status_code=404, detail="Document not found.")

        if provider:
            row = conn.execute(
                """
                SELECT *
                FROM document_ai_analyses
                WHERE document_id = ? AND provider = ? AND analysis_type = 'summary'
                """,
                (document_id, provider.strip().lower()),
            ).fetchone()
        else:
            row = conn.execute(
                """
                SELECT *
                FROM document_ai_analyses
                WHERE document_id = ? AND analysis_type = 'summary'
                  AND provider IN ('ollama', 'openai')
                ORDER BY updated_at DESC, id DESC
                LIMIT 1
                """,
                (document_id,),
            ).fetchone()

    return dict(row) if row else {"document_id": document_id, "status": "NOT_GENERATED"}


@app.post("/api/documents/{document_id}/ai-analysis")
def generate_document_ai_analysis(
    document_id: int,
    request: DocumentAIAnalysisRequest,
):
    provider = request.provider.strip().lower()
    analysis_type = request.analysis_type.strip().lower()
    if provider not in {"ollama", "openai"}:
        raise HTTPException(status_code=400, detail="AI provider unavailable.")
    if analysis_type != "summary":
        raise HTTPException(status_code=400, detail="Unsupported analysis type.")

    with get_db() as conn:
        document = conn.execute(
            """
            SELECT d.id, d.title, s.instruction AS source_instruction
            FROM documents d
            JOIN sources s ON s.id = d.source_id
            WHERE d.id = ?
            """,
            (document_id,),
        ).fetchone()
        if not document:
            raise HTTPException(status_code=404, detail="Document not found.")

        pages = conn.execute(
            """
            SELECT page_text
            FROM document_pages
            WHERE document_id = ?
            ORDER BY page_number
            """,
            (document_id,),
        ).fetchall()

        content = "\n\n".join(
            str(page["page_text"] or "").strip()
            for page in pages
            if str(page["page_text"] or "").strip()
        )
        if not content:
            _record_audit_event(
                conn,
                "AI Analysis (Summary)",
                "Analysis failed because the document has no extracted text.",
                entity_type="document",
                entity_id=document_id,
                reference_id=None,
                action="AI Analysis (Summary)",
                ai_model=(
                    os.getenv("OPENAI_MODEL", "gpt-4o-mini")
                    if provider == "openai"
                    else os.getenv("OLLAMA_MODEL", provider)
                ),
                result="No extracted text available",
                status="FAILED",
            )
            now = datetime.utcnow().isoformat()
            conn.execute(
                """
                INSERT INTO document_ai_analyses
                    (document_id, provider, analysis_type, summary, status,
                     error_message, created_at, updated_at)
                VALUES (?, ?, ?, NULL, 'FAILED', ?, ?, ?)
                ON CONFLICT(document_id, provider, analysis_type)
                DO UPDATE SET summary = NULL, status = 'FAILED',
                              error_message = excluded.error_message,
                              created_at = excluded.created_at,
                              updated_at = excluded.updated_at
                """,
                (
                    document_id,
                    provider,
                    analysis_type,
                    "This document has no extracted text available for analysis.",
                    now,
                    now,
                ),
            )
        else:
            now = datetime.utcnow().isoformat()
            conn.execute(
                """
                INSERT INTO document_ai_analyses
                    (document_id, provider, analysis_type, summary, status,
                     error_message, created_at, updated_at)
                VALUES (?, ?, ?, NULL, 'GENERATING', NULL, ?, ?)
                ON CONFLICT(document_id, provider, analysis_type)
                DO UPDATE SET summary = NULL, status = 'GENERATING',
                              error_message = NULL,
                              created_at = excluded.created_at,
                              updated_at = excluded.updated_at
                """,
                (document_id, provider, analysis_type, now, now),
            )

    if not content:
        raise HTTPException(
            status_code=422,
            detail="This document has no extracted text available for analysis.",
        )

    try:
        summary_generator = (
            generate_openai_document_summary
            if provider == "openai"
            else generate_document_summary
        )
        result = summary_generator(
            str(document["title"] or "Untitled document"),
            content,
            str(document["source_instruction"] or "").strip() or None,
        )
    except Exception as exc:
        logger.exception("AI analysis failed for document %s", document_id)
        message = "AI analysis failed."
        with get_db() as conn:
            conn.execute(
                """
                UPDATE document_ai_analyses
                SET status = 'FAILED', error_message = ?, updated_at = ?
                WHERE document_id = ? AND provider = ? AND analysis_type = ?
                """,
                (message, datetime.utcnow().isoformat(), document_id, provider, analysis_type),
            )
            document_ref = conn.execute(
                "SELECT reference_id FROM documents WHERE id = ?",
                (document_id,),
            ).fetchone()
            _record_audit_event(
                conn,
                "AI Analysis (Summary)",
                f"AI analysis failed for document {document_id}",
                entity_type="document",
                entity_id=document_id,
                reference_id=document_ref["reference_id"] if document_ref else None,
                action="AI Analysis (Summary)",
                ai_model=provider,
                result=message,
                status="FAILED",
            )
        raise HTTPException(status_code=502, detail=message) from exc

    if result.get("status") != "SUCCESS":
        message = str(result.get("error") or "AI provider unavailable.")
        status = (
            "UNAVAILABLE"
            if result.get("status") in {"UNAVAILABLE", "MODEL_NOT_INSTALLED"}
            else "FAILED"
        )
        with get_db() as conn:
            conn.execute(
                """
                UPDATE document_ai_analyses
                SET status = ?, error_message = ?, updated_at = ?
                WHERE document_id = ? AND provider = ? AND analysis_type = ?
                """,
                (status, message, datetime.utcnow().isoformat(), document_id, provider, analysis_type),
            )
            document_ref = conn.execute(
                "SELECT reference_id FROM documents WHERE id = ?",
                (document_id,),
            ).fetchone()
            _record_audit_event(
                conn,
                "AI Analysis (Summary)",
                f"AI analysis unavailable for document {document_id}",
                entity_type="document",
                entity_id=document_id,
                reference_id=document_ref["reference_id"] if document_ref else None,
                action="AI Analysis (Summary)",
                ai_model=result.get("model") or provider,
                result=message,
                status=status,
            )
        raise HTTPException(
            status_code=503 if status == "UNAVAILABLE" else 502,
            detail=message,
        )

    generated_at = datetime.utcnow().isoformat()
    with get_db() as conn:
        conn.execute(
            """
            UPDATE document_ai_analyses
            SET summary = ?, status = 'COMPLETED', error_message = NULL,
                updated_at = ?
            WHERE document_id = ? AND provider = ? AND analysis_type = ?
            """,
            (result["summary"], generated_at, document_id, provider, analysis_type),
        )
        row = conn.execute(
            """
            SELECT *
            FROM document_ai_analyses
            WHERE document_id = ? AND provider = ? AND analysis_type = ?
            """,
            (document_id, provider, analysis_type),
        ).fetchone()
        document_ref = conn.execute(
            "SELECT reference_id FROM documents WHERE id = ?",
            (document_id,),
        ).fetchone()
        _record_audit_event(
            conn,
            "AI Analysis (Summary)",
            f"AI summary generated for document {document_id}",
            entity_type="document",
            entity_id=document_id,
            reference_id=document_ref["reference_id"] if document_ref else None,
            action="AI Analysis (Summary)",
            ai_model=result.get("model") or provider,
            result="Summary generated",
            status="COMPLETED",
        )

    return dict(row)


@app.get("/api/documents/{document_id}/ai-analysis/download")
def download_document_ai_analysis(document_id: int, provider: str):
    analysis = get_document_ai_analysis(document_id, provider)
    if analysis.get("status") != "COMPLETED" or not analysis.get("summary"):
        raise HTTPException(status_code=404, detail="Completed AI summary not found.")

    with get_db() as conn:
        document = conn.execute(
            """
            SELECT title, publication_date, reference_id
            FROM documents
            WHERE id = ?
            """,
            (document_id,),
        ).fetchone()
    if not document:
        raise HTTPException(status_code=404, detail="Document not found.")

    content = (
        f"Document: {document['title'] or 'Untitled document'}\n"
        f"Reference: {document['reference_id'] or 'Not available'}\n"
        f"Published: {document['publication_date'] or 'Not available'}\n"
        f"AI provider: {analysis['provider']}\n"
        f"Generated: {analysis['updated_at']}\n\n"
        f"Summary\n=======\n{analysis['summary']}\n"
    )
    filename = f"document-{document_id}-ai-summary.txt"
    return Response(
        content=content,
        media_type="text/plain; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _published_date_token(value: Optional[str]) -> str:
    if not value:
        return "UNDATED"

    cleaned = " ".join(str(value).split())
    for date_format in (
        "%Y-%m-%d",
        "%d %B %Y",
        "%d %b %Y",
        "%d-%B-%Y",
        "%d-%b-%Y",
    ):
        try:
            return datetime.strptime(cleaned, date_format).strftime("%Y%m%d")
        except ValueError:
            continue

    match = re.search(r"\b(\d{1,2})[\s-]+([A-Za-z]{3,9})[\s-]+(\d{4})\b", cleaned)
    if match:
        for date_format in ("%d %B %Y", "%d %b %Y"):
            try:
                parsed = datetime.strptime(" ".join(match.groups()), date_format)
                return parsed.strftime("%Y%m%d")
            except ValueError:
                continue

    return "UNDATED"


def _document_reference_id(source_reference: Optional[str], published_date: Optional[str], document_id: int) -> str:
    prefix = re.sub(
        r"[^A-Za-z0-9_-]+",
        "_",
        str(source_reference or "DOCUMENT").strip(),
    ).strip("_")
    return f"{prefix}_{_published_date_token(published_date)}_{document_id:03d}"


def _record_discovered_documents(
    source: dict,
    documents: list[dict],
) -> tuple[int, str]:
    crawled_at = datetime.utcnow().isoformat()
    source_reference = source.get("reference_no") or str(source["id"])

    with get_db() as conn:
        for document in documents:
            url = str(document.get("document_key") or document.get("url") or "").strip()
            parsed_url = urlparse(url)
            if parsed_url.scheme not in {"http", "https"} or not parsed_url.netloc:
                logger.warning("Skipping discovered document with invalid URL: %s", url)
                continue

            publication_date = document.get("publication_date")
            existing = conn.execute(
                """
                SELECT id, reference_id, publication_date
                FROM documents
                WHERE lower(rtrim(url, '/')) = lower(rtrim(?, '/'))
                LIMIT 1
                """,
                (url,),
            ).fetchone()

            is_new = existing is None
            if existing:
                document_id = int(existing["id"])
                conn.execute(
                    """
                    UPDATE documents
                    SET title = COALESCE(?, title),
                        publication_date = ?,
                        meeting_date = COALESCE(?, meeting_date),
                        filename = COALESCE(?, filename),
                        content_type = COALESCE(?, content_type),
                        last_crawled_at = ?
                    WHERE id = ?
                    """,
                    (
                        document.get("title"),
                        publication_date,
                        document.get("meeting_date"),
                        document.get("filename"),
                        document.get("content_type"),
                        crawled_at,
                        document_id,
                    ),
                )
                reference_id = existing["reference_id"]
                previous_generated_id = _document_reference_id(
                    source_reference,
                    existing["publication_date"],
                    document_id,
                )
                refreshed_reference_id = _document_reference_id(
                    source_reference,
                    publication_date,
                    document_id,
                )
                if not reference_id or reference_id == previous_generated_id:
                    reference_id = refreshed_reference_id
                    conn.execute(
                        "UPDATE documents SET reference_id = ? WHERE id = ?",
                        (reference_id, document_id),
                    )
            else:
                cursor = conn.execute(
                    """
                    INSERT OR IGNORE INTO documents
                    (
                        source_id,
                        title,
                        url,
                        meeting_date,
                        publication_date,
                        filename,
                        discovered_date,
                        last_crawled_at,
                        content_type,
                        status,
                        created_at
                    )
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 'NEW', ?)
                    """,
                    (
                        int(source["id"]),
                        document.get("title") or document.get("filename") or url,
                        url,
                        document.get("meeting_date"),
                        publication_date,
                        document.get("filename"),
                        crawled_at,
                        crawled_at,
                        document.get("content_type"),
                        crawled_at,
                    ),
                )
                if cursor.rowcount:
                    document_id = int(cursor.lastrowid)
                    reference_id = None
                else:
                    existing = conn.execute(
                        """
                        SELECT id, reference_id, publication_date
                        FROM documents
                        WHERE lower(rtrim(url, '/')) = lower(rtrim(?, '/'))
                        LIMIT 1
                        """,
                        (url,),
                    ).fetchone()
                    if not existing:
                        raise RuntimeError(
                            f"Discovered document could not be stored: {url}"
                        )
                    document_id = int(existing["id"])
                    reference_id = existing["reference_id"]
                    is_new = False
                    conn.execute(
                        """
                        UPDATE documents
                        SET publication_date = ?,
                            content_type = COALESCE(?, content_type),
                            last_crawled_at = ?
                        WHERE id = ?
                        """,
                        (
                            publication_date,
                            document.get("content_type"),
                            crawled_at,
                            document_id,
                        ),
                    )

            if not reference_id:
                reference_id = _document_reference_id(
                    source_reference,
                    publication_date,
                    document_id,
                )
                conn.execute(
                    "UPDATE documents SET reference_id = ? WHERE id = ?",
                    (reference_id, document_id),
                )

            document.update(
                {
                    "id": document_id,
                    "reference_id": reference_id,
                    "publication_date": publication_date,
                    "date_crawled": crawled_at,
                    "status": "NEW" if is_new else "ALREADY_CRAWLED",
                    "new_document": is_new,
                }
            )

        conn.execute(
            """
            UPDATE sources
            SET last_crawled_at = ?,
                last_crawl_status = 'success',
                last_error_message = NULL,
                updated_at = ?
            WHERE id = ?
            """,
            (crawled_at, crawled_at, int(source["id"])),
        )

    return sum(1 for document in documents if document.get("new_document")), crawled_at


@app.post("/api/crawl/discover")
def discover_crawl_documents(request: CrawlDiscoverRequest):
    with get_db() as conn:
        source = conn.execute(
            """
            SELECT *
            FROM sources
            WHERE id = ?
              AND active = 1
            """,
            (request.source_id,),
        ).fetchone()

    if not source:
        raise HTTPException(
            status_code=404,
            detail="Active source not found.",
        )

    try:
        source_data = serialize_source_row(source)
        configured_content_types = normalize_content_types(
            source_data.get("content_types")
        )
        content_types = (
            _crawl_content_types(source_data, request.content_types)
            if request.content_types is not None or configured_content_types
            else None
        )
        if content_types is not None and "Other published documents" not in content_types:
            content_types.append("Other published documents")
        result = crawl_source(
            source_url=str(source["url"]),
            source_id=None,
            content_types=content_types,
            source_name=source_data.get("name"),
            reference_no=source_data.get("reference_no"),
        )
    except HTTPException:
        raise
    except Exception as exc:
        logger.exception(
            "Document discovery failed for source %s",
            request.source_id,
        )
        raise HTTPException(
            status_code=502,
            detail=f"Document discovery failed: {exc}",
        )

    if not result["success"]:
        crawl_error = result.get("error") or "Unknown crawl failure"
        raise HTTPException(
            status_code=502,
            detail=f"Document discovery failed: {crawl_error}",
        )

    documents = deduplicate_documents(list(result.get("documents", [])))
    documents = sorted(
        documents,
        key=lambda item: (
            _published_date_token(item.get("publication_date")) != "UNDATED",
            _published_date_token(item.get("publication_date")),
            str(item.get("title") or "").casefold(),
        ),
        reverse=True,
    )
    source_reference = source_data.get("reference_no") or str(request.source_id)
    discovered_at = datetime.utcnow().isoformat()
    new_documents = 0
    already_crawled = 0
    with get_db() as conn:
        stored_rows = conn.execute(
            """
            SELECT id, url, reference_id, publication_date, status, filename
            FROM documents
            WHERE source_id = ?
            """,
            (request.source_id,),
        ).fetchall()
        next_document_id = (
            conn.execute(
                "SELECT COALESCE(MAX(id), 0) + 1 FROM documents"
            ).fetchone()[0]
        )

    stored_by_url = {
        str(row["url"] or "").strip().rstrip("/").casefold(): dict(row)
        for row in stored_rows
    }
    for index, discovered in enumerate(documents):
        url = str(discovered.get("document_key") or discovered.get("url") or "").strip()
        stored = stored_by_url.get(url.rstrip("/").casefold())
        discovered["source_id"] = int(request.source_id)
        discovered["source_name"] = source_data.get("name")
        discovered["discovered_at"] = discovered_at
        if not stored:
            discovered["status"] = "NEW"
            discovered["new_document"] = True
            discovered["already_crawled"] = False
            discovered["reference_id"] = _document_reference_id(
                source_reference,
                discovered.get("publication_date"),
                int(next_document_id) + index,
            )
            new_documents += 1
            continue

        stored_status = str(stored.get("status") or "").upper()
        processed = stored_status in {"PROCESSED", "EXTRACTED"}
        failed = stored_status in {"FAILED", "EXTRACTION_FAILED"}
        discovered["id"] = int(stored["id"])
        discovered["reference_id"] = stored.get("reference_id") or _document_reference_id(
            source_reference,
            discovered.get("publication_date") or stored.get("publication_date"),
            int(stored["id"]),
        )
        discovered["filename"] = discovered.get("filename") or stored.get("filename")
        discovered["status"] = "ALREADY_CRAWLED" if processed else "FAILED" if failed else "NEW"
        discovered["new_document"] = False
        discovered["already_crawled"] = processed
        if processed:
            already_crawled += 1

    groups_by_type = {}
    for document in documents:
        group = str(document.get("content_type") or "Other published documents")
        groups_by_type[group] = groups_by_type.get(group, 0) + 1

    return {
        "status": "success",
        "source": source_data,
        "documents_found": len(documents),
        "new_documents": new_documents,
        "already_crawled": already_crawled,
        "date_discovered": discovered_at,
        "groups": [
            {"content_type": name, "count": count}
            for name, count in sorted(groups_by_type.items(), key=lambda item: item[0].casefold())
        ],
        "documents": documents,
    }


@app.post("/api/crawl/run")
def run_manual_crawl(request: CrawlDiscoverRequest):
    with get_db() as conn:
        source = conn.execute(
            """
            SELECT *
            FROM sources
            WHERE id = ?
              AND active = 1
            """,
            (request.source_id,),
        ).fetchone()

    if not source:
        raise HTTPException(
            status_code=404,
            detail="Active source not found.",
        )

    source_data = serialize_source_row(source)
    try:
        result = crawl_source(
            source_url=str(source["url"]),
            source_id=int(request.source_id),
            content_types=_crawl_content_types(
                source_data,
                request.content_types,
            ),
            source_name=source_data.get("name"),
            reference_no=source_data.get("reference_no"),
        )
    except HTTPException:
        raise
    except Exception as exc:
        logger.exception(
            "Manual crawl failed for source %s",
            request.source_id,
        )
        result = {
            "success": False,
            "error": str(exc),
            "documents_found": 0,
            "documents": [],
        }

    if result["success"]:
        result["documents"] = deduplicate_documents(list(result.get("documents", [])))
        result["documents_found"] = len(result["documents"])
        new_documents, _ = _record_discovered_documents(
            source_data,
            result.get("documents", []),
        )
        with get_db() as conn:
            completed_at = datetime.utcnow().isoformat()
            total_documents = conn.execute(
                "SELECT COUNT(*) AS count FROM documents WHERE source_id = ?",
                (request.source_id,),
            ).fetchone()["count"]
            conn.execute(
                """
                INSERT INTO crawl_history
                (source_id, started_at, completed_at, status, documents_found,
                 new_documents, total_documents, error_message)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    request.source_id,
                    completed_at,
                    completed_at,
                    "success",
                    result.get("documents_found", 0),
                    new_documents,
                    total_documents,
                    None,
                ),
            )

        return {
            "success": True,
            "source": source_data,
            "documents_found": result.get("documents_found", 0),
            "new_documents": new_documents,
            "documents": result.get("documents", []),
            "message": "Manual crawl completed successfully.",
        }

    logger.error(
        "Manual crawl failed for source %s: %s",
        request.source_id,
        result.get("error") or "Manual crawl failed",
    )
    with get_db() as conn:
        completed_at = datetime.utcnow().isoformat()
        conn.execute(
            """
            INSERT INTO crawl_history
            (source_id, started_at, completed_at, status, documents_found,
             new_documents, total_documents, error_message)
            VALUES (?, ?, ?, ?, ?, 0,
                (SELECT COUNT(*) FROM documents WHERE source_id = ?), ?)
            """,
            (
                request.source_id,
                completed_at,
                completed_at,
                "failed",
                0,
                request.source_id,
                result.get("error") or "Manual crawl failed",
            ),
        )
        conn.execute(
            """
            UPDATE sources
            SET last_crawled_at = ?, last_crawl_status = 'failed',
                last_error_message = ?, updated_at = ?
            WHERE id = ?
            """,
            (
                completed_at,
                result.get("error") or "Manual crawl failed",
                completed_at,
                request.source_id,
            ),
        )

    raise HTTPException(
        status_code=502,
        detail=result.get("error") or "Manual crawl failed.",
    )


selected_processing_jobs = {}


@app.post("/api/crawl/process")
def start_selected_document_processing(
    request: CrawlProcessRequest,
    background_tasks: BackgroundTasks,
):
    with get_db() as conn:
        source = conn.execute(
            """
            SELECT *
            FROM sources
            WHERE id = ?
              AND active = 1
            """,
            (request.source_id,),
        ).fetchone()

    if not source:
        raise HTTPException(
            status_code=404,
            detail="Active source not found.",
        )

    discovered_url = str(request.document_url).strip()

    try:
        parsed = urlparse(discovered_url)
    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail=f"Invalid document URL: {exc}",
        )

    if parsed.scheme not in {"http", "https"} or not parsed.netloc:
        raise HTTPException(
            status_code=400,
            detail="Document URL must be an HTTP or HTTPS URL.",
        )

    filename = (
        Path(parsed.path).name
    )

    if not parsed.path.lower().endswith(".pdf"):
        raise HTTPException(
            status_code=400,
            detail="Selected document URL must point to a supported PDF file.",
        )

    job_id = str(uuid.uuid4())
    selected_processing_jobs[job_id] = {
        "job_id": job_id,
        "status": "queued",
        "progress": 0,
        "stage": "Queued",
        "message": "Selected document processing queued.",
        "selected_document": {
            "title": request.title or filename,
            "meeting_date": request.meeting_date,
            "url": discovered_url,
            "filename": filename,
        },
        "pages_extracted": 0,
        "findings_detected": 0,
        "new_information_detected": 0,
        "topics_extracted": 0,
        "processing_status": None,
        "document_id": None,
        "error": None,
    }

    background_tasks.add_task(
        process_selected_document,
        job_id,
        request.source_id,
        discovered_url,
        request.title,
        filename,
        request.meeting_date,
        request.publication_date,
        request.content_type,
        request.reference_id,
    )

    return {
        "job_id": job_id,
        "status": "queued",
        "message": "Selected document processing started.",
    }


@app.get("/api/crawl/process/status/{job_id}")
def selected_document_processing_status(job_id: str):
    job = selected_processing_jobs.get(job_id)

    if not job:
        raise HTTPException(
            status_code=404,
            detail="Selected document processing job not found.",
        )

    return job


def process_selected_document(
    job_id: str,
    source_id: int,
    document_url: str,
    title: Optional[str],
    filename: str,
    meeting_date: Optional[str],
    publication_date: Optional[str],
    content_type: Optional[str],
    reference_id: Optional[str],
):
    job = selected_processing_jobs[job_id]
    crawl_started = datetime.utcnow().isoformat()
    document_id = None
    document_reference_id = None
    extraction_started = None
    new_document_created = False

    def update(**values):
        job.update(values)

    try:
        update(
            status="running",
            progress=5,
            stage="Validating selection",
            message="Validating the selected document.",
        )

        with get_db() as conn:
            existing = conn.execute(
                """
                SELECT *
                FROM documents
                WHERE url = ?
                """,
                (document_url,),
            ).fetchone()

            if existing and str(existing["status"]).upper() in {
                "PROCESSED",
                "EXTRACTED",
            }:
                page_count = conn.execute(
                    """
                    SELECT COUNT(*)
                    FROM document_pages
                    WHERE document_id = ?
                    """,
                    (existing["id"],),
                ).fetchone()[0]
                findings = conn.execute(
                    """
                    SELECT comparison_status
                    FROM findings
                    WHERE document_id = ?
                    """,
                    (existing["id"],),
                ).fetchall()
                existing_pages = conn.execute(
                    """
                    SELECT page_number, page_text
                    FROM document_pages
                    WHERE document_id = ?
                    ORDER BY page_number
                    """,
                    (existing["id"],),
                ).fetchall()
                # Reuse the stored extraction and record this successful crawl
                # request without duplicating pages or findings.
                reuse_started = datetime.utcnow().isoformat()
                reuse_completed = datetime.utcnow().isoformat()

                conn.execute(
                    """
                    INSERT INTO extraction_history
                    (
                        document_id,
                        started_at,
                        completed_at,
                        status,
                        pages_extracted,
                        error_message
                    )
                    VALUES (?, ?, ?, 'success', ?, NULL)
                    """,
                    (
                        existing["id"],
                        reuse_started,
                        reuse_completed,
                        page_count,
                    ),
                )

                existing_extraction = conn.execute(
                    """
                    SELECT completed_at
                    FROM extraction_history
                    WHERE document_id = ?
                    ORDER BY id DESC
                    LIMIT 1
                    """,
                    (existing["id"],),
                ).fetchone()
                conn.execute(
                    "UPDATE documents SET last_crawled_at = ? WHERE id = ?",
                    (reuse_completed, existing["id"]),
                )
                conn.execute(
                    """
                    INSERT INTO crawl_history
                    (
                        source_id, started_at, completed_at, status,
                        documents_found, new_documents, total_documents,
                        error_message
                    )
                    VALUES (
                        ?, ?, ?, 'success', 1, 0,
                        (SELECT COUNT(*) FROM documents WHERE source_id = ?),
                        NULL
                    )
                    """,
                    (source_id, reuse_started, reuse_completed, source_id),
                )
                conn.execute(
                    """
                    UPDATE sources
                    SET last_crawled_at = ?, last_crawl_status = 'success',
                        last_error_message = NULL, updated_at = ?
                    WHERE id = ?
                    """,
                    (reuse_completed, reuse_completed, source_id),
                )
                new_count = sum(
                    1
                    for finding in findings
                    if str(finding["comparison_status"] or "").upper()
                    in {"NEW", "NEW_INFORMATION"}
                )

                document_id = existing["id"]
                update(
                    status="completed",
                    progress=100,
                    stage="Completed",
                    message="Document was already crawled; no download or extraction was needed.",
                    selected_document={
                        "id": document_id,
                        "reference_id": existing["reference_id"],
                        "title": existing["title"],
                        "meeting_date": existing["meeting_date"],
                        "url": existing["url"],
                        "filename": existing["filename"],
                    },
                    pages_extracted=page_count,
                    findings_detected=len(findings),
                    new_information_detected=new_count,
                    processing_status="PROCESSED",
                    document_id=document_id,
                    topics_extracted=len(
                        extract_meeting_topics(
                            [dict(page) for page in existing_pages]
                        )
                    ),
                    last_crawled=reuse_completed,
                    last_processed=(
                        existing_extraction["completed_at"]
                        if existing_extraction
                        else None
                    ),
                    summary_status="PENDING",
                    already_processed=True,
                )
                return

            if existing:
                document_id = existing["id"]
                document_reference_id = existing["reference_id"]
                conn.execute(
                    """
                    UPDATE documents
                    SET title = COALESCE(?, title),
                        meeting_date = COALESCE(?, meeting_date),
                        publication_date = COALESCE(?, publication_date),
                        filename = COALESCE(?, filename),
                        content_type = COALESCE(?, content_type)
                    WHERE id = ?
                    """,
                    (
                        title,
                        meeting_date,
                        publication_date,
                        filename,
                        content_type,
                        document_id,
                    ),
                )
            else:
                new_document_created = True
                cursor = conn.execute(
                    """
                    INSERT INTO documents
                    (
                        source_id,
                        title,
                        url,
                        meeting_date,
                        publication_date,
                        filename,
                        discovered_date,
                        content_type,
                        status,
                        created_at
                    )
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, 'NEW', ?)
                    """,
                    (
                        source_id,
                        title or filename,
                        document_url,
                        meeting_date,
                        publication_date,
                        filename,
                        datetime.utcnow().isoformat(),
                        content_type,
                        datetime.utcnow().isoformat(),
                    ),
                )
                document_id = cursor.lastrowid
                source_row = conn.execute(
                    "SELECT reference_no FROM sources WHERE id = ?",
                    (source_id,),
                ).fetchone()
                reference_id = reference_id or _document_reference_id(
                    source_row["reference_no"] if source_row else None,
                    publication_date,
                    int(document_id),
                )
                document_reference_id = reference_id
                conn.execute(
                    "UPDATE documents SET reference_id = ? WHERE id = ?",
                    (reference_id, document_id),
                )

        update(
            progress=15,
            stage="Downloading PDF",
            message=f"Downloading {filename}.",
            document_id=document_id,
        )

        pdf_path = PDF_DIR / filename
        downloaded_path = download_pdf(
            document_url,
            str(pdf_path),
        )

        if not downloaded_path:
            raise RuntimeError("PDF download failed.")

        update(
            progress=35,
            stage="Extracting pages",
            message="Extracting text from the selected PDF.",
        )
        extraction_started = datetime.utcnow().isoformat()
        processed = process_pdf_safe(downloaded_path)

        if not processed.get("success"):
            raise RuntimeError(
                processed.get("error")
                or "PDF text extraction failed."
            )

        pages = processed.get("pages", [])
        sha256 = processed.get("sha256")

        with get_db() as conn:
            for page in pages:
                conn.execute(
                    """
                    INSERT OR REPLACE INTO document_pages
                    (document_id, page_number, page_text)
                    VALUES (?, ?, ?)
                    """,
                    (
                        document_id,
                        page["page_number"],
                        page.get("text", ""),
                    ),
                )

            conn.execute(
                """
                UPDATE documents
                SET filename = ?, sha256 = ?, status = 'EXTRACTED'
                WHERE id = ?
                """,
                (filename, sha256, document_id),
            )

            conn.execute(
                """
                INSERT INTO extraction_history
                (
                    document_id,
                    started_at,
                    completed_at,
                    status,
                    pages_extracted,
                    error_message
                )
                VALUES (?, ?, ?, 'success', ?, NULL)
                """,
                (
                    document_id,
                    extraction_started,
                    datetime.utcnow().isoformat(),
                    len(pages),
                ),
            )

        update(
            progress=65,
            stage="Detecting safety findings",
            message="Applying configured safety triggers.",
            pages_extracted=len(pages),
        )
        findings = remove_duplicate_findings(
            detect_document(pages, load_trigger_config())
        )
        topics = extract_meeting_topics(pages)

        with get_db() as conn:
            historical_rows = conn.execute(
                """
                SELECT
                    product_name,
                    active_substance,
                    safety_topic,
                    trigger_name,
                    evidence
                FROM findings
                WHERE document_id != ?
                """,
                (document_id,),
            ).fetchall()
            historical_findings = [dict(row) for row in historical_rows]

        update(
            progress=80,
            stage="Comparing historical findings",
            message="Checking findings for new information.",
        )

        compared_findings = [
            compare_finding(finding, historical_findings)
            for finding in findings
        ]
        new_count = sum(
            1
            for finding in compared_findings
            if finding.get("comparison_status") == "NEW_INFORMATION"
        )

        with get_db() as conn:
            for finding in compared_findings:
                conn.execute(
                    """
                    INSERT INTO findings
                    (
                        document_id,
                        page_number,
                        product_name,
                        active_substance,
                        safety_topic,
                        trigger_name,
                        evidence,
                        comparison_status,
                        created_at
                    )
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        document_id,
                        finding.get("page_number"),
                        finding.get("product_name"),
                        finding.get("active_substance"),
                        finding.get("safety_topic"),
                        finding.get("trigger_name"),
                        finding.get("evidence"),
                        finding.get(
                            "comparison_status",
                            "NEW_INFORMATION",
                        ),
                        datetime.utcnow().isoformat(),
                    ),
                )

            completed_at = datetime.utcnow().isoformat()
            conn.execute(
                """
                UPDATE documents
                SET status = 'PROCESSED', last_crawled_at = ?
                WHERE id = ?
                """,
                (completed_at, document_id),
            )
            conn.execute(
                """
                INSERT INTO crawl_history
                (
                    source_id,
                    started_at,
                    completed_at,
                    status,
                    documents_found,
                    new_documents,
                    total_documents,
                    error_message
                )
                VALUES (
                    ?, ?, ?, 'success', 1, ?,
                    (SELECT COUNT(*) FROM documents WHERE source_id = ?),
                    NULL
                )
                """,
                (
                    source_id,
                    crawl_started,
                    completed_at,
                    int(new_document_created),
                    source_id,
                ),
            )
            conn.execute(
                """
                UPDATE sources
                SET last_crawled_at = ?, last_crawl_status = 'success',
                    last_error_message = NULL, updated_at = ?
                WHERE id = ?
                """,
                (completed_at, completed_at, source_id),
            )

        update(
            status="completed",
            progress=100,
            stage="Completed",
            message="Selected document processed successfully.",
            selected_document={
                "id": document_id,
                "reference_id": document_reference_id,
                "title": title or filename,
                "meeting_date": meeting_date,
                "url": document_url,
                "filename": filename,
            },
            findings_detected=len(compared_findings),
            new_information_detected=new_count,
            topics_extracted=len(topics),
            processing_status="PROCESSED",
            document_id=document_id,
            already_processed=False,
        )

        with get_db() as conn:
            extraction = conn.execute(
                """
                SELECT started_at, completed_at
                FROM extraction_history
                WHERE document_id = ?
                ORDER BY id DESC
                LIMIT 1
                """,
                (document_id,),
            ).fetchone()
            crawl = conn.execute(
                """
                SELECT started_at, completed_at, status, documents_found
                FROM crawl_history
                WHERE source_id = ?
                ORDER BY id DESC
                LIMIT 1
                """,
                (source_id,),
            ).fetchone()
        job.update(
            last_crawled=crawl["completed_at"] if crawl else None,
            last_processed=(
                extraction["completed_at"]
                if extraction
                else None
            ),
            summary_status="PENDING",
        )

    except Exception as exc:
        logger.exception(
            "Selected document processing failed for job %s",
            job_id,
        )

        if document_id:
            with get_db() as conn:
                if extraction_started:
                    conn.execute(
                        """
                        INSERT INTO extraction_history
                        (
                            document_id,
                            started_at,
                            completed_at,
                            status,
                            pages_extracted,
                            error_message
                        )
                        VALUES (?, ?, ?, 'failed', 0, ?)
                        """,
                        (
                            document_id,
                            extraction_started,
                            datetime.utcnow().isoformat(),
                            str(exc),
                        ),
                    )
                conn.execute(
                    """
                    UPDATE documents
                    SET status = 'FAILED', last_crawled_at = ?
                    WHERE id = ?
                    """,
                    (datetime.utcnow().isoformat(), document_id),
                )
                failed_at = datetime.utcnow().isoformat()
                conn.execute(
                    """
                    INSERT INTO crawl_history
                    (
                        source_id, started_at, completed_at, status,
                        documents_found, new_documents, total_documents,
                        error_message
                    )
                    VALUES (
                        ?, ?, ?, 'failed', 1, 0,
                        (SELECT COUNT(*) FROM documents WHERE source_id = ?), ?
                    )
                    """,
                    (source_id, crawl_started, failed_at, source_id, str(exc)),
                )
                conn.execute(
                    """
                    UPDATE sources
                    SET last_crawled_at = ?, last_crawl_status = 'failed',
                        last_error_message = ?, updated_at = ?
                    WHERE id = ?
                    """,
                    (failed_at, str(exc), failed_at, source_id),
                )

        update(
            status="failed",
            progress=100,
            stage="Failed",
            message="Selected document processing failed.",
            error=str(exc),
            processing_status="FAILED",
            document_id=document_id,
        )


# ---------------------------------------------------------
# Triggers
# ---------------------------------------------------------

@app.get("/api/triggers")
def get_triggers():
    return [load_trigger_config()]


@app.get("/api/triggers/default")
def get_default_trigger():
    return load_trigger_config()


# ---------------------------------------------------------
# Documents
# ---------------------------------------------------------

def _is_displayable_crawl_document(content_type: Optional[str], url: str) -> bool:
    document_path = urlparse(url).path.lower()
    if content_type in {"Meeting Minutes", "Agendas"}:
        return "/documents/" in document_path and document_path.endswith(".pdf")
    if content_type == "Highlights":
        return "/en/news/meeting-highlights-" in document_path
    if content_type == "PRAC recommendation on signals":
        return document_path.endswith(".pdf")
    return True


def _distinct_finding_products(value: Optional[str]) -> list[str]:
    products = []
    seen = set()
    for candidate in str(value or "").split(","):
        product = candidate.strip()
        normalized = product.casefold()
        if product and normalized not in seen:
            seen.add(normalized)
            products.append(product)
    return products


@app.get("/api/results")
def list_results(
    source_id: Optional[int] = Query(None, ge=1),
    content_type: Optional[str] = None,
    search: Optional[str] = None,
    limit: int = Query(1000, ge=1, le=5000),
):
    conditions = []
    parameters = []

    if source_id is not None:
        conditions.append("d.source_id = ?")
        parameters.append(source_id)

    if content_type:
        conditions.append("d.content_type = ?")
        parameters.append(content_type)

    if search and search.strip():
        query = f"%{search.strip()}%"
        conditions.append(
            """
            (
                d.title LIKE ? OR d.reference_id LIKE ? OR d.url LIKE ?
                OR EXISTS (
                    SELECT 1
                    FROM document_pages search_page
                    WHERE search_page.document_id = d.id
                      AND search_page.page_text LIKE ?
                )
                OR EXISTS (
                    SELECT 1
                    FROM findings search_finding
                    WHERE search_finding.document_id = d.id
                      AND (
                          search_finding.product_name LIKE ?
                          OR search_finding.active_substance LIKE ?
                          OR search_finding.safety_topic LIKE ?
                          OR search_finding.trigger_name LIKE ?
                          OR search_finding.evidence LIKE ?
                      )
                )
            )
            """
        )
        parameters.extend((query, query, query, query, query, query, query, query, query))

    where_clause = f"WHERE {' AND '.join(conditions)}" if conditions else ""
    with get_db() as conn:
        missing_content_types = conn.execute(
            """
            SELECT id, title, url
            FROM documents
            WHERE content_type IS NULL OR content_type = ''
            """
        ).fetchall()
        for document in missing_content_types:
            content_type_value = classify_content_type(
                str(document["title"] or ""),
                str(document["url"] or ""),
            )
            if content_type_value:
                conn.execute(
                    "UPDATE documents SET content_type = ? WHERE id = ?",
                    (content_type_value, document["id"]),
                )

        rows = conn.execute(
            f"""
            SELECT
                d.*,
                s.name AS source_name,
                s.url AS source_url,
                s.reference_no AS source_reference_no,
                s.instruction AS source_instruction,
                (
                    SELECT COUNT(*)
                    FROM document_pages p
                    WHERE p.document_id = d.id
                ) AS pages_extracted,
                (
                    SELECT COUNT(*)
                    FROM extraction_history eh
                    WHERE eh.document_id = d.id
                      AND LOWER(COALESCE(eh.status, '')) IN ('success', 'completed')
                ) AS crawl_count,
                (
                    SELECT COALESCE(eh.completed_at, eh.started_at)
                    FROM extraction_history eh
                    WHERE eh.document_id = d.id
                      AND LOWER(COALESCE(eh.status, '')) IN ('success', 'completed')
                    ORDER BY COALESCE(eh.completed_at, eh.started_at) DESC, eh.id DESC
                    LIMIT 1
                ) AS latest_extraction_at,
                (
                    SELECT COUNT(*)
                    FROM findings f
                    WHERE f.document_id = d.id
                ) AS findings_count,
                (
                    SELECT GROUP_CONCAT(DISTINCT COALESCE(
                        NULLIF(trim(f.product_name), ''),
                        NULLIF(trim(f.active_substance), '')
                    ))
                    FROM findings f
                    WHERE f.document_id = d.id
                      AND COALESCE(
                          NULLIF(trim(f.product_name), ''),
                          NULLIF(trim(f.active_substance), '')
                      ) IS NOT NULL
                ) AS products_found,
                (
                    SELECT a.status
                    FROM document_ai_analyses a
                    WHERE a.document_id = d.id
                      AND a.provider IN ('ollama', 'openai')
                      AND a.analysis_type = 'summary'
                    ORDER BY a.updated_at DESC, a.id DESC
                    LIMIT 1
                ) AS ai_analysis_status,
                (
                    SELECT a.provider
                    FROM document_ai_analyses a
                    WHERE a.document_id = d.id
                      AND a.provider IN ('ollama', 'openai')
                      AND a.analysis_type = 'summary'
                    ORDER BY a.updated_at DESC, a.id DESC
                    LIMIT 1
                ) AS ai_analysis_provider,
                (
                    SELECT a.status
                    FROM document_ai_analyses a
                    WHERE a.document_id = d.id
                      AND a.provider = 'document-extraction'
                      AND a.analysis_type = 'summary'
                    LIMIT 1
                ) AS document_summary_status,
                (
                    SELECT a.updated_at
                    FROM document_ai_analyses a
                    WHERE a.document_id = d.id
                      AND a.provider = 'document-extraction'
                      AND a.analysis_type = 'summary'
                    LIMIT 1
                ) AS document_summary_updated_at
            FROM documents d
            JOIN sources s ON s.id = d.source_id
            {where_clause}
            LIMIT ?
            """,
            (*parameters, limit),
        ).fetchall()

        results = []
        for row in rows:
            document = dict(row)
            content_type_value = document.get("content_type")
            if not _is_displayable_crawl_document(
                content_type_value,
                str(document.get("url") or ""),
            ):
                continue

            reference_id = document.get("reference_id")
            if not reference_id:
                reference_id = _document_reference_id(
                    document.get("source_reference_no"),
                    document.get("publication_date"),
                    int(document["id"]),
                )
                conn.execute(
                    "UPDATE documents SET reference_id = ? WHERE id = ?",
                    (reference_id, document["id"]),
                )

            status = str(document.get("status") or "").upper()
            document_summary_status = str(
                document.get("document_summary_status") or ""
            ).upper()
            if status in {"PROCESSED", "EXTRACTED"}:
                working_status = "Yes"
            elif status in {"FAILED", "EXTRACTION_FAILED"}:
                working_status = "Error"
            else:
                working_status = "Pending"

            if document.get("findings_count"):
                trigger_status = "Found"
            elif status == "PROCESSED":
                trigger_status = "Not Found"
            else:
                trigger_status = "Not Analysed"

            filename = document.get("filename")
            local_pdf = PDF_DIR / filename if filename else None
            has_local_pdf = bool(
                local_pdf
                and local_pdf.is_file()
            )

            document.update(
                {
                    "reference_id": reference_id,
                    "date_crawled": (
                        document.get("last_crawled_at")
                        or document.get("latest_extraction_at")
                    ),
                    "crawled_status": (
                        "Yes"
                        if status in {"PROCESSED", "EXTRACTED"}
                        else "No"
                    ),
                    "working_status": working_status,
                    "safety_trigger_status": trigger_status,
                    "products_found": _distinct_finding_products(
                        document.get("products_found")
                    ),
                    "summary_status": (
                        "AVAILABLE"
                        if document_summary_status == "COMPLETED"
                        else "FAILED"
                        if document_summary_status in {"FAILED", "UNAVAILABLE"}
                        else "PENDING"
                    ),
                    "summary_method": (
                        "Document Extraction / POC Summary"
                        if document_summary_status == "COMPLETED"
                        else None
                    ),
                    "summary_updated_at": document.get("document_summary_updated_at"),
                    "has_local_pdf": has_local_pdf,
                    "local_path": str(local_pdf) if has_local_pdf else None,
                }
            )
            results.append(document)

    def published_sort_key(document):
        value = document.get("publication_date")
        if not value:
            return (0, datetime.min)

        cleaned = " ".join(str(value).split())
        for date_format in (
            "%Y-%m-%d",
            "%d %B %Y",
            "%d %b %Y",
            "%d-%B-%Y",
            "%d-%b-%Y",
        ):
            try:
                return (1, datetime.strptime(cleaned, date_format))
            except ValueError:
                continue

        return (0, datetime.min)

    results.sort(
        key=lambda document: (
            published_sort_key(document),
            int(document["id"]),
        ),
        reverse=True,
    )
    return results


@app.get("/api/results/{document_id}/workflow")
def get_results_workflow(document_id: int):
    with get_db() as conn:
        document = conn.execute(
            """
            SELECT
                d.id, d.title, d.url, d.publication_date, d.content_type,
                d.reference_id, d.last_crawled_at, d.status,
                s.name AS source_name, s.url AS source_url,
                s.instruction AS source_instruction
            FROM documents d
            JOIN sources s ON s.id = d.source_id
            WHERE d.id = ?
            """,
            (document_id,),
        ).fetchone()
        if not document:
            raise HTTPException(status_code=404, detail="Document not found.")

        analysis = conn.execute(
            """
            SELECT id, provider, analysis_type, summary, status,
                   error_message, created_at, updated_at
            FROM document_ai_analyses
            WHERE document_id = ?
              AND provider IN ('ollama', 'openai')
              AND analysis_type = 'summary'
            ORDER BY updated_at DESC, id DESC
            LIMIT 1
            """,
            (document_id,),
        ).fetchone()
        findings = conn.execute(
            """
            SELECT id, product_name, active_substance, trigger_name,
                   safety_topic, evidence, page_number
            FROM findings
            WHERE document_id = ?
            ORDER BY page_number, id
            """,
            (document_id,),
        ).fetchall()
        reviews = conn.execute(
            """
            SELECT r.id, r.finding_id, r.stage, r.reviewer, r.decision,
                   r.comments, r.reviewed_at
            FROM reviews r
            JOIN findings f ON f.id = r.finding_id
            WHERE f.document_id = ?
            ORDER BY r.reviewed_at, r.id
            """,
            (document_id,),
        ).fetchall()
        audit_events = conn.execute(
            """
            SELECT created_at, event_type, details, action, ai_model, result, status
            FROM audit_events
            WHERE entity_type = 'document' AND entity_id = ?
            ORDER BY created_at, id
            """,
            (document_id,),
        ).fetchall()

    analysis_data = dict(analysis) if analysis else None
    if analysis_data:
        model_event = next(
            (
                event for event in reversed(audit_events)
                if str(event["action"] or "").startswith("AI Analysis")
                and event["ai_model"]
            ),
            None,
        )
        analysis_data["model"] = model_event["ai_model"] if model_event else None

    return {
        "document": dict(document),
        "analysis": analysis_data,
        "findings": [dict(row) for row in findings],
        "reviews": [dict(row) for row in reviews],
        "audit_events": [dict(row) for row in audit_events],
        "reviewer_directory": {
            "configured": False,
            "reviewers": [],
            "message": "No configured reviewer directory is available.",
        },
        "notifications_configured": False,
    }


@app.get("/api/documents")
def list_documents(
    limit: int = Query(100, ge=1, le=1000),
):
    with get_db() as conn:
        rows = conn.execute(
            """
            SELECT
                d.*,
                s.authority,
                s.name AS source_name,
                (
                    SELECT COUNT(*)
                    FROM document_pages p
                    WHERE p.document_id = d.id
                ) AS pages_extracted,
                (
                    SELECT COUNT(*)
                    FROM findings f
                    WHERE f.document_id = d.id
                ) AS findings_count,
                (
                    SELECT COUNT(*)
                    FROM findings f
                    WHERE f.document_id = d.id
                      AND UPPER(COALESCE(f.comparison_status, ''))
                          IN ('NEW', 'NEW_INFORMATION')
                ) AS new_information_count,
                (
                    SELECT COUNT(*)
                    FROM reviews r
                    JOIN findings f
                      ON f.id = r.finding_id
                    WHERE f.document_id = d.id
                ) AS review_count,
                (
                    SELECT COUNT(*)
                    FROM summaries s2
                    JOIN findings f2
                      ON f2.id = s2.finding_id
                    WHERE f2.document_id = d.id
                ) AS summary_count,
                (
                    SELECT ch.completed_at
                    FROM crawl_history ch
                    WHERE ch.source_id = d.source_id
                    ORDER BY ch.id DESC
                    LIMIT 1
                ) AS last_crawled,
                (
                    SELECT eh.completed_at
                    FROM extraction_history eh
                    WHERE eh.document_id = d.id
                    ORDER BY eh.id DESC
                    LIMIT 1
                ) AS last_processed
            FROM documents d
            LEFT JOIN sources s
                ON d.source_id = s.id
            ORDER BY d.created_at DESC
            LIMIT ?
            """,
            (limit,),
        ).fetchall()

    documents = []
    with get_db() as conn:
        for row in rows:
            document = dict(row)
            topic_pages = conn.execute(
                """
                SELECT page_number, page_text
                FROM document_pages
                WHERE document_id = ?
                ORDER BY page_number
                """,
                (document["id"],),
            ).fetchall()
            document["topics_count"] = len(
                extract_meeting_topics(
                    [dict(page) for page in topic_pages]
                )
            )
            documents.append(document)

    return documents


def _publication_date_sort_sql() -> str:
    return """
        CASE
            WHEN date(d.publication_date) IS NOT NULL
                THEN date(d.publication_date)
            WHEN length(d.publication_date) >= 11
                 AND substr(d.publication_date, 3, 1) = '-'
                 AND lower(substr(d.publication_date, 4, 3))
                     IN ('jan', 'feb', 'mar', 'apr', 'may', 'jun',
                         'jul', 'aug', 'sep', 'oct', 'nov', 'dec')
                THEN substr(d.publication_date, 8, 4) || '-' ||
                    CASE lower(substr(d.publication_date, 4, 3))
                        WHEN 'jan' THEN '01' WHEN 'feb' THEN '02'
                        WHEN 'mar' THEN '03' WHEN 'apr' THEN '04'
                        WHEN 'may' THEN '05' WHEN 'jun' THEN '06'
                        WHEN 'jul' THEN '07' WHEN 'aug' THEN '08'
                        WHEN 'sep' THEN '09' WHEN 'oct' THEN '10'
                        WHEN 'nov' THEN '11' WHEN 'dec' THEN '12'
                    END || '-' || substr(d.publication_date, 1, 2)
            WHEN instr(d.publication_date, ' ') > 0
                 AND CAST(substr(
                     d.publication_date, 1,
                     instr(d.publication_date, ' ') - 1
                 ) AS INTEGER) > 0
                 AND lower(substr(
                     d.publication_date,
                     instr(d.publication_date, ' ') + 1, 3
                 )) IN ('jan', 'feb', 'mar', 'apr', 'may', 'jun',
                        'jul', 'aug', 'sep', 'oct', 'nov', 'dec')
                THEN substr(d.publication_date, -4) || '-' ||
                    CASE lower(substr(
                        d.publication_date,
                        instr(d.publication_date, ' ') + 1, 3
                    ))
                        WHEN 'jan' THEN '01' WHEN 'feb' THEN '02'
                        WHEN 'mar' THEN '03' WHEN 'apr' THEN '04'
                        WHEN 'may' THEN '05' WHEN 'jun' THEN '06'
                        WHEN 'jul' THEN '07' WHEN 'aug' THEN '08'
                        WHEN 'sep' THEN '09' WHEN 'oct' THEN '10'
                        WHEN 'nov' THEN '11' WHEN 'dec' THEN '12'
                    END || '-' || printf('%02d', CAST(substr(
                        d.publication_date, 1,
                        instr(d.publication_date, ' ') - 1
                    ) AS INTEGER))
            ELSE NULL
        END
    """


def _document_content_type_sql(alias: str = "d") -> str:
    title = f"lower(coalesce({alias}.title, ''))"
    url = f"lower(coalesce({alias}.url, ''))"
    stored_type = f"trim(coalesce({alias}.content_type, ''))"
    return f"""
        CASE
            WHEN {title} LIKE '%new product information wording extracts%'
              OR {url} LIKE '%new-product-information-wording-extracts%'
                THEN 'Other published documents'
            WHEN {title} LIKE '%signal management%'
              OR {url} LIKE '%/signal-management#%'
              OR {url} LIKE '%/signal-management?%'
              OR {url} LIKE '%/signal-management'
                THEN 'Signal Management'
            WHEN {title} LIKE '%list of safety signals%'
              OR {url} LIKE '%#list-of-safety-signals%'
                THEN 'Safety Signals'
            WHEN {title} LIKE '%prac recommendations on safety signals%'
              OR {title} LIKE '%prac recommendation on safety signals%'
              OR {title} LIKE '%prac recommendations signals adopted%'
              OR {url} LIKE '%/documents/prac-recommendation/%'
                THEN 'PRAC Recommendations'
            WHEN lower({stored_type}) = 'prac recommendation on signals'
                THEN 'PRAC Recommendations'
            WHEN {stored_type} != ''
                THEN {alias}.content_type
            WHEN {title} LIKE '%meeting minutes%'
              OR {title} LIKE '%minutes-prac-meeting%'
                THEN 'Meeting Minutes'
            WHEN {title} LIKE '%agenda%'
                THEN 'Agendas'
            WHEN {title} LIKE '%highlight%'
                THEN 'Highlights'
            WHEN {title} LIKE '%work plan%'
              OR {title} LIKE '%work programme%'
                THEN 'Work Plan'
            WHEN {title} LIKE '%rules of procedure%'
                THEN 'Rules of Procedure'
            WHEN {url} LIKE '%.pdf'
              OR {url} LIKE '%.pdf#%'
              OR {url} LIKE '%.pdf?%'
                THEN 'Other published documents'
            ELSE NULL
        END
    """


def _local_document_file(filename: Optional[str]) -> Optional[Path]:
    if not filename or Path(filename).name != filename:
        return None
    pdf_root = PDF_DIR.resolve()
    local_path = (pdf_root / filename).resolve()
    try:
        local_path.relative_to(pdf_root)
    except ValueError:
        return None
    return local_path if local_path.is_file() else None


@app.get("/api/documents-view")
def list_documents_view(
    page: int = Query(1, ge=1),
    page_size: int = Query(10, ge=1, le=50),
    source_id: Optional[int] = Query(None, ge=1),
    content_type: Optional[str] = None,
    status: Optional[str] = None,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    search: Optional[str] = None,
):
    for label, value in (("date_from", date_from), ("date_to", date_to)):
        if value:
            try:
                datetime.strptime(value, "%Y-%m-%d")
            except ValueError as exc:
                raise HTTPException(
                    status_code=422,
                    detail=f"{label} must use YYYY-MM-DD format.",
                ) from exc

    if date_from and date_to and date_from > date_to:
        raise HTTPException(
            status_code=422,
            detail="date_from must not be later than date_to.",
        )

    status_value = (status or "").strip().lower()
    content_type_sql = _document_content_type_sql()
    processed_condition = """
        (
            UPPER(COALESCE(d.status, '')) IN ('PROCESSED', 'EXTRACTED')
            OR EXISTS (
                SELECT 1 FROM extraction_history eh
                WHERE eh.document_id = d.id
                  AND lower(coalesce(eh.status, '')) IN ('success', 'completed')
            )
        )
    """
    analyzed_condition = """
        (
            EXISTS (
                SELECT 1 FROM document_ai_analyses a
                WHERE a.document_id = d.id
                  AND a.analysis_type = 'summary'
                  AND UPPER(COALESCE(a.status, '')) = 'COMPLETED'
            )
            OR EXISTS (
                SELECT 1 FROM findings f
                WHERE f.document_id = d.id
            )
        )
    """
    failed_condition = """
        (
            UPPER(COALESCE(d.status, '')) IN ('FAILED', 'EXTRACTION_FAILED')
            OR EXISTS (
                SELECT 1 FROM extraction_history eh
                WHERE eh.document_id = d.id
                  AND lower(coalesce(eh.status, '')) IN ('failed', 'error')
            )
        )
    """
    status_conditions = {
        "discovered": "(d.discovered_date IS NOT NULL OR d.created_at IS NOT NULL)",
        "crawled": "d.last_crawled_at IS NOT NULL",
        "downloaded": "1 = 1",
        "not_crawled": "d.last_crawled_at IS NULL",
        "processed": processed_condition,
        "working": processed_condition,
        "not_working": f"NOT {processed_condition}",
        "analyzed": analyzed_condition,
        "failed": failed_condition,
    }
    if status_value and status_value not in status_conditions:
        raise HTTPException(status_code=422, detail="Unsupported document status filter.")

    conditions = []
    parameters = []
    if source_id is not None:
        conditions.append("d.source_id = ?")
        parameters.append(source_id)
    if content_type:
        conditions.append(f"({content_type_sql}) = ?")
        parameters.append(content_type.strip())
    if status_value == "downloaded":
        with get_db() as conn:
            file_rows = conn.execute(
                "SELECT id, filename FROM documents WHERE source_id = ?"
                if source_id is not None
                else "SELECT id, filename FROM documents",
                (source_id,) if source_id is not None else (),
            ).fetchall()
        downloaded_ids = [
            int(row["id"])
            for row in file_rows
            if _local_document_file(row["filename"])
        ]
        if downloaded_ids:
            conditions.append(
                f"d.id IN ({', '.join('?' for _ in downloaded_ids)})"
            )
            parameters.extend(downloaded_ids)
        else:
            conditions.append("1 = 0")
    elif status_value:
        conditions.append(status_conditions[status_value])
    if date_from:
        conditions.append(f"({_publication_date_sort_sql()}) >= date(?)")
        parameters.append(date_from)
    if date_to:
        conditions.append(f"({_publication_date_sort_sql()}) <= date(?)")
        parameters.append(date_to)
    if search and search.strip():
        query = f"%{search.strip()}%"
        conditions.append(
            f"""
            (
                d.title LIKE ? OR d.reference_id LIKE ? OR d.url LIKE ?
                OR s.name LIKE ? OR ({content_type_sql}) LIKE ?
                OR EXISTS (
                    SELECT 1 FROM findings f
                    WHERE f.document_id = d.id AND f.product_name LIKE ?
                )
            )
            """
        )
        parameters.extend((query, query, query, query, query, query))
    where_clause = f"WHERE {' AND '.join(conditions)}" if conditions else ""
    publication_date_sort = _publication_date_sort_sql()

    with get_db() as conn:
        rows = conn.execute(
            f"""
            SELECT
                d.*,
                ({content_type_sql}) AS display_content_type,
                s.name AS source_name,
                s.reference_no AS source_reference_no,
                s.url AS source_url,
                (
                    SELECT COUNT(*) FROM findings f
                    WHERE f.document_id = d.id
                ) AS findings_count,
                (
                    SELECT GROUP_CONCAT(DISTINCT COALESCE(
                        NULLIF(trim(f.product_name), ''),
                        NULLIF(trim(f.active_substance), '')
                    ))
                    FROM findings f
                    WHERE f.document_id = d.id
                      AND COALESCE(
                          NULLIF(trim(f.product_name), ''),
                          NULLIF(trim(f.active_substance), '')
                      ) IS NOT NULL
                ) AS products_found,
                (
                    SELECT a.status FROM document_ai_analyses a
                    WHERE a.document_id = d.id
                      AND a.provider IN ('ollama', 'openai')
                      AND a.analysis_type = 'summary'
                    ORDER BY a.updated_at DESC, a.id DESC LIMIT 1
                ) AS ai_analysis_status
            FROM documents d
            LEFT JOIN sources s ON s.id = d.source_id
            {where_clause}
            ORDER BY
                CASE WHEN ({publication_date_sort}) IS NULL THEN 1 ELSE 0 END,
                ({publication_date_sort}) DESC,
                d.id DESC
            LIMIT ? OFFSET ?
            """,
            (*parameters, page_size, (page - 1) * page_size),
        ).fetchall()

        active_sources = conn.execute(
            "SELECT COUNT(*) AS total FROM sources WHERE active = 1"
        ).fetchone()["total"]
        total_sources = conn.execute(
            "SELECT COUNT(*) AS total FROM sources"
        ).fetchone()["total"]
        source_scope = "WHERE source_id = ?" if source_id is not None else ""
        source_parameters = (source_id,) if source_id is not None else ()
        document_stats = conn.execute(
            f"""
            SELECT
                COUNT(*) AS total_documents,
                SUM(CASE
                    WHEN lower(substr(url, 1, CASE
                        WHEN instr(url, '?') > 0 THEN instr(url, '?') - 1
                        WHEN instr(url, '#') > 0 THEN instr(url, '#') - 1
                        ELSE length(url)
                    END)) LIKE '%.pdf'
                      OR lower(coalesce(filename, '')) LIKE '%.pdf'
                    THEN 1 ELSE 0
                END) AS pdf_documents,
                SUM(CASE
                    WHEN lower(substr(url, 1, CASE
                        WHEN instr(url, '?') > 0 THEN instr(url, '?') - 1
                        WHEN instr(url, '#') > 0 THEN instr(url, '#') - 1
                        ELSE length(url)
                    END)) NOT LIKE '%.pdf'
                     AND lower(coalesce(filename, '')) NOT LIKE '%.pdf'
                    THEN 1 ELSE 0
                END) AS html_documents,
                SUM(CASE
                    WHEN date(COALESCE(discovered_date, created_at)) >= date('now', '-6 days')
                    THEN 1 ELSE 0
                END) AS new_documents
            FROM documents
            {source_scope}
            """,
            source_parameters,
        ).fetchone()
        total_documents = document_stats["total_documents"] or 0
        new_documents = document_stats["new_documents"] or 0
        safety_triggers = conn.execute(
            """
            SELECT COUNT(DISTINCT document_id) AS total
            FROM findings
            """
        ).fetchone()["total"]
        source_rows = conn.execute(
            "SELECT id, name, reference_no, url FROM sources ORDER BY name, id"
        ).fetchall()
        source_filter = "WHERE d.source_id = ?" if source_id is not None else ""
        content_type_rows = conn.execute(
            f"""
            SELECT DISTINCT ({content_type_sql}) AS content_type
            FROM documents d
            {source_filter}
            {"AND" if source_filter else "WHERE"}
                ({content_type_sql}) IS NOT NULL
                AND ({content_type_sql}) != ''
            ORDER BY content_type
            """,
            source_parameters,
        ).fetchall()

        available_content_types = {
            str(row["content_type"]).strip()
            for row in content_type_rows
            if str(row["content_type"] or "").strip()
        }
        available_content_types = sorted(
            available_content_types,
            key=str.casefold,
        )

    documents = []
    for row in rows:
        item = dict(row)
        item["content_type"] = item.pop("display_content_type")
        document_status = str(item.get("status") or "").upper()
        if not item.get("reference_id"):
            item["reference_id"] = _document_reference_id(
                item.get("source_reference_no"),
                item.get("publication_date"),
                int(item["id"]),
            )
        item["date_crawled"] = item.get("last_crawled_at") or item.get("discovered_date")
        item["crawled_status"] = "Yes" if item.get("last_crawled_at") else "No"
        item["downloaded_status"] = (
            "Yes" if _local_document_file(item.get("filename")) else "No"
        )
        item["processing_status"] = (
            "Failed"
            if document_status in {"FAILED", "EXTRACTION_FAILED"}
            else "Processed"
            if document_status in {"PROCESSED", "EXTRACTED"}
            else "Discovered"
        )
        item["working_status"] = "Yes" if item["processing_status"] == "Processed" else "No"
        item["analysis_status"] = (
            "Analyzed"
            if str(item.get("ai_analysis_status") or "").upper() == "COMPLETED"
            or int(item.get("findings_count") or 0) > 0
            else "Analysis failed"
            if str(item.get("ai_analysis_status") or "").upper()
            in {"FAILED", "UNAVAILABLE"}
            else "Not analyzed"
        )
        item["safety_trigger_status"] = (
            "Detected"
            if item.get("findings_count")
            else "Not analyzed"
            if item["analysis_status"] == "Not analyzed"
            else "Not detected"
            if item["analysis_status"] == "Analyzed"
            else "Not analyzed"
        )
        item["products_found"] = _distinct_finding_products(
            item.get("products_found")
        )
        filename = item.get("filename")
        local_file = _local_document_file(filename)
        item["has_local_document"] = local_file is not None
        item["has_local_pdf"] = bool(
            local_file and local_file.suffix.lower() == ".pdf"
        )
        item["file_format"] = (
            "PDF"
            if urlparse(str(item.get("url") or "")).path.lower().endswith(".pdf")
            or str(filename or "").lower().endswith(".pdf")
            else "HTML/article"
        )
        documents.append(item)

    with get_db() as conn:
        filtered_count = conn.execute(
            f"""
            SELECT COUNT(*) AS total FROM documents d
            LEFT JOIN sources s ON s.id = d.source_id
            {where_clause}
            """,
            parameters,
        ).fetchone()["total"]

    return {
        "documents": documents,
        "page": page,
        "page_size": page_size,
        "total": filtered_count,
        "stats": {
            "total_documents": total_documents,
            "pdf_documents": document_stats["pdf_documents"] or 0,
            "html_documents": document_stats["html_documents"] or 0,
            "new_documents": new_documents,
            "safety_triggers": safety_triggers,
            "total_sources": total_sources,
            "active_sources": active_sources,
        },
        "filters": {
            "sources": [dict(row) for row in source_rows],
            "content_types": available_content_types,
        },
    }


@app.get("/api/sources/{source_id}/documents")
def list_source_documents(
    source_id: int,
    limit: int = Query(1000, ge=1, le=5000),
):
    with get_db() as conn:
        source = conn.execute(
            """
            SELECT id, reference_no, name
            FROM sources
            WHERE id = ?
            """,
            (source_id,),
        ).fetchone()
        if not source:
            raise HTTPException(status_code=404, detail="Source not found.")

        rows = conn.execute(
            """
            SELECT
                d.id,
                d.source_id,
                d.title,
                d.url,
                d.meeting_date,
                d.publication_date,
                d.filename,
                d.discovered_date,
                d.last_crawled_at,
                d.content_type,
                d.reference_id,
                d.status,
                d.created_at
            FROM documents d
            WHERE d.source_id = ?
            ORDER BY
                CASE WHEN d.publication_date IS NULL OR d.publication_date = ''
                    THEN 1 ELSE 0 END,
                d.publication_date DESC,
                d.id DESC
            LIMIT ?
            """,
            (source_id, limit),
        ).fetchall()

    documents = []
    source_reference = source["reference_no"] or source_id
    for row in rows:
        item = dict(row)
        item["content_type"] = item.get("content_type") or classify_content_type(
            str(item.get("title") or ""),
            str(item.get("url") or ""),
        )
        if not _is_displayable_crawl_document(
            item["content_type"],
            str(item.get("url") or ""),
        ):
            continue
        if not item.get("reference_id"):
            item["reference_id"] = _document_reference_id(
                source_reference,
                item.get("publication_date"),
                int(item["id"]),
            )
        status = str(item.get("status") or "").upper()
        item["source_name"] = source["name"]
        item["source_reference_no"] = source_reference
        item["crawled_status"] = (
            "Yes"
            if item.get("discovered_date") or item.get("last_crawled_at")
            else "No"
        )
        if status in {"PROCESSED", "EXTRACTED"}:
            item["working_status"] = "Yes"
        elif status in {"FAILED", "EXTRACTION_FAILED"}:
            item["working_status"] = "Error"
        else:
            item["working_status"] = "Pending"
        documents.append(item)

    documents.sort(
        key=lambda item: (
            _published_date_token(item.get("publication_date")) != "UNDATED",
            _published_date_token(item.get("publication_date")),
            int(item["id"]),
        ),
        reverse=True,
    )
    return documents


@app.get("/api/documents/{document_id}")
def get_document(document_id: int):
    with get_db() as conn:
        row = conn.execute(
            """
            SELECT
                d.*,
                s.authority,
                s.name AS source_name,
                (
                    SELECT COUNT(*)
                    FROM document_pages p
                    WHERE p.document_id = d.id
                ) AS pages_extracted,
                (
                    SELECT COUNT(*)
                    FROM findings f
                    WHERE f.document_id = d.id
                ) AS findings_count,
                (
                    SELECT COUNT(*)
                    FROM findings f
                    WHERE f.document_id = d.id
                      AND UPPER(COALESCE(f.comparison_status, ''))
                          IN ('NEW', 'NEW_INFORMATION')
                ) AS new_information_count,
                (
                    SELECT COUNT(*)
                    FROM reviews r
                    JOIN findings f
                      ON f.id = r.finding_id
                    WHERE f.document_id = d.id
                ) AS review_count,
                (
                    SELECT COUNT(*)
                    FROM summaries s2
                    JOIN findings f2
                      ON f2.id = s2.finding_id
                    WHERE f2.document_id = d.id
                ) AS summary_count,
                (
                    SELECT ch.completed_at
                    FROM crawl_history ch
                    WHERE ch.source_id = d.source_id
                    ORDER BY ch.id DESC
                    LIMIT 1
                ) AS last_crawled,
                (
                    SELECT eh.completed_at
                    FROM extraction_history eh
                    WHERE eh.document_id = d.id
                    ORDER BY eh.id DESC
                    LIMIT 1
                ) AS last_processed
            FROM documents d
            LEFT JOIN sources s
                ON d.source_id = s.id
            WHERE d.id = ?
            """,
            (document_id,),
        ).fetchone()
        topic_pages = conn.execute(
            """
            SELECT page_number, page_text
            FROM document_pages
            WHERE document_id = ?
            ORDER BY page_number
            """,
            (document_id,),
        ).fetchall()

    if not row:
        raise HTTPException(
            status_code=404,
            detail="Document not found.",
        )

    document = dict(row)
    document["topics_count"] = len(
        extract_meeting_topics(
            [dict(page) for page in topic_pages]
        )
    )
    return document


@app.get("/api/documents/{document_id}/pages")
def get_document_pages(document_id: int):
    with get_db() as conn:
        rows = conn.execute(
            """
            SELECT *
            FROM document_pages
            WHERE document_id = ?
            ORDER BY page_number
            """,
            (document_id,),
        ).fetchall()

    return [dict(row) for row in rows]


@app.get("/api/documents/{document_id}/topics")
def get_document_topics(document_id: int):
    with get_db() as conn:
        document = conn.execute(
            """
            SELECT id, title, meeting_date
            FROM documents
            WHERE id = ?
            """,
            (document_id,),
        ).fetchone()

        if not document:
            raise HTTPException(
                status_code=404,
                detail="Document not found.",
            )

        pages = conn.execute(
            """
            SELECT page_number, page_text
            FROM document_pages
            WHERE document_id = ?
            ORDER BY page_number
            """,
            (document_id,),
        ).fetchall()

    topics = extract_meeting_topics(
        [dict(page) for page in pages]
    )

    return {
        "document_id": document_id,
        "document_title": document["title"],
        "meeting_date": document["meeting_date"],
        "topics": topics,
        "message": (
            None
            if topics
            else "No meeting topics could be extracted from this document."
        ),
    }


@app.get("/api/documents/{document_id}/summary")
def get_document_summary(document_id: int):
    with get_db() as conn:
        document = conn.execute(
            """
            SELECT d.*, s.authority, s.name AS source_name,
                   s.source_type, s.url AS source_url, s.active,
                   s.instruction AS source_instruction
            FROM documents d
            LEFT JOIN sources s ON s.id = d.source_id
            WHERE d.id = ?
            """,
            (document_id,),
        ).fetchone()
        if not document:
            raise HTTPException(status_code=404, detail="Document not found.")

        pages = conn.execute(
            """
            SELECT page_number, page_text
            FROM document_pages
            WHERE document_id = ?
            ORDER BY page_number
            """,
            (document_id,),
        ).fetchall()
        findings = conn.execute(
            """
            SELECT f.*, EXISTS(
                SELECT 1 FROM reviews r WHERE r.finding_id = f.id
            ) AS reviewed
            FROM findings f
            WHERE f.document_id = ?
            ORDER BY f.page_number, f.id
            """,
            (document_id,),
        ).fetchall()
        extraction = conn.execute(
            """
            SELECT started_at, completed_at, status, pages_extracted
            FROM extraction_history
            WHERE document_id = ?
            ORDER BY id DESC LIMIT 1
            """,
            (document_id,),
        ).fetchone()
        crawl = conn.execute(
            """
            SELECT ch.started_at, ch.completed_at, ch.status,
                   ch.documents_found
            FROM crawl_history ch
            WHERE ch.source_id = ?
            ORDER BY ch.id DESC LIMIT 1
            """,
            (document["source_id"],),
        ).fetchone()
        stored_summary = conn.execute(
            """
            SELECT status
            FROM document_ai_analyses
            WHERE document_id = ?
              AND provider = 'document-extraction'
              AND analysis_type = 'summary'
            """,
            (document_id,),
        ).fetchone()

    topics = extract_meeting_topics([dict(page) for page in pages])
    section_summaries = extract_section_summaries(
        [dict(page) for page in pages],
        topics,
    )
    finding_items = [dict(row) for row in findings]
    new_information = [
        finding for finding in finding_items
        if str(finding.get("comparison_status") or "").upper()
        in {"NEW", "NEW_INFORMATION"}
    ]
    key_points = _extractive_summary_sentences(
        [dict(page) for page in pages],
        limit=5,
    )
    executive_summary = _document_executive_summary(
        document["content_type"],
        section_summaries,
        key_points,
    )
    status = (
        "AVAILABLE"
        if stored_summary and str(stored_summary["status"]).upper() == "COMPLETED"
        else "FAILED"
        if stored_summary and str(stored_summary["status"]).upper() == "FAILED"
        else "PENDING"
    )
    return {
        "document": dict(document),
        "topics": topics,
        "section_summaries": section_summaries,
        "executive_summary": executive_summary,
        "key_points": key_points,
        "source_instruction": document["source_instruction"],
        "findings": finding_items,
        "new_information": new_information,
        "pages_extracted": len(pages),
        "topics_count": len(topics),
        "findings_count": len(finding_items),
        "new_information_count": len(new_information),
        "crawl": dict(crawl) if crawl else None,
        "processing": dict(extraction) if extraction else None,
        "configuration": {
            "source_instruction": document["source_instruction"],
            "section_xpath": PRAC_SECTION_XPATH,
            "meeting_pdf_xpath": PRAC_MEETING_XPATH,
            "trigger": load_trigger_config(),
        },
        "summary_status": status,
        "summary_method": "Document Extraction / POC Summary",
        "summary_note": (
            "This structured summary is extracted from document headings, "
            "key sentences, and stored safety findings. No AI provider was used."
        ),
    }


def _document_executive_summary(
    content_type: Optional[str],
    section_summaries: list[dict],
    key_points: list[str],
) -> str:
    if key_points:
        return "Key safety and regulatory points include: " + " ".join(key_points[:2])
    titles = [
        str(section["title"])
        for section in section_summaries[:8]
        if section.get("title")
    ]
    if titles:
        return f"The {str(content_type or 'document').lower()} covers " + ", ".join(titles) + "."
    return "Source text was insufficient to generate a detailed deterministic summary."


def _extractive_summary_sentences(pages: list[dict], limit: int = 5) -> list[str]:
    safety_terms = (
        "safety", "adverse", "risk", "recommend", "signal", "product",
        "medicinal", "benefit", "assessment", "monitor", "action",
    )
    candidates = []
    seen = set()
    seen_semantic = set()
    for page in pages:
        text = str(page.get("page_text") or page.get("text") or "")
        text = re.sub(
            r"(?im)^(?:pharmacovigilance risk assessment committee|ema/prac/|page\s+\d+).*$",
            "",
            text,
        )
        for sentence in re.split(r"(?<=[.!?])\s+", re.sub(r"\s+", " ", text)):
            cleaned = sentence.strip(" -\t\r\n")
            key = cleaned.casefold()
            if len(cleaned) < 45 or len(cleaned) > 420 or key in seen:
                continue
            if re.search(r"\.{3,}\s*\d+", cleaned):
                continue
            seen.add(key)
            semantic_key = re.sub(
                r"benefit[- ]risk balance of .+? remains unchanged",
                "benefit-risk balance remains unchanged",
                key,
            )
            if semantic_key in seen_semantic:
                continue
            seen_semantic.add(semantic_key)
            relevance = sum(term in key for term in safety_terms)
            candidates.append((relevance, len(candidates), cleaned))
    candidates.sort(key=lambda item: (-item[0], item[1]))
    return [sentence for _, _, sentence in candidates[:limit]]


@app.post("/api/documents/{document_id}/summary/generate")
def generate_poc_document_summary(document_id: int):
    summary_data = get_document_summary(document_id)
    if (
        str(summary_data["document"]["status"] or "").upper()
        not in {"PROCESSED", "EXTRACTED"}
        or not summary_data["pages_extracted"]
    ):
        raise HTTPException(
            status_code=409,
            detail="Crawl and extract the document before generating its summary.",
        )
    if not (
        summary_data["section_summaries"]
        or summary_data["key_points"]
        or summary_data["topics"]
        or summary_data["findings"]
    ):
        failed_at = datetime.utcnow().isoformat()
        with get_db() as conn:
            conn.execute(
                """
                INSERT INTO document_ai_analyses
                    (document_id, provider, analysis_type, summary, status,
                     error_message, created_at, updated_at)
                VALUES (?, 'document-extraction', 'summary', NULL, 'FAILED',
                        ?, ?, ?)
                ON CONFLICT(document_id, provider, analysis_type)
                DO UPDATE SET summary = NULL, status = 'FAILED',
                              error_message = excluded.error_message,
                              updated_at = excluded.updated_at
                """,
                (
                    document_id,
                    "Insufficient extracted content to generate a summary.",
                    failed_at,
                    failed_at,
                ),
            )
        raise HTTPException(
            status_code=422,
            detail="The extracted document does not contain enough content to summarize.",
        )

    generated_at = datetime.utcnow().isoformat()
    stored_summary = {
        "method": summary_data["summary_method"],
        "source_instruction": summary_data["source_instruction"],
        "executive_summary": summary_data["executive_summary"],
        "section_summaries": summary_data["section_summaries"],
        "topics": summary_data["topics"],
        "key_points": summary_data["key_points"],
        "safety_information": summary_data["findings"],
        "products": sorted({
            str(finding.get("product_name") or finding.get("active_substance")).strip()
            for finding in summary_data["findings"]
            if str(finding.get("product_name") or finding.get("active_substance") or "").strip()
        }),
        "generated_at": generated_at,
    }
    encoded_summary = json.dumps(stored_summary, ensure_ascii=False)
    with get_db() as conn:
        conn.execute(
            """
            INSERT INTO document_ai_analyses
                (document_id, provider, analysis_type, summary, status,
                 error_message, created_at, updated_at)
            VALUES (?, 'document-extraction', 'summary', ?, 'COMPLETED',
                    NULL, ?, ?)
            ON CONFLICT(document_id, provider, analysis_type)
            DO UPDATE SET summary = excluded.summary, status = 'COMPLETED',
                          error_message = NULL, updated_at = excluded.updated_at
            """,
            (document_id, encoded_summary, generated_at, generated_at),
        )
    return {
        "document_id": document_id,
        "status": "COMPLETED",
        "summary_method": summary_data["summary_method"],
        "summary": stored_summary,
        "updated_at": generated_at,
    }


@app.get("/api/documents/{document_id}/metadata")
def document_metadata(document_id: int):
    with get_db() as conn:
        row = conn.execute(
            """
            SELECT filename
            FROM documents
            WHERE id = ?
            """,
            (document_id,),
        ).fetchone()

    if not row:
        raise HTTPException(
            status_code=404,
            detail="Document not found.",
        )

    filename = row["filename"]

    if not filename:
        raise HTTPException(
            status_code=404,
            detail="PDF filename is not available.",
        )

    pdf_path = PDF_DIR / filename

    if not pdf_path.exists():
        raise HTTPException(
            status_code=404,
            detail="PDF file not found.",
        )

    return get_pdf_metadata(pdf_path)


@app.get("/api/documents/{document_id}/pdf")
def download_document_pdf(document_id: int, inline: bool = False):
    with get_db() as conn:
        row = conn.execute(
            """
            SELECT d.filename, d.title, d.url, s.url AS source_url
            FROM documents d
            JOIN sources s ON s.id = d.source_id
            WHERE d.id = ?
            """,
            (document_id,),
        ).fetchone()

    if not row:
        raise HTTPException(
            status_code=404,
            detail="Document not found.",
        )

    filename = row["filename"]

    pdf_path = PDF_DIR / filename if filename else None
    temporary_download = False
    if not pdf_path or not pdf_path.is_file():
        document_url = str(row["url"] or "")
        source_host = (urlparse(str(row["source_url"] or "")).hostname or "").lower()
        document_host = (urlparse(document_url).hostname or "").lower()
        if not urlparse(document_url).path.lower().endswith(".pdf"):
            raise HTTPException(
                status_code=404,
                detail="This document is not an available PDF download.",
            )
        if not source_host or document_host != source_host:
            raise HTTPException(
                status_code=400,
                detail="Document download host does not match its configured source.",
            )

        download_name = f"document-{document_id}-{uuid.uuid4().hex}.pdf"
        pdf_path = PDF_DIR / "downloads" / download_name
        try:
            downloaded_path = download_pdf(document_url, str(pdf_path))
            if not downloaded_path:
                raise HTTPException(
                    status_code=502,
                    detail="Unable to download the original document.",
                )
            pdf_path = Path(downloaded_path)
            temporary_download = True
        except HTTPException:
            pdf_path.unlink(missing_ok=True)
            pdf_path.with_suffix(".part").unlink(missing_ok=True)
            raise
        except Exception as exc:
            logger.exception("Original PDF download failed for document %s", document_id)
            pdf_path.unlink(missing_ok=True)
            pdf_path.with_suffix(".part").unlink(missing_ok=True)
            raise HTTPException(
                status_code=502,
                detail="Unable to download the original document.",
            ) from exc

    return FileResponse(
        path=str(pdf_path),
        media_type="application/pdf",
        filename=filename or Path(urlparse(str(row["url"])).path).name,
        content_disposition_type="inline" if inline else "attachment",
        background=(
            BackgroundTask(pdf_path.unlink, missing_ok=True)
            if temporary_download
            else None
        ),
    )


# ---------------------------------------------------------
# Pages
# ---------------------------------------------------------

@app.get("/api/pages/{document_id}/{page_number}")
def get_page(
    document_id: int,
    page_number: int,
):
    with get_db() as conn:
        row = conn.execute(
            """
            SELECT *
            FROM document_pages
            WHERE document_id = ?
              AND page_number = ?
            """,
            (
                document_id,
                page_number,
            ),
        ).fetchone()

    if not row:
        raise HTTPException(
            status_code=404,
            detail="Page not found.",
        )

    return dict(row)


# ---------------------------------------------------------
# Findings
# ---------------------------------------------------------

def _finding_review_history(conn, finding_ids):
    if not finding_ids:
        return {}
    placeholders = ",".join("?" for _ in finding_ids)
    rows = conn.execute(
        f"""
        SELECT * FROM reviews
        WHERE finding_id IN ({placeholders})
        ORDER BY reviewed_at, id
        """,
        tuple(finding_ids),
    ).fetchall()
    history = {}
    for row in rows:
        review = dict(row)
        finding_reviews = history.setdefault(int(review["finding_id"]), [])
        finding_reviews.append(review)
    return history


def _configured_product_class_rows(conn):
    rows = conn.execute(
        """
        SELECT p.product_name, pc.class_name
        FROM products p
        JOIN product_classes pc ON pc.id = p.product_class_id
        WHERE p.active = 1 AND pc.active = 1
        ORDER BY pc.class_name, p.product_name
        """
    ).fetchall()
    return [
        {"product_name": row["product_name"], "class_name": row["class_name"]}
        for row in rows
    ]


def _matching_configured_product_classes(finding, configured_products):
    product_name = str(finding.get("product_name") or "").strip().casefold()
    if not product_name:
        return []
    return [
        item
        for item in configured_products
        if str(item["product_name"] or "").strip().casefold() == product_name
    ]


@app.get("/api/findings")
def list_findings(
    limit: int = Query(100, ge=1, le=1000),
    document_id: Optional[int] = Query(None, ge=1),
    trigger_name: Optional[str] = Query(None),
    comparison_status: Optional[str] = Query(None),
    review_status: Optional[str] = Query(None),
):
    with get_db() as conn:
        query = """
            SELECT
                f.*,
                d.title AS document_title,
                d.filename,
                d.url AS document_url
            FROM findings f
            LEFT JOIN documents d
                ON f.document_id = d.id
        """
        parameters = []
        conditions = []
        if document_id is not None:
            conditions.append("f.document_id = ?")
            parameters.append(document_id)
        if trigger_name:
            conditions.append("LOWER(COALESCE(f.trigger_name, '')) = LOWER(?)")
            parameters.append(trigger_name)
        if comparison_status:
            if comparison_status.upper() == "NEW_INFORMATION":
                conditions.append(
                    "UPPER(COALESCE(f.comparison_status, '')) IN ('NEW', 'NEW_INFORMATION')"
                )
            else:
                conditions.append(
                    "UPPER(COALESCE(f.comparison_status, '')) = UPPER(?)"
                )
                parameters.append(comparison_status)
        if review_status:
            normalized_review_status = review_status.upper()
            if normalized_review_status == "REVIEWED":
                conditions.append(
                    """
                    EXISTS (
                        SELECT 1 FROM reviews rr
                        WHERE rr.finding_id = f.id AND rr.stage = 'FINAL'
                    )
                    """
                )
            elif normalized_review_status == "AWAITING_FINAL_REVIEW":
                conditions.append(
                    """
                    EXISTS (
                        SELECT 1 FROM reviews ri
                        WHERE ri.finding_id = f.id AND ri.stage = 'INITIAL'
                    )
                    AND NOT EXISTS (
                        SELECT 1 FROM reviews rf
                        WHERE rf.finding_id = f.id AND rf.stage = 'FINAL'
                    )
                    """
                )
            elif normalized_review_status == "PENDING":
                conditions.append(
                    """
                    NOT EXISTS (
                        SELECT 1 FROM reviews ri
                        WHERE ri.finding_id = f.id AND ri.stage = 'INITIAL'
                    )
                    """
                )
        if conditions:
            query += " WHERE " + " AND ".join(conditions)
        query += " ORDER BY f.created_at DESC LIMIT ?"
        parameters.append(limit)
        rows = conn.execute(query, parameters).fetchall()

    findings = []
    with get_db() as conn:
        review_history = _finding_review_history(
            conn,
            [int(row["id"]) for row in rows],
        )
        configured_products = _configured_product_class_rows(conn)
        for row in rows:
            finding = dict(row)
            finding["configured_product_classes"] = _matching_configured_product_classes(
                finding,
                configured_products,
            )
            for field in (
                "relevant_terms",
                "section",
                "evidence_hash",
            ):
                finding.setdefault(field, None)
            summary = conn.execute(
                """
                SELECT *
                FROM summaries
                WHERE finding_id = ?
                ORDER BY created_at DESC
                LIMIT 1
                """,
                (finding["id"],),
            ).fetchone()
            finding["summary"] = dict(summary) if summary else None
            reviews = review_history.get(int(finding["id"]), [])
            finding["review_history"] = reviews
            finding["initial_review"] = next(
                (review for review in reversed(reviews) if review.get("stage") == "INITIAL"),
                None,
            )
            finding["final_review"] = next(
                (review for review in reversed(reviews) if review.get("stage") == "FINAL"),
                None,
            )
            finding["latest_review"] = reviews[-1] if reviews else None
            finding["review_status"] = (
                "REVIEWED" if finding["final_review"]
                else "AWAITING_FINAL_REVIEW" if finding["initial_review"]
                else "PENDING"
            )
            if finding["summary"]:
                finding["summary"]["model"] = finding["summary"].get("model") or os.getenv(
                    "OLLAMA_MODEL", "llama3.2"
                )
                finding["summary"]["status"] = "GENERATED"
            findings.append(finding)

    return findings


@app.get("/api/findings/new")
def list_new_information(
    limit: int = Query(100, ge=1, le=1000),
):
    with get_db() as conn:
        rows = conn.execute(
            """
            SELECT
                f.*,
                d.title AS document_title,
                d.filename,
                d.url AS document_url
            FROM findings f
            LEFT JOIN documents d
                ON f.document_id = d.id
            WHERE UPPER(COALESCE(f.comparison_status, ''))
                IN ('NEW', 'NEW_INFORMATION')
            ORDER BY f.created_at DESC
            LIMIT ?
            """,
            (limit,),
        ).fetchall()

        findings = []
        for row in rows:
            finding = dict(row)
            for field in (
                "relevant_terms",
                "section",
                "evidence_hash",
            ):
                finding.setdefault(field, None)
            summary = conn.execute(
                """
                SELECT *
                FROM summaries
                WHERE finding_id = ?
                ORDER BY created_at DESC
                LIMIT 1
                """,
                (finding["id"],),
            ).fetchone()
            finding["summary"] = dict(summary) if summary else None
            review = conn.execute(
                """
                SELECT id, reviewer, decision, comments, reviewed_at
                FROM reviews
                WHERE finding_id = ?
                ORDER BY reviewed_at DESC, id DESC
                LIMIT 1
                """,
                (finding["id"],),
            ).fetchone()
            finding["latest_review"] = dict(review) if review else None
            finding["review_status"] = "REVIEWED" if review else "PENDING"
            if finding["summary"]:
                finding["summary"]["model"] = finding["summary"].get("model") or os.getenv(
                    "OLLAMA_MODEL", "llama3.2"
                )
                finding["summary"]["status"] = "GENERATED"
            findings.append(finding)

    return findings


@app.get("/api/new-information")
def list_new_information_documents(
    page: int = Query(1, ge=1),
    page_size: int = Query(5, ge=1, le=50),
    date_range: str = Query("all"),
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    source_id: Optional[int] = Query(None, ge=1),
    content_type: Optional[str] = None,
    product_class_id: Optional[int] = Query(None, ge=1),
    trigger_name: Optional[str] = None,
    safety_trigger: Optional[str] = None,
    search: Optional[str] = None,
):
    for label, value in (("date_from", date_from), ("date_to", date_to)):
        if value:
            try:
                datetime.strptime(value, "%Y-%m-%d")
            except ValueError as exc:
                raise HTTPException(
                    status_code=422,
                    detail=f"{label} must use YYYY-MM-DD format.",
                ) from exc
    if date_from and date_to and date_from > date_to:
        raise HTTPException(
            status_code=422,
            detail="date_from must not be later than date_to.",
        )

    publication_date_sql = _publication_date_sort_sql()
    conditions = [
        "UPPER(COALESCE(f_new.comparison_status, '')) IN ('NEW', 'NEW_INFORMATION')"
    ]
    parameters = []
    if date_range == "custom":
        if date_from:
            conditions.append(f"date({publication_date_sql}) >= date(?)")
            parameters.append(date_from)
        if date_to:
            conditions.append(f"date({publication_date_sql}) <= date(?)")
            parameters.append(date_to)
    elif date_range in {"7d", "30d", "90d"}:
        days = int(date_range[:-1])
        conditions.append(
            f"date({publication_date_sql}) >= date('now', ?)"
        )
        parameters.append(f"-{days} days")
    elif date_range != "all":
        raise HTTPException(status_code=422, detail="Unsupported date range.")
    if source_id is not None:
        conditions.append("d.source_id = ?")
        parameters.append(source_id)
    if content_type:
        conditions.append("LOWER(trim(d.content_type)) = LOWER(trim(?))")
        parameters.append(content_type.strip())
    if product_class_id is not None:
        conditions.append(
            """
            EXISTS (
                SELECT 1
                FROM findings f_class
                JOIN products p_class
                  ON p_class.product_class_id = ?
                 AND p_class.active = 1
                WHERE f_class.document_id = d.id
                  AND (
                    instr(
                        lower(COALESCE(f_class.product_name, '')),
                        lower(p_class.product_name)
                    ) > 0
                    OR instr(
                        lower(COALESCE(f_class.active_substance, '')),
                        lower(p_class.product_name)
                    ) > 0
                  )
            )
            """
        )
        parameters.append(product_class_id)
    if trigger_name:
        conditions.append(
            """
            EXISTS (
                SELECT 1 FROM findings f_trigger
                WHERE f_trigger.document_id = d.id
                  AND UPPER(COALESCE(f_trigger.comparison_status, ''))
                      IN ('NEW', 'NEW_INFORMATION')
                  AND lower(trim(COALESCE(f_trigger.trigger_name, ''))) = lower(trim(?))
            )
            """
        )
        parameters.append(trigger_name.strip())
    safety_value = (safety_trigger or "").strip().lower()
    if safety_value == "detected":
        conditions.append(
            """
            EXISTS (
                SELECT 1 FROM findings f_detected
                WHERE f_detected.document_id = d.id
                  AND UPPER(COALESCE(f_detected.comparison_status, ''))
                      IN ('NEW', 'NEW_INFORMATION')
            )
            """
        )
    elif safety_value == "not_detected":
        conditions.append(
            """
            NOT EXISTS (
                SELECT 1 FROM findings f_detected
                WHERE f_detected.document_id = d.id
                  AND UPPER(COALESCE(f_detected.comparison_status, ''))
                      IN ('NEW', 'NEW_INFORMATION')
            )
            """
        )
    elif safety_value:
        raise HTTPException(
            status_code=422,
            detail="Unsupported safety trigger filter.",
        )
    if search and search.strip():
        search_like = f"%{search.strip()}%"
        conditions.append(
            """
            (
                d.title LIKE ? OR d.reference_id LIKE ? OR s.name LIKE ?
                OR d.content_type LIKE ?
                OR EXISTS (
                    SELECT 1 FROM findings f_search
                    WHERE f_search.document_id = d.id
                      AND UPPER(COALESCE(f_search.comparison_status, ''))
                          IN ('NEW', 'NEW_INFORMATION')
                      AND (
                          f_search.product_name LIKE ?
                          OR f_search.active_substance LIKE ?
                          OR f_search.trigger_name LIKE ?
                          OR f_search.safety_topic LIKE ?
                          OR f_search.evidence LIKE ?
                      )
                )
            )
            """
        )
        parameters.extend([search_like] * 9)
    where_clause = " AND ".join(conditions)

    with get_db() as conn:
        total = conn.execute(
            f"""
            SELECT COUNT(*) AS total
            FROM documents d
            JOIN sources s ON s.id = d.source_id
            WHERE EXISTS (
                SELECT 1 FROM findings f_new
                WHERE f_new.document_id = d.id
                  AND {where_clause}
            )
            """,
            parameters,
        ).fetchone()["total"]
        document_rows = conn.execute(
            f"""
            SELECT d.*, s.name AS source_name, s.reference_no AS source_reference_no,
                   s.url AS source_url,
                   (
                       SELECT COUNT(*)
                       FROM findings f_all
                       WHERE f_all.document_id = d.id
                   ) AS findings_count,
                   (
                       SELECT GROUP_CONCAT(DISTINCT COALESCE(
                           NULLIF(trim(f_product.product_name), ''),
                           NULLIF(trim(f_product.active_substance), '')
                       ))
                       FROM findings f_product
                       WHERE f_product.document_id = d.id
                         AND COALESCE(
                             NULLIF(trim(f_product.product_name), ''),
                             NULLIF(trim(f_product.active_substance), '')
                         ) IS NOT NULL
                   ) AS products_found,
                   (
                       SELECT COUNT(*)
                       FROM findings f_new_count
                       WHERE f_new_count.document_id = d.id
                         AND UPPER(COALESCE(f_new_count.comparison_status, ''))
                             IN ('NEW', 'NEW_INFORMATION')
                   ) AS new_trigger_count
            FROM documents d
            JOIN sources s ON s.id = d.source_id
            WHERE EXISTS (
                SELECT 1 FROM findings f_new
                WHERE f_new.document_id = d.id
                  AND {where_clause}
            )
            ORDER BY
                CASE WHEN ({publication_date_sql}) IS NULL THEN 1 ELSE 0 END,
                ({publication_date_sql}) DESC,
                d.id DESC
            LIMIT ? OFFSET ?
            """,
            (*parameters, page_size, (page - 1) * page_size),
        ).fetchall()
        source_rows = conn.execute(
            "SELECT id, name, reference_no, url FROM sources ORDER BY name, id"
        ).fetchall()
        content_type_rows = conn.execute(
            """
            SELECT DISTINCT d.content_type
            FROM documents d
            JOIN findings f ON f.document_id = d.id
            WHERE UPPER(COALESCE(f.comparison_status, ''))
                IN ('NEW', 'NEW_INFORMATION')
              AND d.content_type IS NOT NULL AND trim(d.content_type) != ''
            ORDER BY d.content_type
            """
        ).fetchall()
        trigger_rows = conn.execute(
            """
            SELECT DISTINCT f.trigger_name
            FROM findings f
            WHERE UPPER(COALESCE(f.comparison_status, ''))
                IN ('NEW', 'NEW_INFORMATION')
              AND f.trigger_name IS NOT NULL AND trim(f.trigger_name) != ''
            ORDER BY f.trigger_name
            """
        ).fetchall()
        product_classes = conn.execute(
            """
            SELECT id, class_name FROM product_classes
            WHERE active = 1 ORDER BY class_name
            """
        ).fetchall()
        stats_rows = conn.execute(
            f"""
            SELECT d.source_id,
                   s.name AS source_name,
                   COUNT(DISTINCT d.id) AS document_count,
                   COUNT(f_new.id) AS finding_count,
                   GROUP_CONCAT(DISTINCT COALESCE(
                       NULLIF(trim(f_new.product_name), ''),
                       NULLIF(trim(f_new.active_substance), '')
                   )) AS products
            FROM documents d
            JOIN sources s ON s.id = d.source_id
            JOIN findings f_new ON f_new.document_id = d.id
            WHERE {where_clause}
            GROUP BY d.source_id, s.name
            """,
            parameters,
        ).fetchall()
        current_week_clause = "date(f.created_at) >= date('now', '-7 days')"
        previous_week_clause = (
            "date(f.created_at) >= date('now', '-14 days') "
            "AND date(f.created_at) < date('now', '-7 days')"
        )
        recent_count = conn.execute(
            f"""
            SELECT COUNT(DISTINCT d.id) AS total
            FROM documents d JOIN findings f ON f.document_id = d.id
            WHERE UPPER(COALESCE(f.comparison_status, ''))
                IN ('NEW', 'NEW_INFORMATION')
              AND {current_week_clause}
            """
        ).fetchone()["total"]
        previous_count = conn.execute(
            f"""
            SELECT COUNT(DISTINCT d.id) AS total
            FROM documents d JOIN findings f ON f.document_id = d.id
            WHERE UPPER(COALESCE(f.comparison_status, ''))
                IN ('NEW', 'NEW_INFORMATION')
              AND {previous_week_clause}
            """
        ).fetchone()["total"]

    products = set()
    sources = {}
    potential_triggers = 0
    for row in stats_rows:
        potential_triggers += int(row["finding_count"] or 0)
        sources[str(row["source_id"])] = row["source_name"]
        products.update(_distinct_finding_products(row["products"]))

    documents = []
    for row in document_rows:
        item = dict(row)
        item["reference_id"] = item.get("reference_id") or _document_reference_id(
            item.get("source_reference_no"),
            item.get("publication_date"),
            int(item["id"]),
        )
        item["date_crawled"] = item.get("last_crawled_at") or item.get("discovered_date")
        item["crawled_status"] = (
            "Yes"
            if item["date_crawled"]
            or str(item.get("status") or "").upper()
                in {"PROCESSED", "EXTRACTED", "ALREADY_CRAWLED"}
            else "No"
        )
        item["working_status"] = (
            "Yes"
            if str(item.get("status") or "").upper() in {"PROCESSED", "EXTRACTED"}
            else "No"
        )
        item["products_found"] = _distinct_finding_products(item.get("products_found"))
        item["safety_trigger_count"] = int(item.get("new_trigger_count") or 0)
        item["safety_trigger_status"] = (
            "Detected" if item["safety_trigger_count"] else "Not detected"
        )
        documents.append(item)

    return {
        "records": documents,
        "page": page,
        "page_size": page_size,
        "total": total,
        "stats": {
            "new_documents": sum(int(row["document_count"] or 0) for row in stats_rows),
            "potential_safety_triggers": potential_triggers,
            "products_mentioned": len(products),
            "sources": len(sources),
            "source_names": list(dict.fromkeys(sources.values())),
            "this_week": recent_count,
            "last_week": previous_count,
        },
        "filters": {
            "sources": [dict(row) for row in source_rows],
            "content_types": [row["content_type"] for row in content_type_rows],
            "product_classes": [dict(row) for row in product_classes],
            "safety_triggers": [row["trigger_name"] for row in trigger_rows],
        },
    }


@app.get("/api/new-information/{document_id}")
def get_new_information_document(document_id: int):
    with get_db() as conn:
        document = conn.execute(
            """
            SELECT d.*, s.name AS source_name,
                   s.reference_no AS source_reference_no,
                   (
                       SELECT COUNT(*) FROM findings f
                       WHERE f.document_id = d.id
                   ) AS findings_count,
                   (
                       SELECT GROUP_CONCAT(DISTINCT COALESCE(
                           NULLIF(trim(f.product_name), ''),
                           NULLIF(trim(f.active_substance), '')
                       ))
                       FROM findings f
                       WHERE f.document_id = d.id
                         AND COALESCE(
                             NULLIF(trim(f.product_name), ''),
                             NULLIF(trim(f.active_substance), '')
                         ) IS NOT NULL
                   ) AS products_found
            FROM documents d
            JOIN sources s ON s.id = d.source_id
            WHERE d.id = ?
              AND EXISTS (
                  SELECT 1 FROM findings fn
                  WHERE fn.document_id = d.id
                    AND UPPER(COALESCE(fn.comparison_status, ''))
                        IN ('NEW', 'NEW_INFORMATION')
              )
            """,
            (document_id,),
        ).fetchone()
        if not document:
            raise HTTPException(status_code=404, detail="New information document not found.")
        findings = conn.execute(
            """
            SELECT f.id, f.product_name, f.active_substance, f.trigger_name,
                   f.safety_topic, f.evidence, f.page_number, f.created_at,
                   s.summary, s.prac_assessment, s.outcome
            FROM findings f
            LEFT JOIN summaries s ON s.id = (
                SELECT ss.id FROM summaries ss
                WHERE ss.finding_id = f.id
                ORDER BY ss.created_at DESC, ss.id DESC LIMIT 1
            )
            WHERE f.document_id = ?
              AND UPPER(COALESCE(f.comparison_status, ''))
                  IN ('NEW', 'NEW_INFORMATION')
            ORDER BY f.page_number, f.id
            """,
            (document_id,),
        ).fetchall()
        ai_analysis = conn.execute(
            """
            SELECT summary, status, provider, updated_at
            FROM document_ai_analyses
            WHERE document_id = ? AND analysis_type = 'summary'
              AND provider IN ('ollama', 'openai')
            ORDER BY updated_at DESC, id DESC
            LIMIT 1
            """,
            (document_id,),
        ).fetchone()
        configured_products = conn.execute(
            """
            SELECT p.product_name, pc.class_name
            FROM products p
            JOIN product_classes pc ON pc.id = p.product_class_id
            WHERE p.active = 1 AND pc.active = 1
            ORDER BY p.product_name
            """
        ).fetchall()

    item = dict(document)
    item["reference_id"] = item.get("reference_id") or _document_reference_id(
        item.get("source_reference_no"),
        item.get("publication_date"),
        int(item["id"]),
    )
    item["date_crawled"] = item.get("last_crawled_at") or item.get("discovered_date")
    item["crawled_status"] = (
        "Yes"
        if item["date_crawled"]
        or str(item.get("status") or "").upper()
            in {"PROCESSED", "EXTRACTED", "ALREADY_CRAWLED"}
        else "No"
    )
    item["working_status"] = (
        "Yes"
        if str(item.get("status") or "").upper() in {"PROCESSED", "EXTRACTED"}
        else "No"
    )
    item["products_found"] = _distinct_finding_products(item.get("products_found"))
    product_map = {
        str(row["product_name"]).casefold(): row["class_name"]
        for row in configured_products
    }
    item["product_classes"] = list(dict.fromkeys(
        product_map[product.casefold()]
        for product in item["products_found"]
        if product.casefold() in product_map
    ))
    finding_items = [dict(row) for row in findings]
    item["safety_trigger_count"] = len(finding_items)
    item["safety_triggers"] = [
        {
            "trigger_name": finding["trigger_name"],
            "safety_topic": finding["safety_topic"],
            "evidence": finding["evidence"],
            "page_number": finding["page_number"],
        }
        for finding in finding_items
    ]
    summaries = []
    for finding in finding_items:
        text = finding.get("summary") or finding.get("prac_assessment") or finding.get("outcome")
        if text and text not in summaries:
            summaries.append(text)
    if ai_analysis and str(ai_analysis["status"]).upper() == "COMPLETED" and ai_analysis["summary"]:
        item["safety_information_summary"] = ai_analysis["summary"]
    else:
        item["safety_information_summary"] = "\n".join(summaries) or None
    return item


@app.get("/api/findings/{finding_id}")
def get_finding(finding_id: int):
    with get_db() as conn:
        row = conn.execute(
            """
            SELECT
                f.*,
                d.title AS document_title,
                d.filename,
                d.url AS document_url
            FROM findings f
            LEFT JOIN documents d
                ON f.document_id = d.id
            WHERE f.id = ?
            """,
            (finding_id,),
        ).fetchone()
        summary_row = conn.execute(
            """
            SELECT *
            FROM summaries
            WHERE finding_id = ?
            ORDER BY created_at DESC
            LIMIT 1
            """,
            (finding_id,),
        ).fetchone()

        page_row = conn.execute(
            """
            SELECT
                p.page_text,
                d.title AS document_title,
                d.url AS document_url,
                d.filename
            FROM findings f
            LEFT JOIN document_pages p
                ON p.document_id = f.document_id
               AND p.page_number = f.page_number
            LEFT JOIN documents d
                ON d.id = f.document_id
            WHERE f.id = ?
            """,
            (finding_id,),
        ).fetchone()

    if not row:
        raise HTTPException(
            status_code=404,
            detail="Finding not found.",
        )

    result = dict(row)
    for field in (
        "relevant_terms",
        "section",
        "evidence_hash",
    ):
        result.setdefault(field, None)
    result["summary"] = dict(summary_row) if summary_row else None
    if result["summary"]:
        result["summary"]["model"] = os.getenv(
            "OLLAMA_MODEL",
            "llama3.2",
        )
        result["summary"]["status"] = "GENERATED"
    with get_db() as conn:
        review_row = conn.execute(
            """
            SELECT id, reviewer, decision, comments, reviewed_at
            FROM reviews
            WHERE finding_id = ?
            ORDER BY reviewed_at DESC, id DESC
            LIMIT 1
            """,
            (finding_id,),
        ).fetchone()
    result["latest_review"] = dict(review_row) if review_row else None
    result["review_status"] = "REVIEWED" if review_row else "PENDING"

    page_text = page_row["page_text"] if page_row else None
    result["evidence_detail"] = {
        "exact_text": result.get("evidence"),
        "evidence_hash": result.get("evidence_hash"),
        "document_id": result.get("document_id"),
        "page_number": result.get("page_number"),
        "page_text": page_text,
        "document_title": (
            page_row["document_title"]
            if page_row
            else result.get("document_title")
        ),
        "document_url": (
            page_row["document_url"]
            if page_row
            else result.get("document_url")
        ),
        "filename": (
            page_row["filename"]
            if page_row
            else result.get("filename")
        ),
        "warning": (
            None
            if page_text is not None
            else "The exact source page is unavailable."
        ),
    }

    return result


# ---------------------------------------------------------
# Summaries
# ---------------------------------------------------------

@app.get("/api/summaries")
def list_summaries(
    limit: int = Query(100, ge=1, le=1000),
):
    with get_db() as conn:
        rows = conn.execute(
            """
            SELECT *
            FROM summaries
            ORDER BY created_at DESC
            LIMIT ?
            """,
            (limit,),
        ).fetchall()

    summaries = []
    for row in rows:
        summary = dict(row)
        summary["model"] = summary.get("model") or os.getenv("OLLAMA_MODEL", "llama3.2")
        summary["status"] = "GENERATED"
        summaries.append(summary)

    return summaries


def _configured_product_class_matches(conn, finding):
    return _matching_configured_product_classes(
        finding,
        _configured_product_class_rows(conn),
    )


@app.get("/api/findings/export.xlsx")
def export_topic_assessments():
    with get_db() as conn:
        rows = conn.execute(
            """
            SELECT
                f.*,
                d.title AS document_title,
                d.url AS document_url,
                d.content_type,
                d.publication_date,
                d.last_crawled_at,
                d.reference_id,
                d.source_id,
                s.name AS source_name,
                s.instruction AS source_instruction,
                summary.summary AS topic_summary,
                summary.prac_assessment,
                summary.outcome,
                summary.model AS summary_model,
                summary.created_at AS summary_created_at,
                (
                    SELECT COUNT(*) FROM extraction_history eh
                    WHERE eh.document_id = d.id
                      AND UPPER(COALESCE(eh.status, '')) = 'SUCCESS'
                ) AS crawl_count
            FROM findings f
            LEFT JOIN documents d ON d.id = f.document_id
            LEFT JOIN sources s ON s.id = d.source_id
            LEFT JOIN summaries summary ON summary.id = (
                SELECT id FROM summaries
                WHERE finding_id = f.id
                ORDER BY created_at DESC, id DESC
                LIMIT 1
            )
            ORDER BY f.created_at DESC, f.id DESC
            """
        ).fetchall()
        review_rows = conn.execute(
            """
            SELECT r.*, f.safety_topic, f.product_name,
                   d.title AS document_title, d.reference_id
            FROM reviews r
            LEFT JOIN findings f ON f.id = r.finding_id
            LEFT JOIN documents d ON d.id = f.document_id
            ORDER BY r.finding_id, r.stage, r.reviewed_at, r.id
            """
        ).fetchall()
        configured_products = _configured_product_class_rows(conn)
        product_matches = {
            int(row["id"]): _matching_configured_product_classes(
                dict(row),
                configured_products,
            )
            for row in rows
        }

    latest_reviews = {}
    for row in review_rows:
        latest_reviews.setdefault(int(row["finding_id"]), {})[row["stage"]] = dict(row)

    workbook = Workbook()
    topics = workbook.active
    topics.title = "Topic Assessments"
    topic_headers = [
        "Finding ID", "Published Date", "Reference", "Document Title",
        "Content Type", "Source", "Document URL", "Date Crawled", "Crawl Count",
        "Product", "Active Substance", "Configured Product Class",
        "Safety Topic", "Exact Evidence", "Concise Topic Summary",
        "Product Impact / PI-SmPC Assessment", "Proposed Outcome",
        "Summary Model", "Summary Date", "Source Instruction",
        "Initial Review Status", "Initial Reviewer", "Initial Decision",
        "Initial Comments", "Initial Review Date", "Final Review Status",
        "Final Reviewer", "Final Decision", "Final Comments", "Final Review Date",
    ]
    topics.append(topic_headers)

    def excel_values(values):
        return [
            "'" + value if isinstance(value, str) and value.startswith(("=", "+", "-", "@")) else value
            for value in values
        ]

    for row in rows:
        finding_id = int(row["id"])
        reviews = latest_reviews.get(finding_id, {})
        initial = reviews.get("INITIAL", {})
        final = reviews.get("FINAL", {})
        matches = product_matches.get(finding_id, [])
        topics.append(excel_values([
            finding_id,
            row["publication_date"],
            row["reference_id"],
            row["document_title"],
            row["content_type"],
            row["source_name"],
            row["document_url"],
            row["last_crawled_at"],
            row["crawl_count"] or 0,
            row["product_name"],
            row["active_substance"],
            "; ".join(sorted({match["class_name"] for match in matches})),
            row["safety_topic"] or row["trigger_name"],
            row["evidence"],
            row["topic_summary"],
            row["prac_assessment"],
            row["outcome"],
            row["summary_model"],
            row["summary_created_at"],
            row["source_instruction"],
            "Available" if initial else "Pending",
            initial.get("reviewer"),
            initial.get("decision"),
            initial.get("comments"),
            initial.get("reviewed_at"),
            "Available" if final else "Pending",
            final.get("reviewer"),
            final.get("decision"),
            final.get("comments"),
            final.get("reviewed_at"),
        ]))

    history = workbook.create_sheet("Review History")
    history.append(excel_values([
        "Finding ID", "Reference", "Document Title", "Safety Topic",
        "Stage", "Reviewer", "Decision", "Comments", "Reviewed At",
    ]))
    for row in review_rows:
        history.append([
            row["finding_id"],
            row["reference_id"],
            row["document_title"],
            row["safety_topic"],
            row["stage"],
            row["reviewer"],
            row["decision"],
            row["comments"],
            row["reviewed_at"],
        ])

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            header_font = copy(cell.font)
            header_font.bold = True
            cell.font = header_font
        for column in sheet.columns:
            width = min(
                48,
                max(12, max(len(str(cell.value or "")) for cell in column[:50]) + 2),
            )
            sheet.column_dimensions[column[0].column_letter].width = width

    output = BytesIO()
    workbook.save(output)
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": 'attachment; filename="topic-assessments.xlsx"'
        },
    )


@app.post("/api/summaries/generate/{finding_id}")
def generate_summary(
    finding_id: int,
    provider: Optional[str] = Query(None),
):

    with get_db() as conn:
        row = conn.execute(
            """
            SELECT * FROM findings WHERE id = ?
            """,
            (finding_id,),
        ).fetchone()
        configured_matches = (
            _configured_product_class_matches(conn, dict(row))
            if row
            else []
        )

    if not row:
        raise HTTPException(
            status_code=404,
            detail="Finding not found.",
        )

    finding = dict(row)
    finding["configured_product_classes"] = configured_matches
    finding["product_class_context"] = (
        "; ".join(
            f"{item['product_name']} → {item['class_name']}"
            for item in configured_matches
        )
        or "No exact configured product match"
    )
    selected_provider = (provider or "").strip().lower()
    if not selected_provider:
        selected_provider = (
            "openai"
            if openai_provider_status()["available"]
            else "ollama"
        )
    if selected_provider not in {"ollama", "openai"}:
        raise HTTPException(status_code=400, detail="Unsupported AI provider.")

    evidence = str(finding.get("evidence") or "").strip()
    if not evidence:
        return {
            "success": False,
            "status": "NO_EVIDENCE",
            "message": "The finding has no exact evidence to summarize.",
        }

    try:
        result = (
            generate_openai_finding_summary(finding)
            if selected_provider == "openai"
            else generate_summary_from_finding(finding)
        )

    except Exception as exc:
        logger.exception("%s finding summary generation failed", selected_provider)

        return {
            "success": False,
            "status": "ERROR",
            "message": "AI summary generation failed.",
        }

    if not result:
        return {
            "success": False,
            "status": "UNAVAILABLE",
            "message": "AI summary was not generated.",
        }

    if not isinstance(result, dict):
        return {
            "success": False,
            "status": "INVALID_RESPONSE",
            "message": "Ollama returned an invalid response.",
        }

    status = result.get("status", "ERROR")
    model = result.get("model") or (
        os.getenv("OPENAI_MODEL", "gpt-4o-mini")
        if selected_provider == "openai"
        else os.getenv("OLLAMA_MODEL", "llama3.2")
    )
    if status != "SUCCESS":
        return {
            "success": False,
            "status": status,
            "model": model,
            "message": result.get("error") or "AI summary was not generated.",
        }

    summary_text = result.get("summary")
    prac_assessment = result.get("prac_assessment")
    outcome = result.get("outcome")

    if not isinstance(summary_text, str) or not summary_text.strip():
        return {
            "success": False,
            "status": "INVALID_RESPONSE",
            "model": model,
            "message": "AI provider returned no usable summary.",
        }

    if not all(
        value is None or isinstance(value, str)
        for value in (
            result.get("product_name"),
            result.get("active_substance"),
            result.get("safety_topic"),
            prac_assessment,
            outcome,
            summary_text,
        )
    ):
        return {
            "success": False,
            "status": "INVALID_RESPONSE",
            "model": model,
            "message": "AI provider returned fields with invalid types.",
        }

    now = datetime.utcnow().isoformat()

    with get_db() as conn:
        conn.execute(
            """
            INSERT INTO summaries
            (
                finding_id,
                product_name,
                active_substance,
                safety_topic,
                prac_assessment,
                outcome,
                summary,
                model,
                created_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                finding_id,
                result.get("product_name") or finding.get("product_name"),
                result.get("active_substance") or finding.get("active_substance"),
                result.get("safety_topic") or finding.get("safety_topic"),
                prac_assessment,
                outcome,
                summary_text.strip(),
                model,
                now,
            ),
        )

    return {
        "success": True,
        "status": "GENERATED",
        "summary": summary_text.strip(),
        "model": model,
        "created_at": now,
        "prac_assessment": prac_assessment,
        "outcome": outcome,
    }


# ---------------------------------------------------------
# Ollama
# ---------------------------------------------------------

@app.get("/api/ollama/test")
def ollama_test():
    try:
        return test_ollama_connection()

    except Exception as exc:
        logger.exception("Ollama connection test failed")

        return {
            "available": False,
            "model_available": False,
            "message": str(exc),
            "error": str(exc),
        }


# ---------------------------------------------------------
# Reviews
# ---------------------------------------------------------

@app.post("/api/reviews")
def create_review(review: ReviewCreate):

    now = datetime.utcnow().isoformat()

    with get_db() as conn:

        finding = conn.execute(
            """
            SELECT id
            FROM findings
            WHERE id = ?
            """,
            (review.finding_id,),
        ).fetchone()

        if not finding:
            raise HTTPException(
                status_code=404,
                detail="Finding not found.",
            )
        if review.stage == "FINAL":
            initial_review = conn.execute(
                """
                SELECT id FROM reviews
                WHERE finding_id = ? AND stage = 'INITIAL'
                LIMIT 1
                """,
                (review.finding_id,),
            ).fetchone()
            if not initial_review:
                raise HTTPException(
                    status_code=409,
                    detail="Complete the initial review before the final review.",
                )

        cursor = conn.execute(
            """
            INSERT INTO reviews
            (
                finding_id,
                stage,
                reviewer,
                decision,
                comments,
                reviewed_at
            )
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            (
                review.finding_id,
                review.stage,
                review.reviewer,
                review.decision,
                review.comments,
                now,
            ),
        )

        review_id = cursor.lastrowid

    return {
        "success": True,
        "id": review_id,
        "message": "Review recorded successfully.",
    }


@app.get("/api/reviews")
def list_reviews():
    with get_db() as conn:
        rows = conn.execute(
            """
            SELECT
                r.*,
                f.product_name,
                f.active_substance,
                f.safety_topic,
                f.page_number,
                d.title AS document_title,
                d.filename
            FROM reviews r
            LEFT JOIN findings f
                ON r.finding_id = f.id
            LEFT JOIN documents d
                ON f.document_id = d.id
            ORDER BY r.reviewed_at DESC
            """
        ).fetchall()

    return [dict(row) for row in rows]


@app.put("/api/reviews/{review_id}")
def update_review(review_id: int, review: ReviewUpdate):
    now = datetime.utcnow().isoformat()

    with get_db() as conn:
        existing = conn.execute(
            "SELECT id, finding_id, stage FROM reviews WHERE id = ?",
            (review_id,),
        ).fetchone()
        if not existing:
            raise HTTPException(status_code=404, detail="Review not found.")

        if review.stage != existing["stage"]:
            raise HTTPException(
                status_code=409,
                detail="A review's stage cannot be changed after it is recorded.",
            )

        if review.stage == "FINAL":
            initial_review = conn.execute(
                """
                SELECT id FROM reviews
                WHERE finding_id = ? AND stage = 'INITIAL'
                LIMIT 1
                """,
                (existing["finding_id"],),
            ).fetchone()
            if not initial_review:
                raise HTTPException(
                    status_code=409,
                    detail="Complete the initial review before the final review.",
                )
        elif existing["stage"] == "INITIAL":
            final_review = conn.execute(
                """
                SELECT id FROM reviews
                WHERE finding_id = ? AND stage = 'FINAL'
                LIMIT 1
                """,
                (existing["finding_id"],),
            ).fetchone()
            if final_review:
                raise HTTPException(
                    status_code=409,
                    detail="A completed final review prevents changes to the initial review.",
                )

        cursor = conn.execute(
            """
            UPDATE reviews
            SET stage = ?, reviewer = ?, decision = ?, comments = ?, reviewed_at = ?
            WHERE id = ?
            """,
            (
                review.stage,
                review.reviewer,
                review.decision,
                review.comments,
                now,
                review_id,
            ),
        )

    return {
        "success": True,
        "message": "Review updated successfully.",
    }


@app.delete("/api/reviews/{review_id}")
def delete_review(review_id: int):
    with get_db() as conn:
        review = conn.execute(
            "SELECT finding_id, stage FROM reviews WHERE id = ?",
            (review_id,),
        ).fetchone()
        if not review:
            raise HTTPException(status_code=404, detail="Review not found.")
        if review["stage"] == "INITIAL":
            final_review = conn.execute(
                """
                SELECT id FROM reviews
                WHERE finding_id = ? AND stage = 'FINAL'
                LIMIT 1
                """,
                (review["finding_id"],),
            ).fetchone()
            if final_review:
                raise HTTPException(
                    status_code=409,
                    detail="A completed final review prevents deleting the initial review.",
                )
        cursor = conn.execute(
            "DELETE FROM reviews WHERE id = ?",
            (review_id,),
        )

    return {
        "success": True,
        "message": "Review deleted successfully.",
    }

def create_document_summary_pdf(document_id: int) -> Path:
    """
    Create a traceable review-support PDF for a PRAC document.

    The report uses only information already stored by the POC.
    It includes document dates, extracted page count, findings,
    evidence, comparison status, deterministic topic summaries,
    crawler XPath configuration, and reviewer decisions.
    """

    with get_db() as conn:

        document = conn.execute(
            """
            SELECT
                d.*,
                s.authority,
                s.name AS source_name,
                s.source_type,
                s.active AS source_active,
                s.url AS source_url,
                s.instruction
            FROM documents d
            LEFT JOIN sources s
                ON d.source_id = s.id
            WHERE d.id = ?
            """,
            (document_id,),
        ).fetchone()

        if not document:
            raise ValueError("Document not found.")

        findings = conn.execute(
            """
            SELECT *
            FROM findings
            WHERE document_id = ?
            ORDER BY page_number, id
            """,
            (document_id,),
        ).fetchall()

        summaries = conn.execute(
            """
            SELECT
                s.*,
                f.page_number,
                f.evidence,
                f.comparison_status
            FROM summaries s
            LEFT JOIN findings f
                ON s.finding_id = f.id
            WHERE f.document_id = ?
            ORDER BY f.page_number, s.id
            """,
            (document_id,),
        ).fetchall()

        reviews = conn.execute(
            """
            SELECT
                r.*,
                f.page_number,
                f.product_name,
                f.active_substance,
                f.safety_topic
            FROM reviews r
            LEFT JOIN findings f
                ON r.finding_id = f.id
            WHERE f.document_id = ?
            ORDER BY r.reviewed_at DESC
            """,
            (document_id,),
        ).fetchall()

        page_count = conn.execute(
            """
            SELECT COUNT(*) AS count
            FROM document_pages
            WHERE document_id = ?
            """,
            (document_id,),
        ).fetchone()["count"]

        new_information_count = conn.execute(
            """
            SELECT COUNT(*) AS count
            FROM findings
            WHERE document_id = ?
              AND UPPER(COALESCE(comparison_status, ''))
                  IN ('NEW', 'NEW_INFORMATION')
            """,
            (document_id,),
        ).fetchone()["count"]

        topic_pages = conn.execute(
            """
            SELECT page_number, page_text
            FROM document_pages
            WHERE document_id = ?
            ORDER BY page_number
            """,
            (document_id,),
        ).fetchall()
        topics = extract_meeting_topics(
            [dict(page) for page in topic_pages]
        )
        section_summaries = extract_section_summaries(
            [dict(page) for page in topic_pages],
            topics,
        )
        summary_points = _extractive_summary_sentences(
            [dict(page) for page in topic_pages],
            limit=5,
        )
        extraction = conn.execute(
            """
            SELECT started_at, completed_at, status, pages_extracted
            FROM extraction_history
            WHERE document_id = ?
            ORDER BY id DESC LIMIT 1
            """,
            (document_id,),
        ).fetchone()
        crawl = conn.execute(
            """
            SELECT started_at, completed_at, status, documents_found
            FROM crawl_history
            WHERE source_id = ?
            ORDER BY id DESC LIMIT 1
            """,
            (document["source_id"],),
        ).fetchone()

    title = (
        document["title"]
        or document["filename"]
        or "Published Document"
    )
    content_type = str(document["content_type"] or "Published Document")

    safe_title = re.sub(
        r"[^A-Za-z0-9_-]+",
        "_",
        title,
    ).strip("_")

    if not safe_title:
        safe_title = "Published_Document"

    output_path = (
        SUMMARY_DIR
        / f"summary_{document_id}_{safe_title}.pdf"
    )

    styles = getSampleStyleSheet()

    styles.add(
        ParagraphStyle(
            name="ReportTitle",
            parent=styles["Title"],
            fontName="Helvetica-Bold",
            fontSize=18,
            leading=22,
            spaceAfter=12,
            textColor=colors.HexColor("#10213f"),
        )
    )

    styles.add(
        ParagraphStyle(
            name="SectionTitle",
            parent=styles["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=12,
            leading=15,
            spaceBefore=12,
            spaceAfter=8,
            textColor=colors.HexColor("#10213f"),
        )
    )

    styles.add(
        ParagraphStyle(
            name="BodySmall",
            parent=styles["BodyText"],
            fontName="Helvetica",
            fontSize=8.5,
            leading=12,
            spaceAfter=5,
        )
    )

    styles.add(
        ParagraphStyle(
            name="BodySmallBold",
            parent=styles["BodySmall"],
            fontName="Helvetica-Bold",
        )
    )

    styles.add(
        ParagraphStyle(
            name="Evidence",
            parent=styles["BodySmall"],
            fontName="Helvetica",
            fontSize=8,
            leading=11,
            leftIndent=8,
            spaceAfter=5,
        )
    )

    doc = SimpleDocTemplate(
        str(output_path),
        pagesize=A4,
        rightMargin=16 * mm,
        leftMargin=16 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
        title=f"{content_type} Summary - {title}",
        author="Safety Trigger Monitoring POC",
    )

    story = []

    story.append(
        Paragraph(
            "SAFETY TRIGGER MONITORING",
            styles["BodySmallBold"],
        )
    )

    story.append(
        Paragraph(
            f"{content_type} Summary",
            styles["ReportTitle"],
        )
    )

    story.append(
        Paragraph(
            "MEETING SUMMARY" if content_type == "Meeting Minutes" else "DOCUMENT SUMMARY",
            styles["SectionTitle"],
        )
    )

    metadata = [
        ["Document", str(title)],
        ["Reference ID", str(document["reference_id"] or "Not available")],
        ["Content type", content_type],
        ["Source instruction", str(document["instruction"] or "Not available")],
        [
            "Health authority",
            str(
                document["authority"]
                or "European Medicines Agency"
            ),
        ],
        [
            "Source",
            str(
                document["source_name"]
                or "EMA PRAC"
            ),
        ],
        [
            "Meeting date",
            str(
                document["meeting_date"]
                or "Not extracted"
            ),
        ],
        [
            "Publication date",
            str(
                document["publication_date"]
                or "Not available"
            ),
        ],
        [
            "PDF filename",
            str(
                document["filename"]
                or "Not available"
            ),
        ],
        [
            "Pages extracted",
            str(page_count),
        ],
        [
            "Safety findings",
            str(len(findings)),
        ],
        [
            "NEW INFORMATION",
            str(new_information_count),
        ],
    ]

    metadata_table = Table(
        metadata,
        colWidths=[42 * mm, 130 * mm],
        repeatRows=0,
    )

    metadata_table.setStyle(
        TableStyle(
            [
                (
                    "FONTNAME",
                    (0, 0),
                    (0, -1),
                    "Helvetica-Bold",
                ),
                (
                    "FONTNAME",
                    (1, 0),
                    (1, -1),
                    "Helvetica",
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    8,
                ),
                (
                    "LEADING",
                    (0, 0),
                    (-1, -1),
                    11,
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "TOP",
                ),
                (
                    "BACKGROUND",
                    (0, 0),
                    (0, -1),
                    colors.HexColor("#eef3f9"),
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.25,
                    colors.HexColor("#d4dce8"),
                ),
                (
                    "LEFTPADDING",
                    (0, 0),
                    (-1, -1),
                    6,
                ),
                (
                    "RIGHTPADDING",
                    (0, 0),
                    (-1, -1),
                    6,
                ),
                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    5,
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    5,
                ),
            ]
        )
    )

    story.append(metadata_table)
    story.append(Spacer(1, 8))

    history_data = [
        ["Crawl started", str(crawl["started_at"] if crawl else "Not available")],
        ["Crawl completed", str(crawl["completed_at"] if crawl else "Not available")],
        ["Processing started", str(extraction["started_at"] if extraction else "Not available")],
        ["Processing completed", str(extraction["completed_at"] if extraction else "Not available")],
        ["Crawl status", str(crawl["status"] if crawl else "Not available")],
        ["Processing status", str(extraction["status"] if extraction else "Not available")],
        ["Topics extracted", str(len(topics))],
        ["Summary method", "Document Extraction / POC Summary (deterministic)"],
    ]
    history_table = Table(history_data, colWidths=[42 * mm, 130 * mm])
    history_table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#eef3f9")),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#d4dce8")),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(history_table)
    story.append(Spacer(1, 8))

    executive_summary = _document_executive_summary(
        content_type,
        section_summaries,
        summary_points,
    )
    story.append(Paragraph("EXECUTIVE SUMMARY", styles["SectionTitle"]))
    story.append(Paragraph(
        html_escape(executive_summary),
        styles["BodySmall"],
    ))
    story.append(Paragraph(
        "KEY MEETING TOPICS" if content_type == "Meeting Minutes" else "KEY TOPICS",
        styles["SectionTitle"],
    ))
    for section in section_summaries:
        story.append(Paragraph(
            f"{html_escape(str(section.get('section_number') or ''))} "
            f"{html_escape(str(section.get('title') or ''))}",
            styles["BodySmallBold"],
        ))
        story.append(Paragraph(
            html_escape(str(section.get("summary") or "")),
            styles["BodySmall"],
        ))
        story.append(Paragraph(
            f"Source pages: {section.get('source_page_start')}–"
            f"{section.get('source_page_end')}",
            styles["BodySmall"],
        ))

    story.append(Paragraph(
        "DETAILED MEETING TOPICS" if content_type == "Meeting Minutes" else "EXTRACTED DOCUMENT TOPICS",
        styles["SectionTitle"],
    ))
    story.append(Paragraph(
        "Topics and key points were extracted deterministically from the PDF text. "
        "This is not an AI-generated regulatory conclusion.",
        styles["BodySmall"],
    ))
    if topics:
        for topic in topics:
            title_text = html_escape(str(topic.get("title") or "Untitled topic"))
            page_text = html_escape(str(topic.get("page_number") or "—"))
            source_text = topic.get("source_text")
            detail = (
                f"<br/><i>{html_escape(str(source_text))}</i>"
                if source_text and source_text != topic.get("title")
                else ""
            )
            story.append(Paragraph(
                f"{html_escape(str(topic.get('section_number') or ''))} "
                f"{title_text} — Source page: {page_text}{detail}",
                styles["BodySmall"],
            ))
    else:
        for point in summary_points:
            story.append(Paragraph(f"• {html_escape(point)}", styles["BodySmall"]))
        if not summary_points:
            story.append(Paragraph(
                "No topics or key points could be extracted from this document.",
                styles["BodySmall"],
            ))

    story.append(
        Paragraph(
            "Source URL",
            styles["SectionTitle"],
        )
    )

    story.append(
        Paragraph(
            str(document["url"]),
            styles["BodySmall"],
        )
    )

    story.append(
        Paragraph(
            "COMPARISON RESULT",
            styles["SectionTitle"],
        )
    )

    if new_information_count > 0:
        result_text = (
            f"NEW INFORMATION detected: "
            f"{new_information_count} finding(s) were classified as NEW "
            f"after comparison with existing findings."
        )
    elif findings:
        result_text = (
            "No NEW INFORMATION was identified by the comparison "
            "step for the detected findings."
        )
    else:
        result_text = (
            "No trigger-relevant safety findings were detected."
        )

    story.append(
        Paragraph(
            result_text,
            styles["BodySmall"],
        )
    )

    story.append(
        Paragraph(
            "DETECTED SAFETY INFORMATION",
            styles["SectionTitle"],
        )
    )

    if not findings:

        story.append(
            Paragraph(
                "No trigger-relevant findings are stored for this document.",
                styles["BodySmall"],
            )
        )

    else:

        for index, finding in enumerate(
            findings,
            start=1,
        ):

            finding_data = [
                [
                    Paragraph(
                        f"Finding {index}",
                        styles["BodySmallBold"],
                    ),
                    "",
                ],
                [
                    "Product",
                    str(
                        finding["product_name"]
                        or "Not identified"
                    ),
                ],
                [
                    "Active substance",
                    str(
                        finding["active_substance"]
                        or "Not identified"
                    ),
                ],
                [
                    "Safety topic",
                    str(
                        finding["safety_topic"]
                        or "Not identified"
                    ),
                ],
                [
                    "Page",
                    str(
                        finding["page_number"]
                        or "—"
                    ),
                ],
                [
                    "Comparison",
                    str(
                        finding["comparison_status"]
                        or "Not compared"
                    ),
                ],
                [
                    "Relevant terms",
                    str(
                        (
                            finding["relevant_terms"]
                            if "relevant_terms" in finding.keys()
                            else None
                        )
                        or "—"
                    ),
                ],
            ]

            finding_table = Table(
                finding_data,
                colWidths=[42 * mm, 130 * mm],
            )

            finding_table.setStyle(
                TableStyle(
                    [
                        (
                            "SPAN",
                            (0, 0),
                            (1, 0),
                        ),
                        (
                            "FONTNAME",
                            (0, 0),
                            (-1, 0),
                            "Helvetica-Bold",
                        ),
                        (
                            "FONTNAME",
                            (0, 1),
                            (0, -1),
                            "Helvetica-Bold",
                        ),
                        (
                            "FONTSIZE",
                            (0, 0),
                            (-1, -1),
                            8,
                        ),
                        (
                            "VALIGN",
                            (0, 0),
                            (-1, -1),
                            "TOP",
                        ),
                        (
                            "BACKGROUND",
                            (0, 0),
                            (-1, 0),
                            colors.HexColor("#eef3f9"),
                        ),
                        (
                            "GRID",
                            (0, 0),
                            (-1, -1),
                            0.25,
                            colors.HexColor("#d4dce8"),
                        ),
                        (
                            "LEFTPADDING",
                            (0, 0),
                            (-1, -1),
                            6,
                        ),
                        (
                            "RIGHTPADDING",
                            (0, 0),
                            (-1, -1),
                            6,
                        ),
                        (
                            "TOPPADDING",
                            (0, 0),
                            (-1, -1),
                            4,
                        ),
                        (
                            "BOTTOMPADDING",
                            (0, 0),
                            (-1, -1),
                            4,
                        ),
                    ]
                )
            )

            story.append(finding_table)
            story.append(Spacer(1, 5))

            story.append(
                Paragraph(
                    "Exact evidence:",
                    styles["BodySmallBold"],
                )
            )

            evidence = (
                finding["evidence"]
                or "No evidence text available."
            )

            evidence_text = str(evidence).replace(
                "\n",
                "<br/>",
            )

            story.append(
                Paragraph(
                    evidence_text,
                    styles["Evidence"],
                )
            )

    if summaries:

        story.append(
            Paragraph(
                "OLLAMA EVIDENCE SUMMARIES",
                styles["SectionTitle"],
            )
        )

        for index, summary in enumerate(
            summaries,
            start=1,
        ):

            story.append(
                Paragraph(
                    (
                        f"Summary {index} "
                        f"(Page {summary['page_number'] or '—'})"
                    ),
                    styles["BodySmallBold"],
                )
            )

            summary_text = (
                summary["summary"]
                or "No summary text available."
            )

            story.append(
                Paragraph(
                    str(summary_text).replace(
                        "\n",
                        "<br/>",
                    ),
                    styles["BodySmall"],
                )
            )

            story.append(
                Paragraph(
                    (
                        f"Model: "
                        f"{summary['model'] or 'Not available'} | "
                        f"Status: "
                        f"{summary['status'] or '—'}"
                    ),
                    styles["BodySmall"],
                )
            )

    story.append(
        Paragraph(
            "REVIEW HISTORY",
            styles["SectionTitle"],
        )
    )

    if reviews:

        review_data = [
            [
                "Reviewed",
                "Reviewer",
                "Decision",
                "Page",
                "Comments",
            ]
        ]

        for review in reviews:

            review_data.append(
                [
                    str(
                        review["reviewed_at"]
                        or "—"
                    ),
                    str(
                        review["reviewer"]
                        or "—"
                    ),
                    str(
                        review["decision"]
                        or "—"
                    ),
                    str(
                        review["page_number"]
                        or "—"
                    ),
                    str(
                        review["comments"]
                        or "—"
                    ),
                ]
            )

        review_table = Table(
            review_data,
            colWidths=[
                28 * mm,
                28 * mm,
                28 * mm,
                15 * mm,
                71 * mm,
                ],
            repeatRows=1,
        )

        review_table.setStyle(
            TableStyle(
                [
                    (
                        "BACKGROUND",
                        (0, 0),
                        (-1, 0),
                        colors.HexColor("#eef3f9"),
                    ),
                    (
                        "FONTNAME",
                        (0, 0),
                        (-1, 0),
                        "Helvetica-Bold",
                    ),
                    (
                        "FONTSIZE",
                        (0, 0),
                        (-1, -1),
                        7,
                    ),
                    (
                        "GRID",
                        (0, 0),
                        (-1, -1),
                        0.25,
                        colors.HexColor("#d4dce8"),
                    ),
                    (
                        "VALIGN",
                        (0, 0),
                        (-1, -1),
                        "TOP",
                    ),
                    (
                        "LEFTPADDING",
                        (0, 0),
                        (-1, -1),
                        4,
                    ),
                    (
                        "RIGHTPADDING",
                        (0, 0),
                        (-1, -1),
                        4,
                    ),
                    (
                        "TOPPADDING",
                        (0, 0),
                        (-1, -1),
                        4,
                    ),
                    (
                        "BOTTOMPADDING",
                        (0, 0),
                        (-1, -1),
                        4,
                    ),
                ]
            )
        )

        story.append(review_table)

    else:

        story.append(
            Paragraph(
                "No human review has been recorded for this document.",
                styles["BodySmall"],
            )
        )

    story.append(
        Paragraph(
            "CONFIGURATION USED",
            styles["SectionTitle"],
        )
    )
    trigger = load_trigger_config()
    configuration_data = [
        ["Health authority", str(document["authority"] or "Not available")],
        ["Source name", str(document["source_name"] or "Not available")],
        ["Source URL", str(document["source_url"] or "Not available")],
        ["Source type", str(document["source_type"] or "Not available")],
        ["Active status", "Active" if document["source_active"] else "Inactive"],
        ["Meeting Minutes XPath", PRAC_SECTION_XPATH],
        ["Meeting PDF XPath", PRAC_MEETING_XPATH],
        ["Safety trigger", str(trigger.get("name") or "Not available")],
        [
            "Safety terms",
            ", ".join(str(term) for term in trigger.get("safety_terms", [])),
        ],
    ]
    configuration_table = Table(
        configuration_data,
        colWidths=[48 * mm, 124 * mm],
    )
    configuration_table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7.5),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#eef3f9")),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#d4dce8")),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(configuration_table)
    story.append(Spacer(1, 8))

    story.append(
        Paragraph(
            "TRACEABILITY",
            styles["SectionTitle"],
        )
    )

    trace_data = [
        [
            "PRAC section XPath",
            PRAC_SECTION_XPATH,
        ],
        [
            "Meeting PDF XPath",
            PRAC_MEETING_XPATH,
        ],
        [
            "Document page count",
            str(page_count),
        ],
        [
            "Meeting date extraction",
            (
                "Extracted from the discovered meeting-minute "
                "document title/filename when available."
            ),
        ],
        [
            "Source discovery",
            (
                "The configured source URL is crawled and all "
                "matching PRAC meeting-minute PDF links are discovered."
            ),
        ],
    ]

    trace_table = Table(
        trace_data,
        colWidths=[48 * mm, 124 * mm],
    )

    trace_table.setStyle(
        TableStyle(
            [
                (
                    "FONTNAME",
                    (0, 0),
                    (0, -1),
                    "Helvetica-Bold",
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    7.5,
                ),
                (
                    "LEADING",
                    (0, 0),
                    (-1, -1),
                    10,
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "TOP",
                ),
                (
                    "BACKGROUND",
                    (0, 0),
                    (0, -1),
                    colors.HexColor("#eef3f9"),
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.25,
                    colors.HexColor("#d4dce8"),
                ),
                (
                    "LEFTPADDING",
                    (0, 0),
                    (-1, -1),
                    5,
                ),
                (
                    "RIGHTPADDING",
                    (0, 0),
                    (-1, -1),
                    5,
                ),
                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    4,
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    4,
                ),
            ]
        )
    )

    story.append(trace_table)
    story.append(Spacer(1, 8))

    story.append(
        Paragraph(
            "This report is a review-support artifact. "
            "It summarizes extracted evidence and comparison results "
            "and does not make a final medical or regulatory decision.",
            styles["BodySmall"],
        )
    )

    doc.build(story)

    return output_path


# ---------------------------------------------------------
# History
# ---------------------------------------------------------
@app.get("/api/history")
def history():
    """
    Return complete audit history.

    This endpoint is intentionally tolerant of older SQLite schemas.
    It does not assume optional summary timestamp columns exist.
    """

    try:
        with get_db() as conn:

            # -------------------------------------------------
            # CRAWL HISTORY
            # -------------------------------------------------
            crawl = conn.execute(
                """
                SELECT
                    ch.*,
                    s.authority,
                    s.name AS source_name,
                    s.url AS source_url
                FROM crawl_history ch
                LEFT JOIN sources s
                    ON ch.source_id = s.id
                ORDER BY ch.id DESC
                LIMIT 100
                """
            ).fetchall()

            # -------------------------------------------------
            # EXTRACTION HISTORY
            # -------------------------------------------------
            extraction = conn.execute(
                """
                SELECT
                    eh.*,
                    d.title AS document_title,
                    d.filename,
                    d.meeting_date,
                    d.publication_date
                FROM extraction_history eh
                LEFT JOIN documents d
                    ON eh.document_id = d.id
                ORDER BY eh.id DESC
                LIMIT 100
                """
            ).fetchall()

            # -------------------------------------------------
            # DETECTION HISTORY
            # -------------------------------------------------
            detection = conn.execute(
                """
                SELECT
                    d.id AS document_id,
                    d.title AS document_title,
                    d.filename,
                    d.meeting_date,
                    d.publication_date,
                    MIN(f.created_at) AS detected_at,
                    MAX(f.created_at) AS last_detected_at,
                    COUNT(f.id) AS findings_count,
                    SUM(
                        CASE
                            WHEN UPPER(
                                COALESCE(
                                    f.comparison_status,
                                    ''
                                )
                            ) IN (
                                'NEW',
                                'NEW_INFORMATION'
                            )
                            THEN 1
                            ELSE 0
                        END
                    ) AS new_information_count
                FROM documents d
                JOIN findings f
                    ON d.id = f.document_id
                GROUP BY
                    d.id,
                    d.title,
                    d.filename,
                    d.meeting_date,
                    d.publication_date,
                    d.created_at
                ORDER BY d.id DESC
                LIMIT 100
                """
            ).fetchall()

            # -------------------------------------------------
            # SUMMARY HISTORY
            # -------------------------------------------------
            #
            # IMPORTANT:
            # Do not reference s.generated_at here because
            # older versions of the SQLite database may not
            # contain that column.
            #
            summary = conn.execute(
                """
                SELECT
                    s.*,
                    f.document_id,
                    f.page_number,
                    d.title AS document_title
                FROM summaries s
                LEFT JOIN findings f
                    ON s.finding_id = f.id
                LEFT JOIN documents d
                    ON f.document_id = d.id
                ORDER BY s.id DESC
                LIMIT 100
                """
            ).fetchall()

            # -------------------------------------------------
            # REVIEW HISTORY
            # -------------------------------------------------
            reviews = conn.execute(
                """
                SELECT
                    r.*,
                    f.product_name,
                    f.active_substance,
                    f.safety_topic,
                    f.document_id,
                    f.page_number,
                    d.title AS document_title,
                    d.filename
                FROM reviews r
                LEFT JOIN findings f
                    ON r.finding_id = f.id
                LEFT JOIN documents d
                    ON f.document_id = d.id
                ORDER BY r.id DESC
                LIMIT 100
                """
            ).fetchall()

        return {
            "monitoring_config": {
                "section_xpath": PRAC_SECTION_XPATH,
                "meeting_pdf_xpath": PRAC_MEETING_XPATH,
                "source_strategy": (
                    "Crawl the configured source URL and discover "
                    "all matching PRAC meeting-minute PDF links."
                )
            },

            "crawl_history": [
                dict(row)
                for row in crawl
            ],

            "extraction_history": [
                dict(row)
                for row in extraction
            ],

            "detection_history": [
                dict(row)
                for row in detection
            ],

            "summary_history": [
                dict(row)
                for row in summary
            ],

            "review_history": [
                dict(row)
                for row in reviews
            ]
        }

    except Exception as exc:
        logger.exception("History API failed")

        raise HTTPException(
            status_code=500,
            detail=f"History API failed: {str(exc)}"
        )


@app.get("/api/audit-history")
def get_audit_history_page(
    section: str = Query(..., pattern="^(configuration|extraction|crawled|actions)$"),
    page: int = Query(1, ge=1),
    page_size: int = Query(4, ge=1, le=50),
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    source_id: Optional[int] = Query(None, ge=1),
    type_filter: Optional[str] = None,
    search: Optional[str] = None,
):
    for value, field in ((date_from, "date_from"), (date_to, "date_to")):
        if value:
            try:
                datetime.strptime(value, "%Y-%m-%d")
            except ValueError as exc:
                raise HTTPException(
                    status_code=400,
                    detail=f"{field} must use YYYY-MM-DD format.",
                ) from exc
    if date_from and date_to and date_from > date_to:
        raise HTTPException(status_code=400, detail="date_from must not be after date_to.")

    search_term = f"%{search.strip()}%" if search and search.strip() else None
    where = []
    parameters = []
    if section == "actions":
        where.append("history_action.action IS NOT NULL")

    if date_from:
        timestamp_column = {
            "configuration": "history_config.timestamp",
            "extraction": "COALESCE(eh.completed_at, eh.started_at)",
            "crawled": "COALESCE(ch.completed_at, ch.started_at)",
            "actions": "history_action.timestamp",
        }[section]
        where.append(f"DATE(COALESCE({timestamp_column}, '')) >= DATE(?)")
        parameters.append(date_from)
    if date_to:
        timestamp_column = {
            "configuration": "history_config.timestamp",
            "extraction": "COALESCE(eh.completed_at, eh.started_at)",
            "crawled": "COALESCE(ch.completed_at, ch.started_at)",
            "actions": "history_action.timestamp",
        }[section]
        where.append(f"DATE(COALESCE({timestamp_column}, '')) <= DATE(?)")
        parameters.append(date_to)
    if source_id and section in {"extraction", "crawled"}:
        source_column = "d.source_id" if section == "extraction" else "ch.source_id"
        where.append(f"{source_column} = ?")
        parameters.append(source_id)
    elif source_id and section in {"configuration", "actions"}:
        source_column = "history_config.source_id" if section == "configuration" else "history_action.source_id"
        where.append(f"{source_column} = ?")
        parameters.append(source_id)
    if type_filter:
        type_column = "history_config.type" if section == "configuration" else "history_action.action"
        if section in {"configuration", "actions"}:
            where.append(f"{type_column} = ?")
            parameters.append(type_filter)
    if search_term:
        search_columns = {
            "configuration": "history_config.type LIKE ? OR history_config.details LIKE ? OR history_config.reference_id LIKE ? OR history_config.entity_type LIKE ?",
            "extraction": "s.reference_no LIKE ? OR d.reference_id LIKE ? OR d.title LIKE ? OR d.content_type LIKE ?",
            "crawled": "s.reference_no LIKE ? OR s.name LIKE ? OR s.authority LIKE ? OR ch.status LIKE ? OR ch.error_message LIKE ?",
            "actions": "history_action.reference_id LIKE ? OR history_action.action LIKE ? OR history_action.ai_model LIKE ? OR history_action.result LIKE ? OR history_action.details LIKE ?",
        }[section]
        where.append(f"({search_columns})")
        parameters.extend([search_term] * search_columns.count("?"))

    queries = {
        "configuration": (
            """
            SELECT history_config.id, history_config.timestamp,
                   history_config.type, history_config.details,
                   history_config.reference_id, history_config.source_id,
                   history_config.status
            FROM (
                SELECT ae.id, ae.created_at AS timestamp,
                       ae.event_type AS type, ae.details, ae.reference_id,
                       ae.source_id, ae.status, ae.entity_type
                FROM audit_events ae
                UNION ALL
                SELECT -1000000000 - s.id AS id, s.created_at AS timestamp,
                       'Source Created' AS type,
                       COALESCE(s.reference_no, s.name) AS details,
                       s.reference_no AS reference_id, s.id AS source_id,
                       'COMPLETED' AS status, 'source' AS entity_type
                FROM sources s
                WHERE s.created_at IS NOT NULL
                  AND NOT EXISTS (
                      SELECT 1 FROM audit_events existing
                      WHERE existing.entity_type = 'source'
                        AND existing.entity_id = s.id
                        AND existing.event_type = 'Source Created'
                  )
                UNION ALL
                SELECT -2000000000 - pc.id AS id, pc.created_at AS timestamp,
                       'Product Class Added' AS type,
                       'Added ' || pc.class_name AS details,
                       pc.reference_no AS reference_id, NULL AS source_id,
                       'COMPLETED' AS status, 'product_class' AS entity_type
                FROM product_classes pc
                WHERE pc.created_at IS NOT NULL
                  AND NOT EXISTS (
                      SELECT 1 FROM audit_events existing
                      WHERE existing.entity_type = 'product_class'
                        AND existing.entity_id = pc.id
                        AND existing.event_type = 'Product Class Added'
                  )
                UNION ALL
                SELECT -3000000000 - p.id AS id, p.created_at AS timestamp,
                       'Product Added' AS type,
                       'Added ' || p.product_name || ' to ' || pc.class_name AS details,
                       p.reference_no AS reference_id, NULL AS source_id,
                       'COMPLETED' AS status, 'product' AS entity_type
                FROM products p
                JOIN product_classes pc ON pc.id = p.product_class_id
                WHERE p.created_at IS NOT NULL
                  AND NOT EXISTS (
                      SELECT 1 FROM audit_events existing
                      WHERE existing.entity_type = 'product'
                        AND existing.entity_id = p.id
                        AND existing.event_type = 'Product Added'
                  )
            ) history_config
            """,
            "timestamp",
        ),
        "extraction": (
            """
            SELECT eh.id, COALESCE(eh.completed_at, eh.started_at) AS timestamp,
                   s.reference_no AS source_reference, d.source_id,
                   d.reference_id, d.title AS document_title,
                   d.content_type, eh.status, eh.pages_extracted,
                   eh.error_message
            FROM extraction_history eh
            LEFT JOIN documents d ON d.id = eh.document_id
            LEFT JOIN sources s ON s.id = d.source_id
            """,
            "timestamp",
        ),
        "crawled": (
            """
            SELECT ch.id, COALESCE(ch.completed_at, ch.started_at) AS timestamp,
                   ch.source_id, s.reference_no AS source_reference,
                   s.name AS source_name, ch.status, ch.documents_found,
                   ch.new_documents, ch.total_documents, ch.error_message,
                   (SELECT COUNT(*) FROM crawl_history prior
                    WHERE prior.source_id = ch.source_id AND prior.id <= ch.id) AS crawl_count
            FROM crawl_history ch
            LEFT JOIN sources s ON s.id = ch.source_id
            """,
            "timestamp",
        ),
        "actions": (
            """
            SELECT history_action.id, history_action.timestamp,
                   history_action.reference_id, history_action.action,
                   history_action.ai_model, history_action.result,
                   history_action.status, history_action.details,
                   history_action.source_id
            FROM (
                SELECT ae.id, ae.created_at AS timestamp, ae.reference_id,
                       ae.action, ae.ai_model, ae.result, ae.status,
                       ae.details, ae.source_id
                FROM audit_events ae
                WHERE ae.action IS NOT NULL
                UNION ALL
                SELECT -analysis.id AS id, analysis.updated_at AS timestamp,
                       d.reference_id, 'AI Analysis (Summary)' AS action,
                       analysis.provider AS ai_model,
                       CASE WHEN analysis.status = 'COMPLETED'
                            THEN 'Summary generated'
                            ELSE COALESCE(analysis.error_message, analysis.status)
                       END AS result,
                       analysis.status,
                       'Previously stored AI analysis' AS details,
                       d.source_id
                FROM document_ai_analyses analysis
                JOIN documents d ON d.id = analysis.document_id
                WHERE analysis.provider IN ('ollama', 'openai')
                  AND NOT EXISTS (
                    SELECT 1 FROM audit_events prior
                    WHERE prior.entity_type = 'document'
                      AND prior.entity_id = analysis.document_id
                      AND prior.action = 'AI Analysis (Summary)'
                )
            ) history_action
            """,
            "timestamp",
        ),
    }
    select_sql, order_column = queries[section]
    where_clause = " WHERE " + " AND ".join(where) if where else ""
    with get_db() as conn:
        total = conn.execute(
            f"SELECT COUNT(*) AS total FROM ({select_sql}{where_clause}) history_rows",
            parameters,
        ).fetchone()["total"]
        max_page = max(1, (total + page_size - 1) // page_size)
        current_page = min(page, max_page)
        offset = (current_page - 1) * page_size
        records = conn.execute(
            f"SELECT * FROM ({select_sql}{where_clause}) history_rows ORDER BY {order_column} DESC, id DESC LIMIT ? OFFSET ?",
            [*parameters, page_size, offset],
        ).fetchall()
        type_options = []
        if section in {"configuration", "actions"}:
            option_field = "type" if section == "configuration" else "action"
            options_sql = f"""
                SELECT DISTINCT {option_field} AS option
                FROM ({select_sql}) history_options
                WHERE {option_field} IS NOT NULL
                ORDER BY {option_field}
            """
            type_options = [
                row["option"]
                for row in conn.execute(options_sql).fetchall()
                if row["option"]
            ]

    return {
        "section": section,
        "records": [dict(row) for row in records],
        "total": total,
        "page": current_page,
        "page_size": page_size,
        "pages": max_page,
        "type_options": type_options,
    }



# ---------------------------------------------------------
# Summary PDF
# ---------------------------------------------------------

@app.get("/api/documents/{document_id}/summary-pdf")
def view_summary_pdf(document_id: int):

    try:

        pdf_path = create_document_summary_pdf(
            document_id
        )

    except ValueError as exc:

        raise HTTPException(
            status_code=404,
            detail=str(exc),
        )

    except Exception as exc:

        logger.exception(
            "Failed to create summary PDF for document %s",
            document_id,
        )

        raise HTTPException(
            status_code=500,
            detail=str(exc),
        )

    return FileResponse(
        path=str(pdf_path),
        media_type="application/pdf",
        filename=pdf_path.name,
        content_disposition_type="inline",
    )


@app.get("/api/documents/{document_id}/summary-pdf/download")
def download_summary_pdf(document_id: int):

    try:

        pdf_path = create_document_summary_pdf(
            document_id
        )

    except ValueError as exc:

        raise HTTPException(
            status_code=404,
            detail=str(exc),
        )

    except Exception as exc:

        logger.exception(
            "Failed to create downloadable summary PDF "
            "for document %s",
            document_id,
        )

        raise HTTPException(
            status_code=500,
            detail=str(exc),
        )

    return FileResponse(
        path=str(pdf_path),
        media_type="application/pdf",
        filename=pdf_path.name,
        content_disposition_type="attachment",
    )


# ---------------------------------------------------------
# Monitoring
# ---------------------------------------------------------

monitoring_jobs = {}


@app.post("/api/monitoring/run")
def start_monitoring(background_tasks: BackgroundTasks):

    job_id = str(uuid.uuid4())

    monitoring_jobs[job_id] = {
        "job_id": job_id,
        "status": "queued",
        "progress": 0,
        "message": "Monitoring job queued.",
        "documents_found": 0,
        "new_documents": 0,
        "findings_found": 0,
        "summary_pdfs_created": 0,
        "crawl_documents_discovered": 0,
        "pages_extracted": 0,
        "new_information": 0,
        "total_sources": 0,
        "sources_processed": 0,
        "sources_failed": 0,
        "sources_without_new_documents": 0,
    }

    background_tasks.add_task(
        monitoring_worker,
        job_id,
    )

    return {
        "job_id": job_id,
        "status": "queued",
        "message": "Monitoring started.",
    }


@app.get("/api/monitoring/status/{job_id}")
def monitoring_status(job_id: str):

    job = monitoring_jobs.get(job_id)

    if not job:
        raise HTTPException(
            status_code=404,
            detail="Monitoring job not found.",
        )

    return job


# ---------------------------------------------------------
# Monitoring Worker
# ---------------------------------------------------------

def monitoring_worker(job_id: str):

    job = monitoring_jobs[job_id]

    try:

        job["status"] = "running"
        job["progress"] = 5
        job["message"] = "Loading active sources..."

        with get_db() as conn:
            sources = conn.execute(
                """
                SELECT *
                FROM sources
                WHERE active = 1
                ORDER BY id
                """
            ).fetchall()
            sources = [dict(source) for source in sources]

        job["total_sources"] = len(sources)
        if not sources:
            job["status"] = "completed"
            job["progress"] = 100
            job["message"] = "No active sources configured."
            return

        total_documents_found = 0
        new_documents = 0
        total_findings = 0

        for source in sources:

            source_id = source["id"]

            job["progress"] = 10
            job["message"] = (
                f"Crawling {source['name']}..."
            )

            crawl_started = datetime.utcnow().isoformat()
            source_new_documents = 0

            try:
                result = crawl_source(
                    str(source["url"]),
                    source_id=source["id"],
                    content_types=normalize_content_types(source.get("content_types")),
                    source_name=source.get("name"),
                    reference_no=source.get("reference_no"),
                )
                if not result.get("success"):
                    raise RuntimeError(
                        result.get("error") or "Source crawl failed."
                    )

                documents = result.get(
                    "documents",
                    []
                )

                total_documents_found += len(documents)
                job["crawl_documents_discovered"] = (
                    job.get("crawl_documents_discovered", 0)
                    + len(documents)
                )

                crawl_status = "success"
                crawl_error = None

            except Exception as exc:

                logger.exception(
                    "Crawler failed for source %s",
                    source_id,
                )

                documents = []
                crawl_status = "failed"
                crawl_error = str(exc)

            crawl_completed = datetime.utcnow().isoformat()

            with get_db() as conn:
                history_cursor = conn.execute(
                    """
                    INSERT INTO crawl_history
                    (
                        source_id,
                        started_at,
                        completed_at,
                        status,
                        documents_found,
                        new_documents,
                        total_documents,
                        error_message
                    )
                    VALUES (
                        ?, ?, ?, ?, ?, 0,
                        (SELECT COUNT(*) FROM documents WHERE source_id = ?),
                        ?
                    )
                    """,
                    (
                        source_id,
                        crawl_started,
                        crawl_completed,
                        crawl_status,
                        len(documents),
                        source_id,
                        crawl_error,
                    ),
                )
                crawl_history_id = history_cursor.lastrowid

            # -------------------------------------------------
            # Process discovered documents
            # -------------------------------------------------

            for discovered in documents:

                url = discovered.get("url")

                if not url:
                    continue

                title = discovered.get(
                    "title",
                    "PRAC Meeting Minutes",
                )

                filename = discovered.get(
                    "filename"
                )

                meeting_date = discovered.get(
                    "meeting_date"
                )

                publication_date = discovered.get(
                    "publication_date"
                )

                content_type = discovered.get("content_type")
                discovered_date = datetime.utcnow().isoformat()

                with get_db() as conn:
                    existing_document_id = None

                    existing = conn.execute(
                        """
                        SELECT *
                        FROM documents
                        WHERE lower(rtrim(url, '/')) = lower(rtrim(?, '/'))
                        """,
                        (url,),
                    ).fetchone()

                    if existing:
                        previous_generated_reference = _document_reference_id(
                            source["reference_no"],
                            existing["publication_date"],
                            int(existing["id"]),
                        )
                        reference_id = existing["reference_id"]
                        if (
                            not reference_id
                            or reference_id == previous_generated_reference
                        ):
                            reference_id = _document_reference_id(
                                source["reference_no"],
                                publication_date,
                                int(existing["id"]),
                            )
                        conn.execute(
                            """
                            UPDATE documents
                            SET title = COALESCE(?, title),
                                publication_date = ?,
                                meeting_date = COALESCE(?, meeting_date),
                                content_type = COALESCE(?, content_type),
                                last_crawled_at = ?,
                                reference_id = ?
                            WHERE id = ?
                            """,
                            (
                                title,
                                publication_date,
                                meeting_date,
                                content_type,
                                discovered_date,
                                reference_id,
                                existing["id"],
                            ),
                        )
                        if str(existing["status"] or "").upper() in {
                            "PROCESSED",
                            "EXTRACTED",
                        }:
                            continue
                        existing_document_id = int(existing["id"])

                    # -----------------------------------------
                    # Download PDF
                    # -----------------------------------------

                    job["message"] = (
                        f"Downloading {title}..."
                    )

                    try:

                        download_result = download_pdf(
                            url,
                            PDF_DIR,
                        )

                        if isinstance(download_result, dict):
                            downloaded_path = download_result.get("path")
                            downloaded_filename = download_result.get(
                                "filename"
                            )
                            downloaded_sha256 = download_result.get(
                                "document_hash"
                            )
                        else:
                            downloaded_path = download_result
                            downloaded_filename = None
                            downloaded_sha256 = None

                    except Exception as exc:

                        logger.exception(
                            "PDF download failed: %s",
                            url,
                        )
                        if existing_document_id is not None:
                            conn.execute(
                                "UPDATE documents SET status = 'FAILED' WHERE id = ?",
                                (existing_document_id,),
                            )

                        continue

                    if not downloaded_path:
                        if existing_document_id is not None:
                            conn.execute(
                                "UPDATE documents SET status = 'FAILED' WHERE id = ?",
                                (existing_document_id,),
                            )
                        continue

                    downloaded_path = Path(
                        downloaded_path
                    )

                    if not downloaded_path.exists():
                        if existing_document_id is not None:
                            conn.execute(
                                "UPDATE documents SET status = 'FAILED' WHERE id = ?",
                                (existing_document_id,),
                            )
                        continue

                    filename = (
                        downloaded_filename
                        or downloaded_path.name
                        or filename
                    )

                    # -----------------------------------------
                    # Calculate hash
                    # -----------------------------------------

                    try:

                        if downloaded_sha256:
                            sha256 = downloaded_sha256
                        else:
                            import hashlib

                            sha256_hash = hashlib.sha256()

                            with open(
                                downloaded_path,
                                "rb",
                            ) as pdf_file:

                                for chunk in iter(
                                    lambda: pdf_file.read(1024 * 1024),
                                    b"",
                                ):
                                    sha256_hash.update(chunk)

                            sha256 = sha256_hash.hexdigest()

                    except Exception as exc:

                        logger.exception(
                            "Hash calculation failed"
                        )
                        if existing_document_id is not None:
                            conn.execute(
                                "UPDATE documents SET status = 'FAILED' WHERE id = ?",
                                (existing_document_id,),
                            )

                        continue

                    # -----------------------------------------
                    # Duplicate hash protection
                    # -----------------------------------------

                    duplicate = conn.execute(
                        """
                        SELECT id
                        FROM documents
                        WHERE sha256 = ?
                        """,
                        (sha256,),
                    ).fetchone()

                    if (
                        duplicate
                        and existing_document_id is None
                    ):
                        continue

                    # -----------------------------------------
                    # Save document metadata
                    # -----------------------------------------

                    if existing_document_id is not None:
                        document_id = existing_document_id
                        existing_reference_id = existing["reference_id"]
                        old_generated_reference = _document_reference_id(
                            source["reference_no"],
                            existing["publication_date"],
                            document_id,
                        )
                        reference_id = (
                            _document_reference_id(
                                source["reference_no"],
                                publication_date,
                                document_id,
                            )
                            if not existing_reference_id
                            or existing_reference_id == old_generated_reference
                            else existing_reference_id
                        )
                        conn.execute(
                            """
                            UPDATE documents
                            SET filename = ?,
                                sha256 = ?,
                                publication_date = ?,
                                content_type = COALESCE(?, content_type),
                                reference_id = ?,
                                status = 'NEW'
                            WHERE id = ?
                            """,
                            (
                                filename,
                                sha256,
                                publication_date,
                                content_type,
                                reference_id,
                                document_id,
                            ),
                        )
                    else:
                        cursor = conn.execute(
                            """
                            INSERT INTO documents
                            (
                                source_id,
                                title,
                                url,
                                meeting_date,
                                publication_date,
                                filename,
                                discovered_date,
                                last_crawled_at,
                                content_type,
                                sha256,
                                status,
                                created_at
                            )
                            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                            """,
                            (
                                source_id,
                                title,
                                url,
                                meeting_date,
                                publication_date,
                                filename,
                                discovered_date,
                                discovered_date,
                                content_type,
                                sha256,
                                "NEW",
                                datetime.utcnow().isoformat(),
                            ),
                        )

                        document_id = cursor.lastrowid
                        reference_id = _document_reference_id(
                            source["reference_no"],
                            publication_date,
                            document_id,
                        )
                        conn.execute(
                            "UPDATE documents SET reference_id = ? WHERE id = ?",
                            (reference_id, document_id),
                        )
                        new_documents += 1
                        source_new_documents += 1

                # ---------------------------------------------
                # Extract PDF
                # ---------------------------------------------

                job["message"] = (
                    f"Extracting {title}..."
                )

                extraction_started = datetime.utcnow().isoformat()

                try:

                    pages = process_pdf_safe(
                        downloaded_path
                    )

                    extraction_status = "success"
                    extraction_error = None

                except Exception as exc:

                    logger.exception(
                        "PDF extraction failed for document %s",
                        document_id,
                    )

                    pages = []
                    extraction_status = "failed"
                    extraction_error = str(exc)

                pages_extracted = 0

                # ---------------------------------------------
                # Store pages
                # ---------------------------------------------

                with get_db() as conn:

                    for page in pages:

                        page_number = page.get(
                            "page_number"
                        )

                        page_text = page.get(
                            "text",
                            "",
                        )

                        if page_number is None:
                            continue

                        conn.execute(
                            """
                            INSERT OR REPLACE INTO document_pages
                            (
                                document_id,
                                page_number,
                                page_text,
                                created_at
                            )
                            VALUES (?, ?, ?, ?)
                            """,
                            (
                                document_id,
                                page_number,
                                page_text,
                                datetime.utcnow().isoformat(),
                            ),
                        )

                        pages_extracted += 1

                    conn.execute(
                        """
                        UPDATE documents
                        SET status = ?
                        WHERE id = ?
                        """,
                        (
                            "EXTRACTED"
                            if extraction_status == "success"
                            else "EXTRACTION_FAILED",
                            document_id,
                        ),
                    )

                    conn.execute(
                        """
                        INSERT INTO extraction_history
                        (
                            document_id,
                            started_at,
                            completed_at,
                            status,
                            pages_extracted,
                            error_message
                        )
                        VALUES (?, ?, ?, ?, ?, ?)
                        """,
                        (
                            document_id,
                            extraction_started,
                            datetime.utcnow().isoformat(),
                            extraction_status,
                            pages_extracted,
                            extraction_error,
                        ),
                    )

                # ---------------------------------------------
                # Detect safety findings
                # ---------------------------------------------

                if pages:

                    job["message"] = (
                        f"Detecting safety information in {title}..."
                    )

                    try:

                        findings = detect_document(
                            pages,
                            load_trigger_config(),
                        )

                    except Exception as exc:

                        logger.exception(
                            "Detection failed for document %s",
                            document_id,
                        )

                        findings = []

                    findings = remove_duplicate_findings(
                        findings
                    )

                    for finding in findings:

                        # Make sure document ID exists
                        finding["document_id"] = document_id

                        try:
                            comparison = compare_finding(
                                finding
                            )

                            if isinstance(
                                comparison,
                                dict,
                            ):
                                comparison_status = (
                                    comparison.get(
                                        "comparison_status",
                                        "NEW",
                                    )
                                )
                            else:
                                comparison_status = "NEW"

                        except Exception:

                            logger.exception(
                                "Comparison failed"
                            )

                            comparison_status = "NEW"

                        finding[
                            "comparison_status"
                        ] = comparison_status

                        with get_db() as conn:

                            cursor = conn.execute(
                                """
                                INSERT INTO findings
                                (
                                    document_id,
                                    page_number,
                                    product_name,
                                    active_substance,
                                    trigger_name,
                                    safety_topic,
                                    relevant_terms,
                                    evidence,
                                    section,
                                    comparison_status,
                                    evidence_hash,
                                    created_at
                                )
                                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                                """,
                                (
                                    document_id,
                                    finding.get(
                                        "page_number"
                                    ),
                                    finding.get(
                                        "product_name"
                                    ),
                                    finding.get(
                                        "active_substance"
                                    ),
                                    finding.get(
                                        "trigger_name",
                                        load_trigger_config().get(
                                            "name",
                                            "Benzodiazepines",
                                        ),
                                    ),
                                    finding.get(
                                        "safety_topic"
                                    ),
                                    finding.get(
                                        "relevant_terms"
                                    ),
                                    finding.get(
                                        "evidence"
                                    ),
                                    finding.get(
                                        "section"
                                    ),
                                    comparison_status,
                                    finding.get(
                                        "evidence_hash"
                                    ),
                                    datetime.utcnow().isoformat(),
                                ),
                            )

                            finding_id = cursor.lastrowid

                        total_findings += 1

                job["findings_found"] = total_findings
                job["pages_extracted"] = (
                    job.get("pages_extracted", 0)
                    + pages_extracted
                )

                # -------------------------------------------------
                # Create document summary PDF
                # -------------------------------------------------
                try:
                    summary_pdf_path = create_document_summary_pdf(
                        document_id
                    )
                    job["summary_pdfs_created"] = (
                        job.get("summary_pdfs_created", 0) + 1
                    )
                    logger.info(
                        "Summary PDF created for document %s: %s",
                        document_id,
                        summary_pdf_path,
                    )
                except Exception:
                    logger.exception(
                        "Summary PDF generation failed for document %s",
                        document_id,
                    )

            with get_db() as conn:
                conn.execute(
                    """
                    UPDATE crawl_history
                    SET new_documents = ?,
                        total_documents = (
                            SELECT COUNT(*) FROM documents WHERE source_id = ?
                        )
                    WHERE id = ?
                    """,
                    (source_new_documents, source_id, crawl_history_id),
                )
                conn.execute(
                    """
                    UPDATE sources
                    SET last_crawled_at = ?, last_crawl_status = ?,
                        last_error_message = ?, updated_at = ?
                    WHERE id = ?
                    """,
                    (
                        crawl_completed,
                        crawl_status,
                        crawl_error,
                        crawl_completed,
                        source_id,
                    ),
                )

            job["sources_processed"] = job.get("sources_processed", 0) + 1
            if crawl_status == "failed":
                job["sources_failed"] = job.get("sources_failed", 0) + 1
            elif source_new_documents == 0:
                job["sources_without_new_documents"] = (
                    job.get("sources_without_new_documents", 0) + 1
                )

        # -----------------------------------------------------
        # Finished
        # -----------------------------------------------------

        job["status"] = "completed"
        job["progress"] = 100
        job["documents_found"] = total_documents_found
        job["new_documents"] = new_documents
        job["findings_found"] = total_findings

        with get_db() as conn:
            job["new_information"] = conn.execute(
                """
                SELECT COUNT(*) AS count
                FROM findings
                WHERE UPPER(COALESCE(comparison_status, ''))
                    IN ('NEW', 'NEW_INFORMATION')
                """
            ).fetchone()["count"]

        job["message"] = (
            "Monitoring completed successfully."
        )

        logger.info(
            "Monitoring job %s completed.",
            job_id,
        )

    except Exception as exc:

        logger.exception(
            "Monitoring job %s failed.",
            job_id,
        )

        job["status"] = "failed"
        job["message"] = str(exc)
        job["error"] = str(exc)


# ---------------------------------------------------------
# Manual document processing
# ---------------------------------------------------------

@app.post("/api/documents/{document_id}/process")
def process_document_endpoint(
    document_id: int,
):

    with get_db() as conn:
        document = conn.execute(
            """
            SELECT *
            FROM documents
            WHERE id = ?
            """,
            (document_id,),
        ).fetchone()

    if not document:
        raise HTTPException(
            status_code=404,
            detail="Document not found.",
        )

    filename = document["filename"]

    if not filename:
        raise HTTPException(
            status_code=404,
            detail="Document PDF filename missing.",
        )

    pdf_path = PDF_DIR / filename

    if not pdf_path.exists():
        raise HTTPException(
            status_code=404,
            detail="PDF file not found.",
        )

    try:

        processed = process_pdf_safe(pdf_path)
        if not processed.get("success"):
            raise HTTPException(
                status_code=422,
                detail=processed.get("error") or "PDF text extraction failed.",
            )
        pages = processed.get("pages", [])

        with get_db() as conn:

            for page in pages:

                conn.execute(
                    """
                    INSERT OR REPLACE INTO document_pages
                    (
                        document_id,
                        page_number,
                        page_text,
                        created_at
                    )
                    VALUES (?, ?, ?, ?)
                    """,
                    (
                        document_id,
                        page["page_number"],
                        page.get("text", ""),
                        datetime.utcnow().isoformat(),
                    ),
                )

            conn.execute(
                """
                UPDATE documents
                SET status = 'EXTRACTED'
                WHERE id = ?
                """,
                (document_id,),
            )

        findings = detect_document(
            pages,
            load_trigger_config(),
        )

        findings = remove_duplicate_findings(
            findings
        )

        inserted = 0

        with get_db() as conn:

            for finding in findings:

                finding["document_id"] = document_id

                try:
                    comparison = compare_finding(
                        finding
                    )

                    if isinstance(
                        comparison,
                        dict,
                    ):
                        comparison_status = comparison.get(
                            "comparison_status",
                            "NEW",
                        )
                    else:
                        comparison_status = "NEW"

                except Exception:
                    comparison_status = "NEW"

                conn.execute(
                    """
                    INSERT INTO findings
                    (
                        document_id,
                        page_number,
                        product_name,
                        active_substance,
                        trigger_name,
                        safety_topic,
                        relevant_terms,
                        evidence,
                        section,
                        comparison_status,
                        evidence_hash,
                        created_at
                    )
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        document_id,
                        finding.get("page_number"),
                        finding.get("product_name"),
                        finding.get("active_substance"),
                        finding.get(
                            "trigger_name",
                            load_trigger_config().get(
                                "name",
                                "Benzodiazepines",
                            ),
                        ),
                        finding.get("safety_topic"),
                        finding.get("relevant_terms"),
                        finding.get("evidence"),
                        finding.get("section"),
                        comparison_status,
                        finding.get("evidence_hash"),
                        datetime.utcnow().isoformat(),
                    ),
                )

                inserted += 1

        # Create the same review-support summary PDF
        # for manually processed documents.
        summary_pdf_path = create_document_summary_pdf(
            document_id
        )

        return {
            "success": True,
            "document_id": document_id,
            "pages_extracted": len(pages),
            "findings_created": inserted,
            "summary_pdf": str(summary_pdf_path),
            "summary_pdf_view_url": (
                f"/api/documents/{document_id}/summary-pdf"
            ),
            "summary_pdf_download_url": (
                f"/api/documents/{document_id}/summary-pdf/download"
            ),
        }

    except Exception as exc:

        logger.exception(
            "Manual document processing failed."
        )

        raise HTTPException(
            status_code=500,
            detail=str(exc),
        )


# ---------------------------------------------------------
# Storage
# ---------------------------------------------------------

@app.get("/api/storage")
def storage_info():

    pdf_files = list(
        PDF_DIR.glob("*.pdf")
    )

    total_size = sum(
        file.stat().st_size
        for file in pdf_files
        if file.exists()
    )

    return {
        "pdf_directory": str(PDF_DIR),
        "pdf_count": len(pdf_files),
        "total_bytes": total_size,
    }


# ---------------------------------------------------------
# Run locally
# ---------------------------------------------------------

if __name__ == "__main__":
    import uvicorn

    uvicorn.run(
        app,
        host="0.0.0.0",
        port=8000,
    )